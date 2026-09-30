"""Uploaded files are hostile data. Each case here is what the reader does today with a package built to hurt it,
measured on 2026-09-29 before it was written down: refused with the reason, or read within a bound, never a crash
and never a network call. A change in any of these is a change in the security envelope and must be deliberate."""

from __future__ import annotations

import io
import time
import zipfile
from collections.abc import Callable

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from app.ingest import UnsupportedFile, ingest


def _base() -> bytes:
    document = DocxDocument()
    document.add_paragraph("1. Term")
    document.add_paragraph("This Agreement commences on the Effective Date.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _rebuild(mutate: Callable[[str, bytes], bytes | None], extra: Callable[[zipfile.ZipFile], None] | None = None) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(_base()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as package:
        for name in source.namelist():
            data = source.read(name)
            package.writestr(name, mutate(name, data) or data)
        if extra:
            extra(package)
    return out.getvalue()


ENTITY_BOMB = (
    '<!DOCTYPE x [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;"><!ENTITY c "&b;&b;&b;&b;&b;&b;&b;&b;&b;&b;">'
    '<!ENTITY d "&c;&c;&c;&c;&c;&c;&c;&c;&c;&c;"><!ENTITY e "&d;&d;&d;&d;&d;&d;&d;&d;&d;&d;"><!ENTITY f "&e;&e;&e;&e;&e;&e;&e;&e;&e;&e;">]>'
)


def test_an_entity_bomb_in_the_document_part_is_refused_at_once_not_expanded() -> None:
    def bomb(name: str, data: bytes) -> bytes | None:
        if name != "word/document.xml":
            return None
        text = data.decode("utf-8")
        return text.replace("?>", "?>" + ENTITY_BOMB, 1).replace("Effective Date", "Effective &f; Date", 1).encode("utf-8")

    started = time.perf_counter()
    with pytest.raises(UnsupportedFile, match=r"not a readable .docx package"):
        ingest("bomb.docx", _rebuild(bomb))
    assert time.perf_counter() - started < 2, "refused before any expansion"


def test_absurd_nesting_in_the_body_is_refused_not_walked() -> None:
    def deep(name: str, data: bytes) -> bytes | None:
        if name != "word/document.xml":
            return None
        nest = "<w:sdt><w:sdtContent>" * 3000 + "<w:p><w:r><w:t>deep</w:t></w:r></w:p>" + "</w:sdtContent></w:sdt>" * 3000
        return data.decode("utf-8").replace("</w:body>", nest + "</w:body>", 1).encode("utf-8")

    started = time.perf_counter()
    with pytest.raises(UnsupportedFile):
        ingest("deep.docx", _rebuild(deep))
    assert time.perf_counter() - started < 2


def test_a_package_with_twenty_thousand_members_is_read_within_a_bound() -> None:
    def members(package: zipfile.ZipFile) -> None:
        for i in range(20_000):
            package.writestr(f"word/media/x{i}.txt", b"x")

    started = time.perf_counter()
    parsed = ingest("many.docx", _rebuild(lambda name, data: None, members))
    assert [s.text for s in parsed.sections] and "Effective Date" in " ".join(s.text for s in parsed.sections)
    assert time.perf_counter() - started < 5, "member count alone does not stall the reader (0.5 s measured)"


def test_hostile_uploads_reach_the_api_as_422_with_the_reason_and_nothing_stored(client: TestClient) -> None:
    def bomb(name: str, data: bytes) -> bytes | None:
        if name != "word/document.xml":
            return None
        return data.decode("utf-8").replace("?>", "?>" + ENTITY_BOMB, 1).replace("Effective Date", "&f;", 1).encode("utf-8")

    response = client.post("/api/documents", files={"file": ("bomb.docx", _rebuild(bomb), "application/octet-stream")})
    assert response.status_code == 422 and "not a readable .docx package" in response.json()["detail"]
    assert client.get("/api/documents").json() == []
