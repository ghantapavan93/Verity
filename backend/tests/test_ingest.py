import io

import pytest
from docx import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.oxml.xmlchemy import BaseOxmlElement

from app.ingest import UnsupportedFile, ingest


def make_docx() -> bytes:
    document = DocxDocument()
    document.add_paragraph("Master Services Agreement")
    document.add_heading("Term", level=1)
    document.add_paragraph("This Agreement commences on the Effective Date.")
    document.add_heading("Termination", level=1)
    document.add_heading("Termination for Convenience", level=2)
    document.add_paragraph("Customer may terminate this Agreement upon fifteen (15) days’ written notice to Provider.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Fee"
    table.rows[0].cells[1].text = "USD 1,000"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _revision(tag: str, rev_id: int, inner: str = "") -> BaseOxmlElement:
    """A w:ins / w:del / w:moveFrom / w:moveTo element, optionally wrapping runs."""
    return parse_xml(f'<w:{tag} {nsdecls("w")} w:id="{rev_id}" w:author="Reviewer" w:date="2026-09-28T09:00:00Z">{inner}</w:{tag}>')


def make_redlined_docx() -> bytes:
    """Negotiated paper: a word swapped under track changes, a hidden drafting note, a clause deleted
    whole, a paragraph mark deleted so two paragraphs join, and a table row deleted."""
    document = DocxDocument()
    document.add_paragraph("Master Services Agreement")
    document.add_heading("Term", level=1)
    swapped = document.add_paragraph("The Subscription continues for ")
    swapped._p.append(_revision("del", 1, "<w:r><w:delText>twelve</w:delText></w:r>"))
    swapped._p.append(_revision("ins", 2, "<w:r><w:t>twenty-four</w:t></w:r>"))
    swapped.add_run(" months.")

    noted = document.add_paragraph("Renewal is automatic. ")
    noted.add_run("Hidden drafting note: check with finance.").font.hidden = True

    document.add_heading("Fees", level=1)
    gone = document.add_paragraph()
    gone._p.append(_revision("del", 3, "<w:r><w:delText>Provider may suspend the Service at any time.</w:delText></w:r>"))
    gone._p.get_or_add_pPr().append(parse_xml(f'<w:rPr {nsdecls("w")}><w:del w:id="4" w:author="Reviewer" w:date="2026-09-28T09:00:00Z"/></w:rPr>'))

    joined = document.add_paragraph("Fees are due within ")
    joined._p.get_or_add_pPr().append(parse_xml(f'<w:rPr {nsdecls("w")}><w:del w:id="5" w:author="Reviewer" w:date="2026-09-28T09:00:00Z"/></w:rPr>'))
    document.add_paragraph("thirty (30) days of invoice.")

    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Support"
    table.rows[0].cells[1].text = "USD 500"
    table.rows[1].cells[0].text = "Training"
    table.rows[1].cells[1].text = "USD 900"
    table.rows[1]._tr.get_or_add_trPr().append(_revision("del", 6))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_is_read_as_the_accepted_view() -> None:
    parsed = ingest("redline.docx", make_redlined_docx())
    by_heading = {s.heading: s.text for s in parsed.sections}

    assert "continues for twenty-four months." in by_heading["Term"], "the inserted word is kept"
    assert "twelve" not in by_heading["Term"], "the deleted word is gone"
    assert "Renewal is automatic." in by_heading["Term"]
    assert "Hidden drafting note" not in by_heading["Term"], "hidden runs are left out"

    assert "Provider may suspend" not in by_heading["Fees"], "a clause deleted whole leaves nothing behind"
    assert "Fees are due within thirty (30) days of invoice." in by_heading["Fees"], "a deleted paragraph mark joins the two paragraphs"
    assert "Support | USD 500" in by_heading["Fees"]
    assert "Training" not in by_heading["Fees"], "a deleted table row is skipped"

    assert parsed.tracked_changes == 6
    assert parsed.hidden_runs == 1


def test_the_same_bytes_parse_the_same_way() -> None:
    data = make_redlined_docx()
    first, second = ingest("a.docx", data), ingest("b.docx", data)
    assert [(s.number, s.heading, s.text) for s in first.sections] == [(s.number, s.heading, s.text) for s in second.sections]
    assert first.sha256 == second.sha256


def test_custom_heading_styles_are_read_by_name_and_by_outline_level() -> None:
    """Government and firm templates name their own heading styles; the outline level is what Word keys on."""
    document = DocxDocument()
    by_name = document.styles.add_style("SchLevel1Heading", WD_STYLE_TYPE.PARAGRAPH)
    by_level = document.styles.add_style("ClauseTitle", WD_STYLE_TYPE.PARAGRAPH)
    by_level.element.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="1"/>'))
    inherited = document.styles.add_style("ClauseTitleIndented", WD_STYLE_TYPE.PARAGRAPH)
    inherited.base_style = by_level
    document.add_paragraph("Special Term 1 - Definitions", style=by_name)
    document.add_paragraph("In these Special Terms the following words have the following meanings.")
    document.add_paragraph("Convictions", style=by_level)
    document.add_paragraph("means any previous or pending prosecutions, convictions or cautions.")
    document.add_paragraph("Locations", style=inherited)
    document.add_paragraph("means any locations, sites or territories relevant to the performance of the Deliverables.")
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = ingest("special-terms.docx", buffer.getvalue())
    assert [(s.number, s.heading) for s in parsed.sections] == [("1", "Special Term 1 - Definitions"), ("1.1", "Convictions"), ("1.2", "Locations")]
    assert parsed.sections[1].text == "means any previous or pending prosecutions, convictions or cautions."


def test_docx_without_revisions_reports_none() -> None:
    parsed = ingest("sample.docx", make_docx())
    assert (parsed.tracked_changes, parsed.hidden_runs) == (0, 0)


def test_docx_becomes_numbered_sections_with_tables_in_order() -> None:
    parsed = ingest("sample.docx", make_docx())
    assert parsed.media_type.endswith("wordprocessingml.document")
    assert [(s.number, s.heading) for s in parsed.sections] == [
        ("", "Master Services Agreement"),
        ("1", "Term"),
        ("2", "Termination"),
        ("2.1", "Termination for Convenience"),
    ]
    assert "fifteen (15) days’ written notice" in parsed.sections[3].text
    assert "Fee | USD 1,000" in parsed.sections[3].text
    assert len(parsed.sha256) == 64


def test_txt_headings_by_number_and_caps() -> None:
    text = "SERVICES AGREEMENT\n\n1. Definitions\n\nAffiliate means an entity under common control.\n\n2. Payment\n\nFees are due within thirty (30) days.\n"
    parsed = ingest("agreement.txt", text.encode("utf-8"))
    assert [(s.number, s.heading) for s in parsed.sections] == [("", "SERVICES AGREEMENT"), ("1", "Definitions"), ("2", "Payment")]
    assert parsed.sections[2].text == "Fees are due within thirty (30) days."


def test_unsupported_extension_is_rejected() -> None:
    with pytest.raises(UnsupportedFile):
        ingest("contract.xlsx", b"nope")
