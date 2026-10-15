"""Reading fixes from the retrieval study of 2026-10-14 (outside the
repository), which read ten unfamiliar agreements in four formats against hand-written headings:

1. A DOCX that uses no heading style (manual numbers on bold Normal paragraphs, or list numbering) had no heading at all:
   0 of 293 headings, one block of text cut into parts with no clause numbers.
2. A numbered title ending in a full stop ("Section 2.7 Purchase Price Adjustment.") was body: a whole Article became one
   unnumbered section.
3. "ARTICLE II" on one line and its title on the next were two headings, the first an empty section.
4. A PDF line wrapping a clause ("4.2 Any Monthly Rent not paid within ten (10) days after its due date bears a late
   charge") was a title, so its operative words were not quotable section text.

Each has controls: what was read right before is read the same.
"""

from __future__ import annotations

import io
from collections.abc import Callable

import pytest
from docx import Document as DocxDocument
from docx.document import Document as WordDocument

from app.ingest import ingest
from tests.test_docx_numbering import _number_paragraph
from tests.test_post_pilot_repair import _list_of


def _sections(name: str, data: bytes) -> list[tuple[str, str, str]]:
    return [(s.number, s.heading, s.text) for s in ingest(name, data).sections]


def _docx(build: Callable[[WordDocument], None]) -> bytes:
    document = DocxDocument()
    build(document)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_a_docx_with_manual_numbers_and_no_heading_style_has_its_headings() -> None:
    def build(document: WordDocument) -> None:
        document.add_paragraph("MASTER SERVICES AGREEMENT")
        for heading, body in (
            ("1. Definitions", '"Services" means the services in Schedule A.'),
            ("2. Fees", "Customer shall pay the Fees within thirty days of invoice."),
            ("2.1 Late Payment", "Late amounts bear interest at one percent per month."),
            ("3. Termination", "Either party may terminate on sixty days' written notice."),
        ):
            document.add_paragraph().add_run(heading).bold = True
            document.add_paragraph(body)

    found = _sections("msa.docx", _docx(build))
    assert [(n, h) for n, h, _ in found if n] == [("1", "Definitions"), ("2", "Fees"), ("2.1", "Late Payment"), ("3", "Termination")]
    assert next(t for n, _, t in found if n == "3") == "Either party may terminate on sixty days' written notice."


def test_a_docx_with_list_numbering_and_no_heading_style_carries_words_numbers() -> None:
    def build(document: WordDocument) -> None:
        num = _list_of(document, [("decimal", "%1.", 1)])
        document.add_paragraph("SUPPLY AGREEMENT")
        for heading, body in (("Delivery", "Supplier shall deliver within ten days."), ("Payment", "Buyer shall pay on delivery.")):
            paragraph = document.add_paragraph(heading)
            _number_paragraph(paragraph, num, "0")
            document.add_paragraph(body)

    found = _sections("supply.docx", _docx(build))
    assert [(n, h) for n, h, _ in found if n] == [("1", "Delivery"), ("2", "Payment")]


def test_control_a_styled_docx_is_read_by_its_styles_alone() -> None:
    def build(document: WordDocument) -> None:
        document.add_heading("Termination", level=1)
        document.add_paragraph("2. Fees")  # a numbered line in a styled document is body, as before
        document.add_paragraph("Either party may terminate on notice.")

    found = _sections("styled.docx", _docx(build))
    assert [h for _, h, _ in found] == ["Termination"] and "2. Fees" in found[0][2]


def test_control_body_sentences_and_figures_in_an_unstyled_docx_stay_body() -> None:
    def build(document: WordDocument) -> None:
        document.add_paragraph("LICENCE")
        document.add_paragraph("1. Grant")
        document.add_paragraph("2024 Fee Schedule applies to every order placed under this Licence.")
        document.add_paragraph("The Licensee shall not sublicense.")

    found = _sections("licence.docx", _docx(build))
    assert [(n, h) for n, h, _ in found if n] == [("1", "Grant")]
    assert "2024 Fee Schedule applies" in found[-1][2] and "The Licensee shall not sublicense." in found[-1][2]


def test_a_numbered_title_ending_in_a_full_stop_is_a_heading() -> None:
    text = (
        "PURCHASE AGREEMENT\n\nSection 2.6 Closing.\n\nThe Closing shall take place on the Closing Date.\n\n"
        "Section 2.7 Purchase Price Adjustment.\n\nThe Purchase Price shall be reduced dollar for dollar.\n\n"
        "Section 6.1 Conduct of Business Prior to the Closing.\n\nSeller shall operate the Business in the ordinary course.\n"
    )
    found = _sections("apa.txt", text.encode())
    assert [(n, h) for n, h, _ in found if n] == [("2.6", "Closing"), ("2.7", "Purchase Price Adjustment"), ("6.1", "Conduct of Business Prior to the Closing")]


def test_control_a_run_in_title_and_a_sentence_stay_as_before() -> None:
    text = "AGREEMENT\n\n12.3 Governing Law. This Agreement is governed by the laws of New York.\n\n12.4 Notices shall be in writing.\n"
    found = _sections("a.txt", text.encode())
    # A numbered line that is a whole sentence is body, as it was: its words are never a title.
    assert [n for n, _, _ in found if n] == []
    body = " ".join(t for _, _, t in found)
    assert "12.3 Governing Law. This Agreement is governed" in body and "12.4 Notices shall be in writing." in body


def test_an_article_line_and_its_title_are_one_heading() -> None:
    text = "ASSET PURCHASE AGREEMENT\n\nARTICLE II\nPURCHASE AND SALE\n\n2.1 Purchased Assets\n\nSeller shall sell the Purchased Assets.\n"
    found = _sections("apa.txt", text.encode())
    headings = [h for _, h, _ in found]
    assert "ARTICLE II PURCHASE AND SALE" in headings and "ARTICLE II" not in headings


def test_a_long_wrapped_clause_line_is_quotable_body() -> None:
    text = (
        "EQUIPMENT LEASE\n\n4 Rent\n\nLessee shall pay Monthly Rent in advance.\n\n"
        "4.2 Any Monthly Rent not paid within ten (10) days after its due date bears a late charge\n"
        "of five percent of the overdue amount.\n"
    )
    found = _sections("lease.txt", text.encode())
    clause = next(s for s in found if s[0] == "4.2")
    assert "Any Monthly Rent not paid within ten (10) days after its due date bears a late charge" in clause[2]
    assert clause[1] == "Rent"  # labelled by the clause that holds it, never by its own operative words


def test_control_a_long_title_cased_heading_stays_a_title() -> None:
    text = "CREDIT AGREEMENT\n\n5.2 Conditions Precedent to Each Loan and Each Letter of Credit\n\nThe obligation of each Lender is subject to the following.\n"
    found = _sections("credit.txt", text.encode())
    assert ("5.2", "Conditions Precedent to Each Loan and Each Letter of Credit") in [(n, h) for n, h, _ in found]


def test_whether_a_docx_uses_heading_styles_is_read_once_per_style(monkeypatch: pytest.MonkeyPatch) -> None:
    """A bounds fixture of six million empty paragraphs took more than five minutes when every paragraph was asked;
    the styles in use are a handful, and each is asked once."""
    from app.ingest import readers

    calls: list[int] = []
    real = readers._heading_level

    def counted(paragraph: object) -> int:
        calls.append(1)
        return real(paragraph)  # type: ignore[arg-type]

    def build(document: WordDocument) -> None:
        for _ in range(3000):
            document.add_paragraph("")
        document.add_paragraph("A body line.")

    data = _docx(build)
    monkeypatch.setattr(readers, "_heading_level", counted)
    document = DocxDocument(io.BytesIO(data))
    assert readers._uses_heading_styles(document, document.element.body) is False
    assert len(calls) <= 2
