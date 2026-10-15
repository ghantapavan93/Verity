"""The repairs after the legal pilot (2026-10-14): five demonstrated public-release defects, each with its regression
case, controls and, where it applies, a property.

1. Verifier v10. A minus sign or the parentheses of an amount in a currency v9 did not name ("Ksh", "RM", "Rp") were
   punctuation, so "Ksh 5M" verified against "-Ksh 5M" and the evidence pack agreed. From v10 a short label after a
   minus sign or inside an opening parenthesis is part of the number.
2. Reader v14, DOCX. A heading's number was a count of headings, so a list starting at 3 read "1", a restarted list
   "3", a Heading 2 on a list's first level "0.1". From v14 a heading carries the number Word prints, from the list's
   own counters, or none.
3. policy-v9. Guidance in Business Days against a quote in calendar days came back "incomparable" and fell through to
   the model's pass: "30 days, which meets the guidance" of 30 Business Days stood as within guidance, with no reason.
4. The memo. "Issues requiring review (0). None. Every finding was within guidance or a plain answer." headed exactly
   that wrong answer. A flag is now what an automated check raised, with its reason, and every finding says whether a
   person reviewed it.
5. Reader v14, all formats. "AMENDMENT NO. 1 TO EQUIPMENT LEASE AGREEMENT" was cut to the title "AMENDMENT NO", and a
   wrapped clause line ("… signed by both parties. If an") was a title: the model's citation was shown as "found in
   AMENDMENT NO.".
"""

from __future__ import annotations

import hashlib
import io
import json
import types
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fastapi.testclient import TestClient
from hypothesis import given, settings
from hypothesis import strategies as st

from app.application.evidence_pack import render_verify_py
from app.ingest import PARSER_VERSION
from app.ingest.sections import Block, build_sections
from app.memo.service import FLAGS_HEAD, NO_FLAGS, NOT_REVIEWED
from app.runs.status import POLICY_VERSION, decide
from app.verify import tokens
from app.verify.spans import VERIFIER_VERSION, bounded, locate
from tests.support import FUZZ_SCALE, upload_and_ask
from tests.test_docx_numbering import _number_paragraph, _read
from tests.test_memo_restraint import CompliantProvider, memo_html


def _checker(folder: Path) -> types.ModuleType:
    module = types.ModuleType("pack_verify")
    module.__dict__["__file__"] = str(folder / "verify.py")
    exec(compile(render_verify_py(), "verify.py", "exec"), module.__dict__)
    return module


# ------------------------------------------------------------------------------------------------ 1. verifier v10

SIGN_DROPS = [
    pytest.param("Ksh 5M", "The Adjustment Amount is -Ksh 5M at Closing.", id="exact-inside-minus-ksh"),
    pytest.param("Adjustment Amount is Ksh 5M", "The Adjustment Amount is -Ksh 5M at Closing.", id="typed-minus-ksh"),
    pytest.param("RM5M", "The Adjustment Amount is (RM5M) at Closing.", id="exact-inside-parenthesised-rm"),
    pytest.param("Adjustment Amount is RM5M", "The Adjustment Amount is -RM5M at Closing.", id="typed-minus-rm"),
    pytest.param("Rp 5M", "The Adjustment Amount is -Rp 5M at Closing.", id="exact-inside-minus-rp"),
    pytest.param("Adjustment Amount is Rp 5M", "The Adjustment Amount is (Rp 5M) at Closing.", id="typed-parenthesised-rp"),
    pytest.param("Adjustment Amount is -Ksh 5M", "The Adjustment Amount is Ksh 5M at Closing.", id="sign-the-text-lacks"),
    pytest.param("Adjustment Amount is (RM5M)", "The Adjustment Amount is RM5M at Closing.", id="parentheses-the-text-lacks"),
    pytest.param("a credit of Tk 500", "a credit of -Tk 500 applies", id="any-label-plain-amount"),
]


@pytest.mark.parametrize(("quote", "text"), SIGN_DROPS)
def test_a_sign_is_never_dropped_whatever_the_currency_label(quote: str, text: str) -> None:
    assert locate(quote, text) is None


GENUINE = [
    pytest.param("Adjustment Amount is -Ksh 5M", "The Adjustment Amount is -Ksh 5M at Closing.", "exact", id="signed-as-written"),
    pytest.param("(RM5M)", "The Adjustment Amount is (RM5M) at Closing.", "exact", id="parenthesised-as-written"),
    pytest.param("(RM 5M)", "The Adjustment Amount is (RM5M) at Closing.", "typed", id="glued-and-spaced-label-one-value"),
    pytest.param("Ksh 5M", "The amount is Ksh 5M today", "exact", id="unsigned-label"),
    pytest.param("RM5M", "The amount is RM5M today", "exact", id="unsigned-glued-label"),
    pytest.param("Amendment No. 1", "under Amendment No. 1 dated", "exact", id="reference-word-unsigned"),
    pytest.param("Amendment (No. 1)", "under Amendment (No. 1) dated", "exact", id="reference-in-parentheses-as-written"),
    pytest.param("$5M", "a fee of $5M is due", "exact", id="v9-control-dollar"),
    pytest.param("thirty (30) days", "within thirty (30) days after", "exact", id="v8-control-parenthesised-number"),
    pytest.param("- Page 1 of 3", "Agreement - Page 1 of 3", "exact", id="a-dash-and-a-space-is-no-sign"),
]


@pytest.mark.parametrize(("quote", "text", "method"), GENUINE)
def test_controls_genuine_quotes_are_found_and_the_pack_agrees(tmp_path: Path, quote: str, text: str, method: str) -> None:
    found = locate(quote, text)
    assert found is not None and found.method == method, found
    checker = _checker(tmp_path)
    assert checker.edges_hold(text, found.start, found.end, 10)
    assert checker.same(text[found.start : found.end], quote, method, 10)


def test_the_pack_refuses_the_sign_drops_for_a_v10_run_and_reads_a_v9_run_as_v9_did(tmp_path: Path) -> None:
    checker = _checker(tmp_path)
    text = "The Adjustment Amount is -Ksh 5M at Closing."
    span = "Adjustment Amount is -Ksh 5M"
    start = text.index(span)
    assert not checker.same(span, "Adjustment Amount is Ksh 5M", "typed", 10)
    assert checker.same(span, "Adjustment Amount is Ksh 5M", "typed", 9), "a v9 run's pack keeps v9's verdict"
    inner = text.index("Ksh 5M")
    assert not checker.edges_hold(text, inner, inner + len("Ksh 5M"), 10)
    assert checker.edges_hold(text, start, start + len(span), 10)


def test_the_verifier_is_v10_and_v8_and_v9_read_as_they_did() -> None:
    assert VERIFIER_VERSION == "v10" and tokens.VERSION == 10
    assert hashlib.sha256(tokens._WRITTEN_V8.pattern.encode()).hexdigest().startswith("cc4fe06d")
    # v9's pattern exactly as every v9 pack carries it (the legal pilot's packs among them).
    assert hashlib.sha256(tokens._WRITTEN_V9.pattern.encode()).hexdigest() == "204a64ef69f50410410ba5585dfc48fe4194dbb0263a82844125528e52b29a2d"
    assert tokens.written_numbers("is -Ksh 5M now", 9) == [(8, 10)] and tokens.written_numbers("is -Ksh 5M now", 10) == [(3, 10)]


LABELS = ["Ksh ", "RM", "Rp ", "Tk ", "KSh ", "Bs. ", "Sh ", "R ", "RM "]
SIGNS = {"": "", "-": "-", "−": "-", "(": "()"}


@settings(max_examples=400 * FUZZ_SCALE, deadline=None)
@given(
    label=st.sampled_from(LABELS),
    digits=st.sampled_from(["5", "250", "1,500", "2.5", "10"]),
    magnitude=st.sampled_from(["M", "K", "bn", ""]),
    quote_sign=st.sampled_from(sorted(SIGNS)),
    text_sign=st.sampled_from(sorted(SIGNS)),
)
def test_an_amount_with_any_label_is_found_exactly_when_its_sign_agrees(label: str, digits: str, magnitude: str, quote_sign: str, text_sign: str) -> None:
    body = f"{label}{digits}{magnitude}"

    def written(sign: str) -> str:
        return f"({body})" if sign == "(" else sign + body

    quote = f"a credit of {written(quote_sign)}"
    text = f"The parties agree a credit of {written(text_sign)} applies."
    found = locate(quote, text)
    assert (found is not None) == (SIGNS[quote_sign] == SIGNS[text_sign]), (quote, text, found)


@settings(max_examples=300 * FUZZ_SCALE, deadline=None)
@given(text=st.text(alphabet="0123456789,.%$()-−– KshRMpTkNoUSDbn‎", min_size=1, max_size=24), data=st.data())
def test_the_pack_reads_every_edge_as_the_workbench_does_under_v10(tmp_path_factory: pytest.TempPathFactory, text: str, data: st.DataObject) -> None:
    checker = _checker(tmp_path_factory.mktemp("edges-v10"))
    start = data.draw(st.integers(0, len(text)))
    end = data.draw(st.integers(start, len(text)))
    assert checker.edges_hold(text, start, end, 10) == bounded(text, start, end), (text, start, end)


@settings(max_examples=300 * FUZZ_SCALE, deadline=None)
@given(a=st.text(alphabet="0123456789,.$()-− KshRMpTk", min_size=1, max_size=20), b=st.text(alphabet="0123456789,.$()-− KshRMpTk", min_size=1, max_size=20))
def test_the_pack_compares_values_as_the_workbench_does_under_v10(tmp_path_factory: pytest.TempPathFactory, a: str, b: str) -> None:
    checker = _checker(tmp_path_factory.mktemp("typed-v10"))

    def runtime(text: str) -> list[tuple[str, str]]:
        return [(str(t.kind), t.canonical) for t in tokens.significant(tokens.signed(text, tokens.tokenize(text)))]

    assert checker.same(a, b, "typed", 10) == (bool(runtime(a)) and runtime(a) == runtime(b)), (a, b)


# ------------------------------------------------------------------------------------------------ 2. reader v14, DOCX numbers

DECIMAL = [("decimal", "%1.", 1), ("decimal", "%1.%2", 1)]


def _list_of(document: object, levels: list[tuple[str, str, int]], override: str = "", level_extra: str = "") -> str:
    """A list whose levels have their own start values (the shared helper's start at 1), and one instance of it;
    ``level_extra`` is more of each level's definition (a w:lvlRestart)."""
    root = document.part.numbering_part.element  # type: ignore[attr-defined]
    key = str(900 + len(root))
    lvls = "".join(
        f'<w:lvl w:ilvl="{i}"><w:start w:val="{start}"/><w:numFmt w:val="{fmt}"/>{level_extra}<w:lvlText w:val="{text}"/></w:lvl>'
        for i, (fmt, text, start) in enumerate(levels)
    )
    root.insert(0, parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{key}">{lvls}</w:abstractNum>'))
    root.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="{key}"><w:abstractNumId w:val="{key}"/>{override}</w:num>'))
    return key


def _heading(document: object, text: str, level: int, num_id: str | None, ilvl: str = "0") -> None:
    paragraph = document.add_heading(text, level=level)  # type: ignore[attr-defined]
    if num_id is not None:
        _number_paragraph(paragraph, num_id, ilvl)
    document.add_paragraph(f"Body of {text}.")  # type: ignore[attr-defined]


def _numbers(document: object, names: tuple[str, ...]) -> list[str]:
    return [number for number, heading in _read(document) if heading in names]  # type: ignore[arg-type]


@pytest.mark.parametrize(
    ("levels", "override", "expected"),
    [
        pytest.param([("decimal", "%1.", 3)], "", ["3", "4", "5"], id="a-list-starting-at-3"),
        pytest.param(
            [("decimal", "%1.", 1)], '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="10"/></w:lvlOverride>', ["10", "11", "12"], id="instance-override-10"
        ),
        pytest.param([("upperRoman", "%1.", 1)], "", ["I", "II", "III"], id="upper-roman"),
        pytest.param([("lowerLetter", "(%1)", 1)], "", ["(a)", "(b)", "(c)"], id="letters-as-printed"),
        pytest.param([("decimalZero", "%1.", 1)], "", ["01", "02", "03"], id="decimal-zero"),
        pytest.param([("decimal", "Article %1", 1)], "", ["1", "2", "3"], id="article-label-cited-by-its-number"),
        pytest.param([("ordinalText", "%1", 1)], "", ["", "", ""], id="unsupported-format-omitted"),
    ],
)
def test_a_heading_carries_the_number_word_prints_or_none(levels: list[tuple[str, str, int]], override: str, expected: list[str]) -> None:
    document = DocxDocument()
    num = _list_of(document, levels, override)
    for name in ("Alpha", "Beta", "Gamma"):
        _heading(document, name, 1, num)
    assert _numbers(document, ("Alpha", "Beta", "Gamma")) == expected


def test_sub_levels_restart_under_their_parent_and_a_heading_level_is_not_a_list_level() -> None:
    document = DocxDocument()
    num = _list_of(document, DECIMAL)
    _heading(document, "Alpha", 2, num, "0")  # a Heading 2 on the list's first level: Word prints "1.", v13 said "0.1"
    _heading(document, "Beta", 3, num, "1")
    _heading(document, "Gamma", 3, num, "1")
    _heading(document, "Delta", 2, num, "0")
    _heading(document, "Epsilon", 3, num, "1")
    assert _numbers(document, ("Alpha", "Beta", "Gamma", "Delta", "Epsilon")) == ["1", "1.1", "1.2", "2", "2.1"]


def test_written_and_unnumbered_headings_do_not_consume_the_lists_count() -> None:
    document = DocxDocument()
    num = _list_of(document, DECIMAL)
    _heading(document, "Alpha", 1, num)
    _heading(document, "7 Beta", 1, None)
    _heading(document, "Gamma", 1, None)
    _heading(document, "Delta", 1, num)
    assert [n for n, h in _read(document) if h in ("Alpha", "Beta", "Gamma", "Delta")] == ["1", "7", "", "2"]


def test_numbered_paragraphs_in_a_table_count_as_word_counts_them() -> None:
    document = DocxDocument()
    num = _list_of(document, DECIMAL)
    _heading(document, "Alpha", 1, num)
    cell = document.add_table(rows=1, cols=1).rows[0].cells[0]
    _number_paragraph(cell.paragraphs[0], num, "0")
    cell.paragraphs[0].add_run("a numbered row")
    _heading(document, "Beta", 1, num)
    assert _numbers(document, ("Alpha", "Beta")) == ["1", "3"]


def test_where_word_s_rule_is_not_reproduced_the_heading_gets_no_number() -> None:
    document = DocxDocument()
    first = _list_of(document, DECIMAL)
    root = document.part.numbering_part.element
    second = "8888"
    root.append(
        parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{second}"><w:abstractNumId w:val="{first}"/>'
            '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'
        )
    )
    _heading(document, "Alpha", 1, first)
    _heading(document, "Beta", 1, second)  # two instances of one list in use: continue or restart is Word's rule
    restarted = _list_of(document, [("decimal", "%1.", 1)], level_extra='<w:lvlRestart w:val="0"/>')  # a restart rule of its own
    _heading(document, "Gamma", 1, restarted)
    assert _numbers(document, ("Alpha", "Beta", "Gamma")) == ["", "", ""]


def test_word_does_not_count_an_empty_paragraph_holding_a_section_break_nor_a_deleted_one() -> None:
    """As Word 16 itself reports, on built cases and on Common Paper's agreement, whose cover page ends in an empty
    numbered Heading 1 holding the section break: Word prints "1. Service" after it, and v14 as first written, counting
    every numbered paragraph, said "2" (the corpus check against Word, 2026-10-14). An empty paragraph without a
    section break does count, in Word and here."""
    document = DocxDocument()
    num = _list_of(document, DECIMAL)
    _heading(document, "Alpha", 1, num)
    holder = document.add_heading("", level=1)
    _number_paragraph(holder, num, "0")
    holder._p.get_or_add_pPr().append(parse_xml(f'<w:sectPr {nsdecls("w")}><w:pgSz w:w="12240" w:h="15840"/></w:sectPr>'))
    _heading(document, "Beta", 1, num)
    deleted = document.add_heading("", level=1)
    _number_paragraph(deleted, num, "0")
    deleted._p.append(parse_xml(f'<w:del {nsdecls("w")} w:id="1" w:author="a"><w:r><w:delText>Removed clause</w:delText></w:r></w:del>'))
    deleted._p.get_or_add_pPr().append(parse_xml(f'<w:rPr {nsdecls("w")}><w:del w:id="2" w:author="a"/></w:rPr>'))
    _heading(document, "Gamma", 1, num)
    empty = document.add_heading("", level=1)
    _number_paragraph(empty, num, "0")
    _heading(document, "Delta", 1, num)
    assert _numbers(document, ("Alpha", "Beta", "Gamma", "Delta")) == ["1", "2", "3", "5"]


def test_a_lettered_sub_item_is_cited_with_the_clause_it_sits_in() -> None:
    """Word prints "(a)" for a sub-item and "8.1" for its clause; the document cites it "Section 8.1(a)", and so does
    the reading (the mission, 2026-10-14: the answer card showed "§(a) · Termination" beside "§5.3"). Only what Word
    prints at each level is used: a level that prints its parents itself, or one under an unnumbered level, is as printed."""
    document = DocxDocument()
    num = _list_of(document, [("decimal", "%1.", 1), ("decimal", "%1.%2", 1), ("lowerLetter", "(%3)", 1), ("lowerRoman", "(%4)", 1)])
    for name, ilvl in (("Alpha", "0"), ("Beta", "1"), ("Gamma", "2"), ("Delta", "3"), ("Epsilon", "2"), ("Zeta", "1"), ("Eta", "2")):
        _heading(document, name, 1, num, ilvl)
    assert _numbers(document, ("Alpha", "Beta", "Gamma", "Delta", "Epsilon", "Zeta", "Eta")) == ["1", "1.1", "1.1(a)", "1.1(a)(i)", "1.1(b)", "1.2", "1.2(a)"]
    full = DocxDocument()
    nested = _list_of(full, [("decimal", "%1.", 1), ("decimal", "%1.%2", 1), ("decimal", "%1.%2.%3", 1)])
    for name, ilvl in (("Alpha", "0"), ("Beta", "1"), ("Gamma", "2")):
        _heading(full, name, 1, nested, ilvl)
    assert _numbers(full, ("Alpha", "Beta", "Gamma")) == ["1", "1.1", "1.1.1"]  # control: Word prints the parents itself
    top = DocxDocument()
    letters = _list_of(top, [("lowerLetter", "(%1)", 1)])
    _heading(top, "Alpha", 1, letters)
    assert _numbers(top, ("Alpha",)) == ["(a)"]  # control: nothing above it to cite


def test_the_reader_is_v14() -> None:
    assert PARSER_VERSION == "v14"


# ------------------------------------------------------------------------------------------------ 3. policy-v9

H4_Q5_GUIDANCE = "Seller must be given at least thirty (30) Business Days to dispute the Closing Statement."
H4_Q5_QUOTE = ["Seller may dispute the Closing Statement by written notice to Buyer within thirty (30) days after receiving it."]


def test_calendar_days_against_business_days_is_no_pass() -> None:
    assert POLICY_VERSION == "policy-v11"  # v11 keeps this v9/v10 decision
    decision = decide("pass", "thirty (30) days", "at least thirty (30) Business Days", True, True, H4_Q5_QUOTE, H4_Q5_GUIDANCE, "Dispute Period")
    assert (decision.status, decision.source) == ("needs_review", "position_check")
    assert "calendar days cannot be compared with business days" in decision.reason and "not established" in decision.reason


def test_months_against_days_is_no_pass() -> None:
    decision = decide(
        "pass",
        "two (2) months",
        "at least 60 days",
        True,
        True,
        ["Either party may terminate on two (2) months' written notice."],
        "Termination requires at least 60 days' written notice.",
        "Termination notice",
    )
    assert decision.status == "needs_review" and decision.source == "position_check"


@pytest.mark.parametrize(
    ("hint", "observed", "required", "quote", "guidance", "expected"),
    [
        pytest.param(
            "pass",
            "thirty (30) Business Days",
            "at least thirty (30) Business Days",
            ["Seller may dispute within thirty (30) Business Days after receiving it."],
            H4_Q5_GUIDANCE,
            ("pass", "model_hint"),
            id="same-unit-pass-stays-the-models",
        ),
        pytest.param(
            "pass",
            "thirty (30) days",
            "at least 45 days",
            H4_Q5_QUOTE,
            "Seller must be given at least 45 days to dispute the Closing Statement.",
            ("needs_review", "computed_days"),
            id="same-unit-shortfall-is-codes",
        ),
        pytest.param(
            "needs_review",
            "thirty (30) days",
            "at least thirty (30) Business Days",
            H4_Q5_QUOTE,
            H4_Q5_GUIDANCE,
            ("needs_review", "model_hint"),
            id="the-models-own-review-stays",
        ),
    ],
)
def test_controls_comparable_periods_decide_as_before(
    hint: str, observed: str, required: str, quote: list[str], guidance: str, expected: tuple[str, str]
) -> None:
    decision = decide(hint, observed, required, True, True, quote, guidance, "Dispute Period")  # type: ignore[arg-type]
    assert (decision.status, decision.source) == expected


# ------------------------------------------------------------------------------------------------ 4. the memo


def test_a_memo_with_no_flag_says_whose_answers_they_are(client: TestClient) -> None:
    client.app.state.provider = CompliantProvider()  # type: ignore[attr-defined]
    run_id = upload_and_ask(client, "What notice period applies to termination for convenience?", with_guidance=False)["run_id"]
    html = memo_html(client, run_id)
    assert f"{FLAGS_HEAD} (0)" in html
    assert "Issues requiring review" not in html and "Every finding was within guidance" not in html
    assert "No automated check flagged a finding" in html and "not that the answers are right" in html
    assert NOT_REVIEWED in html
    memo = client.post("/api/memos", json={"runId": run_id}).json()
    docx = DocxDocument(io.BytesIO(client.get(f"/api/memos/{memo['id']}/docx").content))
    words = "\n".join(p.text for p in docx.paragraphs)
    assert f"{FLAGS_HEAD} (0)" in words and NO_FLAGS in words and NOT_REVIEWED in words and "Issues requiring review" not in words


def test_a_flag_carries_the_automated_checks_own_reason(client: TestClient) -> None:
    run_id = upload_and_ask(client, "What notice period applies to termination for convenience?")["run_id"]
    finding = client.get(f"/api/runs/{run_id}/detail").json()["findings"][0]
    assert finding["status"] == "needs_review" and finding["statusReason"]
    html = memo_html(client, run_id)
    assert f"{FLAGS_HEAD} (1)" in html and "Why flagged." in html
    assert "code: " in html and finding["statusReason"].split(";")[0][:40].replace("'", "&#x27;") in html


def test_a_reviewed_finding_says_who_reviewed_it_and_no_other_says_so(client: TestClient) -> None:
    run_id = upload_and_ask(client, "What notice period applies to termination for convenience?")["run_id"]
    detail = client.get(f"/api/runs/{run_id}/detail").json()
    finding = detail["findings"][0]
    reviewed = client.post(
        f"/api/findings/{finding['id']}/review", json={"verdict": "confirmed", "reviewer": "A. Reviewer", "basedOn": finding["latestReviewId"]}
    )
    assert reviewed.status_code in (200, 201), reviewed.text
    html = memo_html(client, run_id)
    assert "Reviewed." in html and "A. Reviewer" in html and NOT_REVIEWED not in html


# ------------------------------------------------------------------------------------------------ 5. heading labels


def _sections(*blocks: Block) -> list[tuple[str, str, str]]:
    return [(s.number, s.heading, s.text) for s in build_sections(list(blocks))]


def test_an_amendment_heading_is_not_cut_at_its_number() -> None:
    found = _sections(
        Block("heading", "AMENDMENT NO. 1 TO EQUIPMENT LEASE AGREEMENT", 1, implicit_number=False),
        Block("text", "This Amendment No. 1 is made on June 19, 2026."),
    )
    assert found[0][1] == "AMENDMENT NO. 1 TO EQUIPMENT LEASE AGREEMENT"
    assert found[0][2].startswith("This Amendment No. 1")


def test_a_wrapped_clause_line_is_body_not_a_title() -> None:
    found = _sections(
        Block("heading", "11 General", 1, implicit_number=False),
        Block("heading", "11.2 This Lease may be amended only by a written amendment signed by both parties. If an", 2, implicit_number=False),
        Block("text", "amendment conflicts with this Lease, the amendment controls."),
        Block("heading", "8.1 Seller shall indemnify Buyer against Losses arising from (a) any breach of Seller's", 2, implicit_number=False),
        Block("text", "representations and warranties and (b) any Excluded Liability."),
    )
    clause = next(s for s in found if s[0] == "11.2")
    assert clause[1] == "General" and "both parties. If an" in clause[2] and "the amendment controls" in clause[2]
    indemnity = next(s for s in found if s[0] == "8.1")
    assert "Seller shall indemnify" in indemnity[2] and not indemnity[1].endswith("Seller's")


@pytest.mark.parametrize(
    "title",
    [
        "Acme Corp. Master Services Agreement",
        "U.S. Export Controls",
        "May 2026 Amendment",
        "Payment Terms; Late Fees",
        "Representations and Warranties of the Parties",
    ],
)
def test_controls_real_titles_stay_titles(title: str) -> None:
    found = _sections(Block("heading", f"4 {title}", 1, implicit_number=False), Block("text", "The parties agree as follows."))
    assert found[0][:2] == ("4", title)


@pytest.mark.parametrize(
    ("line", "title", "body"),
    [
        pytest.param(
            "13.2 Governing Law. This Agreement is governed by the laws of New York.", "Governing Law", "This Agreement is governed", id="lead-in-title"
        ),
        pytest.param("5. Fees: Customer shall pay the fees in Schedule 2 within thirty days.", "Fees", "Customer shall pay", id="colon-lead-in"),
    ],
)
def test_controls_a_lead_in_title_still_splits_from_its_body(line: str, title: str, body: str) -> None:
    found = _sections(Block("heading", line, 1, implicit_number=False))
    assert found[0][1] == title and found[0][2].startswith(body)


def test_a_wrapped_line_keeps_its_own_title_and_a_parent_labels_only_a_clause_it_holds() -> None:
    """From the corpus check (reader v13 against v14 over 6,311 files): a clause line took the title of whatever
    heading stood a level up ("<PAGE>", the agreement's name) for its label, and "Designation of Subsidiaries. The
    Parent Borrower" lost its own title to its article's."""
    found = _sections(
        Block("heading", "SITE DEVELOPMENT AND HOSTING AGREEMENT", 1, implicit_number=False),
        Block("heading", "1.4 HDI shall provide the Company with access to, and", 2, implicit_number=False),
        Block("text", "use of, the Site."),
        Block("heading", "5 AFFIRMATIVE COVENANTS", 1, implicit_number=False),
        Block("heading", "5.10 Designation of Subsidiaries. The Parent Borrower", 2, implicit_number=False),
        Block("text", "may designate any Subsidiary as an Unrestricted Subsidiary."),
        Block("heading", "5.11 The Seller shall deliver the Goods. The Buyer shall", 2, implicit_number=False),
        Block("text", "pay on delivery."),
        Block("heading", "E.PIPHANY, INC.          HIGH SPEED NET SOLUTIONS, INC.", 2, implicit_number=False),
        Block("text", "SIGNATURE: /s/ Karen Richardson"),
    )
    signatures = found[-1]  # unnumbered: its own words, never the title of the clause above it
    assert signatures[0] == "" and signatures[1].startswith("E.PIPHANY, INC.") and signatures[2].startswith("E.PIPHANY")
    access = next(s for s in found if s[0] == "1.4")
    assert access[1].startswith("HDI shall provide") and access[1].endswith("…") and "HDI shall provide" in access[2]
    designation = next(s for s in found if s[0] == "5.10")
    assert designation[1] == "Designation of Subsidiaries" and designation[2].startswith("The Parent Borrower")
    delivery = next(s for s in found if s[0] == "5.11")  # a sentence is no title, however short: its parent's label
    assert delivery[1] == "AFFIRMATIVE COVENANTS" and delivery[2].startswith("The Seller shall deliver the Goods.")


def test_the_stage_a_answers_are_untouched_by_this_repair() -> None:
    """Stage A's records are historical: their runs keep the versions they were made under (checked by the run
    identity), and nothing here re-decides them. The versions a new run carries are the repaired ones."""
    from app.runs.versions import SEMANTIC_VERSIONS

    # policy-v11 (2026-10-14) keeps every v10 decision and adds the contract's own replacement statements (runs.relations).
    assert SEMANTIC_VERSIONS["verifier_version"] == "v10" and SEMANTIC_VERSIONS["policy_version"] == "policy-v11"
    assert json.dumps(SEMANTIC_VERSIONS, sort_keys=True)
