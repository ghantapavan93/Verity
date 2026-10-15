"""A clause heading the author typed is a heading in a document that styles its others (reader v14, first-visitor
acceptance test, 2026-10-14).

A fresh services agreement styled most clause headings with Heading 1 but typed "2. Term", "7. Forecasts and Purchase
Orders", "11. Termination" and "14. General" as bold Normal paragraphs. Reader v14 read every unstyled paragraph of a
styled document as body, so clause 11 sat inside §10 and a quote from 11.2(b)(ii) was "found at §10 · Liability" on the
card, under code's badge. The invariant: a passage is cited by the clause the document prints it under, whether the
author styled that clause's heading or typed it. Controls: list items, numbers that go back, sentences and sub-numbers
stay body; a document with no heading styles reads as before.
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from hypothesis import given, settings
from hypothesis import strategies as st

from app.ingest.readers import read
from app.ingest.sections import build_sections
from tests.support import FUZZ_SCALE

BODY = "The parties agree that this clause applies throughout the term of this Agreement and binds their successors."


def _docx(paragraphs: list[tuple[str, str, bool]]) -> bytes:
    """(style, text, bold) paragraphs as a .docx."""
    document = Document()
    for style, text, bold in paragraphs:
        paragraph = document.add_paragraph(style=style)
        paragraph.add_run(text).bold = bold or None
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _numbers(paragraphs: list[tuple[str, str, bool]]) -> list[str]:
    result = read("agreement.docx", _docx(paragraphs))
    return [s.number for s in build_sections(result.blocks, result.title) if s.number]


def test_the_acceptance_case_typed_bold_headings_between_styled_ones_are_clauses() -> None:
    paragraphs = [
        ("Heading 1", "10. Liability", False),
        ("Normal", "10.1 Subject to clause 10.2, Supplier's total aggregate liability is limited to the Fees paid.", False),
        ("Normal", "11. Termination", True),
        ("Normal", "11.2 Customer may terminate this Agreement with immediate effect by written notice to Supplier if:", False),
        ("Normal", "(ii) where the breach consists of a Critical Defect, Supplier fails to remedy it within six (6) Business Days;", False),
        ("Heading 1", "13. Notices", False),
        ("Normal", BODY, False),
    ]
    result = read("agreement.docx", _docx(paragraphs))
    sections = {s.number: s for s in build_sections(result.blocks, result.title)}
    assert list(sections) == ["10", "11", "13"]
    assert sections["11"].heading == "Termination" and "six (6) Business Days" in sections["11"].text
    assert "six (6) Business Days" not in sections["10"].text


@pytest.mark.parametrize(
    ("typed", "expected"),
    [
        (("Normal", "2. TERM", False), ["1", "2", "3"]),  # in capitals, not bold
        (("Normal", "2 Term", True), ["1", "2", "3"]),  # no full stop after the number
        (("Normal", "Section 2 Term", True), ["1", "2", "3"]),  # a clause word before the number
        (("Body Text", "2. Term", True), ["1", "2", "3"]),  # another body style
    ],
)
def test_a_typed_heading_is_read_however_it_is_set_apart(typed: tuple[str, str, bool], expected: list[str]) -> None:
    assert (
        _numbers(
            [
                ("Heading 1", "1. Definitions", False),
                ("Normal", BODY, False),
                typed,
                ("Normal", BODY, False),
                ("Heading 1", "3. Fees", False),
                ("Normal", BODY, False),
            ]
        )
        == expected
    )


def test_boundary_one_clause_written_inline_between_is_allowed_two_are_not() -> None:
    inline = ("Normal", "11. Confidentiality. Each party shall keep the other's information confidential.", False)
    one = [("Heading 1", "10. Liability", False), ("Normal", BODY, False), inline, ("Normal", "12. Notices", True), ("Normal", BODY, False)]
    assert _numbers(one) == ["10", "12"]
    two = [("Heading 1", "10. Liability", False), ("Normal", BODY, False), ("Normal", "13. Notices", True), ("Normal", BODY, False)]
    assert _numbers(two) == ["10"]


@pytest.mark.parametrize(
    "paragraph",
    [
        ("Normal", "4. Training Services", False),  # a list item: the next number, but neither bold nor in capitals
        ("Normal", "2. Payment", True),  # bold, but the number goes back
        ("Normal", "4. The Supplier shall deliver the Services in accordance with the Specification at all times.", True),  # a sentence
        ("Normal", "4.1 Delivery", True),  # a sub-number, not a clause heading
        ("Normal", "Background", True),  # no number
    ],
)
def test_control_what_is_not_a_typed_clause_heading_stays_body(paragraph: tuple[str, str, bool]) -> None:
    numbers = _numbers([("Heading 1", "3. Services", False), ("Normal", BODY, False), paragraph, ("Normal", BODY, False)])
    assert numbers == ["3"]


def test_control_a_document_without_heading_styles_reads_as_before() -> None:
    unstyled = [("Normal", "1. Definitions", True), ("Normal", BODY, False), ("Normal", "2. Term", False), ("Normal", BODY, False)]
    assert _numbers(unstyled) == ["1", "2"]  # the unstyled reading (v14) takes both, bold or not


@settings(max_examples=40 * FUZZ_SCALE, deadline=None)
@given(typed=st.lists(st.booleans(), min_size=3, max_size=12), capitals=st.booleans())
def test_every_clause_heading_is_read_whatever_mix_of_styled_and_typed(typed: list[bool], capitals: bool) -> None:
    if all(typed):
        typed[0] = False  # a document with no styled heading is read by the unstyled path, tested above
    paragraphs: list[tuple[str, str, bool]] = []
    for n, is_typed in enumerate(typed, start=1):
        title = f"Clause Title {n}".upper() if capitals and is_typed else f"Clause Title {n}"
        paragraphs.append(("Normal", f"{n}. {title}", not capitals) if is_typed else ("Heading 1", f"{n}. {title}", False))
        paragraphs.append(("Normal", f"{n}.1 {BODY}", False))
    assert _numbers(paragraphs) == [str(n) for n in range(1, len(typed) + 1)]
