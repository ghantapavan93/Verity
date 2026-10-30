"""Reader v16, DOCX: the reading is the text Word displays, in each cell's column, and what it does not read is counted.

Found by the adversarial stress test (2026-10-30), each with Word 16 as the oracle:

(1) The walk read text Word does not display: a nested field's result written inside the outer field's instruction
("thirty (30) days ninety (90) days" where Word shows "ninety (90) days"), both branches of an mc:AlternateContent, an
element of an ignorable namespace Word drops, and a data-bound control's stale cache; and it stored a w:bdo override in
stored order without the direction controls Word holds for it, so nothing disclosed it. The invariant: a DOCX reading
is the text Word displays, with any reordering disclosed (word_display.py).

(2) A legacy check box's state, a drop-down's choice and every w:sym were dropped: a completed ICO IDTA with Scotland
ticked read exactly as the blank form, and Symbol's "±" was lost from "±0.25". The invariant: a choice the document
records is in the reading, and a glyph Word draws is read as that glyph (or as U+FFFD, counted, when unknown).

(3) A table row kept only cells whose text was new in it, so every value after an empty or a repeated cell moved into
the column before ("Licence fee | $100,000 | $150,000" under Year 1 | Year 2 | Year 3). The invariant: a row's text
keeps every cell in its column: one slot per grid column.

(4) Displayed content was dropped while coverage said "read": text boxes, block w:customXml, a control or a w:customXml
around a row's cells, pictures. The invariant: what the reading skipped is reported with a count, and a text box is
read once, as its own text.

(5) Text Word draws so it cannot be seen (white with nothing shaded beneath, half a point) was read, quoted and
verified (DECISIONS 2026-10-30, open). The invariant: it is left out and counted, as hidden runs are.
"""

from __future__ import annotations

import io
import zipfile
from xml.sax.saxutils import escape

import pytest
from docx import Document as DocxDocument
from hypothesis import given, settings
from hypothesis import strategies as st

from app.ingest import ingest
from app.ingest.coverage import Status
from app.verify.spans import locate
from tests.support import FUZZ_SCALE

NAMESPACES = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
    "v": "urn:schemas-microsoft-com:vml",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "x": "urn:example:a-namespace-word-does-not-know",
}
DECLARED = " ".join(f'xmlns:{prefix}="{uri}"' for prefix, uri in NAMESPACES.items())


def docx(body: str, *, ignorable: str = "w14", root: str = "", parts: dict[str, str] | None = None, styles: str = "", head: str = "") -> bytes:
    """A package whose body is ``body`` (``head`` before it, as a w:background), built on python-docx's template;
    ``parts`` adds or replaces package members."""
    buffer = io.BytesIO()
    DocxDocument().save(buffer)
    document = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {DECLARED} mc:Ignorable="{ignorable}" {root}>'
        f"{head}<w:body>{body}<w:sectPr/></w:body></w:document>"
    )
    replace = {"word/document.xml": document, **(parts or {})}
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as original, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as package:
        for name in original.namelist():
            data = original.read(name)
            if name == "word/styles.xml" and styles:
                data = data.replace(b"</w:styles>", styles.encode() + b"</w:styles>")
            if name not in replace:
                package.writestr(name, data)
        for name, text in replace.items():
            package.writestr(name, text)
    return out.getvalue()


def t(text: str, properties: str = "") -> str:
    return f'<w:r>{f"<w:rPr>{properties}</w:rPr>" if properties else ""}<w:t xml:space="preserve">{escape(text)}</w:t></w:r>'


def p(*runs: str, properties: str = "") -> str:
    return f"<w:p>{f'<w:pPr>{properties}</w:pPr>' if properties else ''}{''.join(runs)}</w:p>"


def mark(kind: str, inner: str = "", properties: str = "") -> str:
    return f'<w:r>{f"<w:rPr>{properties}</w:rPr>" if properties else ""}<w:fldChar w:fldCharType="{kind}">{inner}</w:fldChar></w:r>'


def instruction(text: str) -> str:
    return f'<w:r><w:instrText xml:space="preserve">{escape(text)}</w:instrText></w:r>'


def field(code: str, *result: str) -> str:
    return mark("begin") + instruction(code) + mark("separate") + "".join(result) + mark("end")


def body_text(data: bytes) -> str:
    return "\n".join(s.text for s in ingest("contract.docx", data).sections)


def coverage(data: bytes) -> dict[str, tuple[Status, int, str]]:
    return {c.part: (c.status, c.count, c.note) for c in ingest("contract.docx", data).coverage}


# ------------------------------------------------------------------------------------- (1) the text Word displays


def test_a_nested_fields_result_inside_an_instruction_is_not_read() -> None:
    nested = mark("begin") + instruction(" IF ") + field(" DOCPROPERTY NoticeDays ", t("thirty (30) days ")) + instruction(' = "x" "ninety (90) days" ')
    data = docx(p(t("Notice of termination: "), nested, mark("separate"), t("ninety (90) days"), mark("end"), t(".")))
    assert body_text(data) == "Notice of termination: ninety (90) days."
    assert locate("thirty (30)", body_text(data)) is None


def test_a_field_nested_in_a_result_displays_its_own_result() -> None:
    """A table of contents: the TOC's result holds each entry's PAGEREF, whose result Word displays."""
    entry = f'<w:hyperlink w:anchor="_Toc1">{t("Term and termination ")}{field(" PAGEREF _Toc1 ", t("3"))}</w:hyperlink>'
    data = docx(p(t("Contents: ")) + p(mark("begin"), instruction(" TOC \\o "), mark("separate"), entry) + p(mark("end"), t("The term is one year.")))
    text = body_text(data)
    assert "Term and termination 3" in text and "TOC" not in text and "PAGEREF" not in text and "The term is one year." in text


def test_field_marks_in_a_hidden_run_still_move_the_field() -> None:
    hidden = "<w:vanish/>"
    data = docx(
        p(t("Fee: "), mark("begin", properties=hidden), t("thirty (30)"), mark("separate", properties=hidden), t("ninety (90)"), mark("end", properties=hidden))
    )
    assert body_text(data) == "Fee: ninety (90)"


def test_a_deleted_field_is_gone_and_moves_nothing() -> None:
    deleted = f'<w:del w:id="1" w:author="A">{mark("begin")}<w:r><w:delInstrText> REF X </w:delInstrText></w:r>{mark("separate")}</w:del>'
    data = docx(p(t("Notice: "), deleted, t("ninety (90) days")))
    assert body_text(data) == "Notice: ninety (90) days"


def test_a_field_that_never_closes_swallows_nothing_and_is_reported() -> None:
    data = docx(p(t("1. Term "), mark("begin"), instruction(" REF A "), t("of one year.")) + p(t("2. Fees are due monthly.")))
    text = body_text(data)
    assert "of one year." in text and "Fees are due monthly." in text
    status, _, note = coverage(data)["fields"]
    assert status is Status.UNKNOWN and "do not balance" in note


def alternate(choice: str, fallback: str, requires: str = "w14") -> str:
    return f'<mc:AlternateContent><mc:Choice Requires="{requires}">{choice}</mc:Choice><mc:Fallback>{fallback}</mc:Fallback></mc:AlternateContent>'


@pytest.mark.parametrize("level", ["run", "paragraph", "body"])
def test_one_branch_of_alternate_content_is_read_the_one_word_displays(level: str) -> None:
    def build(requires: str) -> bytes:
        if level == "body":
            return docx(p(t("Supplier may terminate.")) + alternate(p(t("on ninety (90) days notice")), p(t("on thirty (30) days notice")), requires))
        if level == "paragraph":
            return docx(p(t("terminate "), alternate(t("on ninety (90)"), t("on thirty (30)"), requires), t(" days notice")))
        inner = alternate('<w:t xml:space="preserve">on ninety (90)</w:t>', '<w:t xml:space="preserve">on thirty (30)</w:t>', requires)
        return docx(p(t("terminate "), f"<w:r>{inner}</w:r>", t(" days notice")))

    shown = body_text(build("w14"))
    assert "ninety (90)" in shown and "thirty (30)" not in shown, "Word understands w14: the Choice"
    fallback = body_text(build("x"))
    assert "thirty (30)" in fallback and "ninety (90)" not in fallback, "a namespace Word does not know: the Fallback"


def test_an_element_of_an_ignorable_namespace_word_does_not_know_is_dropped() -> None:
    wrapped = f"<x:note>{t(' ninety (90) days')}</x:note>"
    assert body_text(docx(p(t("terminate on thirty (30) days"), wrapped), ignorable="w14 x")) == "terminate on thirty (30) days"
    unwrapped = docx(p(t("terminate on"), wrapped), ignorable="w14 x", root='mc:ProcessContent="x:note"')
    assert body_text(unwrapped) == "terminate on ninety (90) days", "mc:ProcessContent: its content is read"


CORE = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><cp:coreProperties '
    'xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/">'
    "<dc:title>ninety (90) days</dc:title></cp:coreProperties>"
)
CORE_BINDING = (
    "<w:dataBinding w:prefixMappings=\"xmlns:ns0='http://purl.org/dc/elements/1.1/' "
    "xmlns:ns1='http://schemas.openxmlformats.org/package/2006/metadata/core-properties'\" "
    'w:xpath="/ns1:coreProperties[1]/ns0:title[1]" w:storeItemID="{6C3C8BC8-F283-45AE-878A-BAB7291924A1}"/>'
)
ITEM_ID = "{0E1C3D52-6A9B-4F1E-8C1D-1A2B3C4D5E6F}"
CUSTOM = {
    "customXml/item9.xml": '<?xml version="1.0"?><terms xmlns="urn:example:terms"><notice>ninety (90) days</notice></terms>',
    "customXml/itemProps9.xml": f'<?xml version="1.0"?><ds:datastoreItem ds:itemID="{ITEM_ID}" xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml"/>',
}
CUSTOM_BINDING = f'<w:dataBinding w:prefixMappings="xmlns:t=\'urn:example:terms\'" w:xpath="/t:terms[1]/t:notice[1]" w:storeItemID="{ITEM_ID}"/>'


def bound(binding: str, cached: str, kind: str = "<w:text/>", block: bool = False) -> str:
    content = p(t(cached)) if block else t(cached)
    return f"<w:sdt><w:sdtPr>{binding}{kind}</w:sdtPr><w:sdtContent>{content}</w:sdtContent></w:sdt>"


@pytest.mark.parametrize(("binding", "parts"), [(CORE_BINDING, {"docProps/core.xml": CORE}), (CUSTOM_BINDING, CUSTOM)])
def test_a_data_bound_control_is_read_as_its_bound_value_not_a_stale_cache(binding: str, parts: dict[str, str]) -> None:
    inline = docx(p(t("terminate on "), bound(binding, "thirty (30) days"), t(" notice")), parts=parts)
    assert body_text(inline) == "terminate on ninety (90) days notice"
    block = docx(p(t("1. Notice")) + bound(binding, "thirty (30) days", block=True), parts=parts)
    assert "ninety (90) days" in body_text(block) and "thirty" not in body_text(block)
    assert coverage(inline)["bound_controls"][:2] == (Status.READ, 1)


def test_a_bound_control_whose_store_is_missing_reads_its_cache_as_word_does() -> None:
    assert body_text(docx(p(t("terminate on "), bound(CUSTOM_BINDING, "thirty (30) days"), t(" notice")))) == "terminate on thirty (30) days notice"


def test_a_bound_value_shown_through_a_format_this_reader_does_not_apply_is_left_out_and_counted() -> None:
    data = docx(p(t("Effective "), bound(CUSTOM_BINDING, "1 March 2026", kind='<w:date w:fullDate="2026-03-01T00:00:00Z"/>')), parts=CUSTOM)
    assert body_text(data) == "Effective \ufffd"
    assert coverage(data)["bound_controls"][:2] == (Status.PARTIAL, 1)


def test_a_bound_control_that_agrees_with_its_store_is_read_as_before() -> None:
    data = docx(p(t("terminate on "), bound(CUSTOM_BINDING, "ninety (90) days"), t(" notice")), parts=CUSTOM)
    assert body_text(data) == "terminate on ninety (90) days notice"


@pytest.mark.parametrize(("element", "direction"), [("bdo", "rtl"), ("bdo", "ltr"), ("dir", "rtl"), ("dir", "ltr")])
def test_a_direction_override_is_disclosed_like_the_same_text_in_a_text_file(element: str, direction: str) -> None:
    data = docx(p(t("terminate on "), f'<w:{element} w:val="{direction}">{t("09")}</w:{element}>', t(" days notice")))
    assert body_text(data) == "terminate on 09 days notice"
    status, count, _ = coverage(data)["format_characters"]
    assert status is Status.EXCLUDED and count == 2, "the control Word holds and its closing pop, removed and counted"


def test_a_plain_paragraph_has_no_format_characters() -> None:
    assert coverage(docx(p(t("terminate on 90 days notice"))))["format_characters"][0] is Status.ABSENT


# Each construct: its markup, and the text Word displays for it (Word 16 COM, stress test 2026-10-30).
CONSTRUCTS: list[tuple[str, str]] = [
    (t("Plain words "), "Plain words "),
    (mark("begin") + instruction(" IF ") + field(" REF A ", t("hidden result ")) + mark("separate") + t("shown result ") + mark("end"), "shown result "),
    (field(" MERGEFIELD Name ", t("Acme Ltd ")), "Acme Ltd "),
    (alternate(t("the choice "), t("the fallback ")), "the choice "),
    (alternate(t("the choice "), t("the fallback "), requires="x"), "the fallback "),
    (f"<x:note>{t('ignored words ')}</x:note>", ""),
    (f'<w:bdo w:val="rtl">{t("09")}</w:bdo>{t(" ")}', "09 "),
    (f'<w:hyperlink w:anchor="a">{t("linked words ")}</w:hyperlink>', "linked words "),
    (f'<w:ins w:id="7" w:author="A">{t("inserted words ")}</w:ins>', "inserted words "),
    ('<w:del w:id="8" w:author="A"><w:r><w:delText>deleted words </w:delText></w:r></w:del>', ""),
]


@settings(max_examples=120 * FUZZ_SCALE, deadline=None)
@given(st.lists(st.sampled_from(range(len(CONSTRUCTS))), min_size=1, max_size=8))
def test_a_paragraph_of_any_mix_of_constructs_reads_as_word_displays_it(picked: list[int]) -> None:
    markup = "".join(CONSTRUCTS[i][0] for i in picked)
    shown = "".join(CONSTRUCTS[i][1] for i in picked).strip()
    data = docx(p(markup) if shown else p(markup) + p(t("Control.")), ignorable="w14 x")
    expected = shown or "Control."
    assert " ".join(body_text(data).split()) == " ".join(expected.split())


# ------------------------------------------------------------------------------------- (2) choices and glyphs


def checkbox(state: str) -> str:
    return (
        mark("begin", f'<w:ffData><w:name w:val="Check1"/><w:checkBox><w:sizeAuto/>{state}</w:checkBox></w:ffData>')
        + instruction(" FORMCHECKBOX ")
        + mark("end")
    )


@pytest.mark.parametrize(
    ("state", "box"),
    [
        ("<w:checked/>", "☒"),
        ('<w:checked w:val="1"/>', "☒"),
        ('<w:checked w:val="0"/>', "☐"),
        ('<w:default w:val="1"/>', "☒"),
        ('<w:default w:val="0"/>', "☐"),
        ('<w:default w:val="1"/><w:checked w:val="0"/>', "☐"),
        ("", "☐"),
    ],
)
def test_a_legacy_check_box_is_read_as_the_box_word_draws(state: str, box: str) -> None:
    assert body_text(docx(p(checkbox(state), t(" Scotland")))) == f"{box} Scotland"


@settings(max_examples=40 * FUZZ_SCALE, deadline=None)
@given(st.integers(min_value=2, max_value=6), st.data())
def test_ticking_any_one_box_changes_the_reading_at_exactly_that_option(boxes: int, data: st.DataObject) -> None:
    """The IDTA's shape: a governing-law row of options, each after its box; ticking box k changes only option k."""
    options = ["England and Wales", "Northern Ireland", "Scotland", "Wales", "Jersey", "Guernsey"][:boxes]
    ticked = data.draw(st.integers(min_value=0, max_value=boxes - 1))

    def row(tick: int | None) -> str:
        cells = "".join(f"<w:tc>{p(checkbox('<w:checked/>' if i == tick else ''), t(' ' + name))}</w:tc>" for i, name in enumerate(options))
        return f"<w:tbl><w:tr><w:tc>{p(t('Governing law:'))}</w:tc>{cells}</w:tr></w:tbl>"

    blank, chosen = body_text(docx(row(None))), body_text(docx(row(ticked)))
    assert blank.count("☐") == boxes and "☒" not in blank
    assert chosen.count("☒") == 1 and f"☒ {options[ticked]}" in chosen
    assert chosen.replace(f"☒ {options[ticked]}", f"☐ {options[ticked]}") == blank


def test_a_drop_down_is_read_as_its_chosen_entry() -> None:
    def dropdown(result: str) -> str:
        entries = "".join(f'<w:listEntry w:val="{v}"/>' for v in ("England and Wales", "Scotland", "Northern Ireland"))
        return mark("begin", f'<w:ffData><w:name w:val="Law"/><w:ddList>{result}{entries}</w:ddList></w:ffData>') + instruction(" FORMDROPDOWN ") + mark("end")

    assert body_text(docx(p(t("Law: "), dropdown('<w:result w:val="1"/>'), t(" is the law.")))) == "Law: Scotland is the law."
    assert body_text(docx(p(t("Law: "), dropdown(""), t(" is the law.")))) == "Law: England and Wales is the law."
    assert coverage(docx(p(dropdown(""))))["fields"][0] is Status.PARTIAL


def sym(font: str, char: str) -> str:
    return f'<w:r><w:sym w:font="{font}" w:char="{char}"/></w:r>'


@pytest.mark.parametrize(
    ("font", "char", "shown"),
    [
        ("Wingdings", "F0FE", "☑"),
        ("Wingdings", "F0A8", "☐"),
        ("Wingdings", "F0FD", "☒"),
        ("Wingdings 2", "F052", "☑"),
        ("Wingdings 2", "F0A3", "☐"),
        ("Symbol", "F0B3", "≥"),
        ("Symbol", "F0A3", "≤"),
        ("Symbol", "F0B1", "±"),
        ("Symbol", "F06D", "μ"),
        ("Arial", "00A9", "©"),
    ],
)
def test_a_symbol_is_read_as_the_glyph_its_font_draws(font: str, char: str, shown: str) -> None:
    data = docx(p(t("a tolerance of "), sym(font, char), t("0.25")))
    assert body_text(data) == f"a tolerance of {shown}0.25"
    assert coverage(data)["symbols"][:2] == (Status.READ, 1)


def test_a_symbol_this_reader_does_not_know_is_read_as_a_replacement_and_counted() -> None:
    data = docx(p(t("Mark: "), sym("Wingdings", "F041"), t(" here")))
    assert body_text(data) == "Mark: \ufffd here"
    assert coverage(data)["symbols"][:2] == (Status.PARTIAL, 1)


def test_text_set_in_a_symbol_font_is_read_as_that_fonts_glyphs() -> None:
    wingdings = '<w:rFonts w:ascii="Wingdings" w:hAnsi="Wingdings"/>'
    data = docx(p(t("þ", wingdings), t(" The Agreement renews automatically.")) + p(t("¨", wingdings), t(" The Agreement does not renew.")))
    assert body_text(data) == "☑ The Agreement renews automatically.\n\n☐ The Agreement does not renew."


# ------------------------------------------------------------------------------------- (3) one slot per grid column


@st.composite
def grids(draw: st.DrawFn) -> tuple[str, list[list[str]]]:
    """A table of random rows: empty and repeated values, spans, vertical merges, w:gridBefore, cells wrapped in a
    content control or a w:customXml. Returns the markup and each row's expected slots."""
    width = draw(st.integers(min_value=2, max_value=5))
    values = st.sampled_from(["", "", "$100,000", "$100,000", "3%", "5%", "Yes", "No", "Included", "N/a"])
    rows_xml, expected = [], []
    for _ in range(draw(st.integers(min_value=1, max_value=5))):
        before = draw(st.integers(min_value=0, max_value=1))
        column, cells, slots = before, [], [""] * before
        while column < width:
            span = draw(st.integers(min_value=1, max_value=min(2, width - column)))
            continuation = column == 0 and draw(st.booleans())
            text = draw(values)
            properties = (f'<w:gridSpan w:val="{span}"/>' if span > 1 else "") + ("<w:vMerge/>" if continuation else "")
            cell = f"<w:tc>{f'<w:tcPr>{properties}</w:tcPr>' if properties else ''}{p(t(text)) if text else '<w:p/>'}</w:tc>"
            wrapper = draw(st.sampled_from(["", "", "sdt", "customXml"]))
            if wrapper == "sdt":
                cell = f"<w:sdt><w:sdtPr/><w:sdtContent>{cell}</w:sdtContent></w:sdt>"
            elif wrapper == "customXml":
                cell = f'<w:customXml w:element="cell">{cell}</w:customXml>'
            cells.append(cell)
            slots.extend(["" if continuation else text] + [""] * (span - 1))
            column += span
        grid_before = f'<w:trPr><w:gridBefore w:val="{before}"/></w:trPr>' if before else ""
        rows_xml.append(f"<w:tr>{grid_before}{''.join(cells)}</w:tr>")
        while slots and not slots[-1]:
            slots.pop()
        expected.append(slots)
    return f"<w:tbl>{''.join(rows_xml)}</w:tbl>", expected


def slots_of(row: str) -> list[str]:
    """A stored row back into its slots: a block's leading space is stripped, so a row opening with an empty slot starts '|'."""
    return (f" {row}" if row.startswith("|") else row).split(" | ")


@settings(max_examples=150 * FUZZ_SCALE, deadline=None)
@given(grids())
def test_every_cell_of_a_row_stays_in_its_column(grid: tuple[str, list[list[str]]]) -> None:
    markup, expected = grid
    rows = ingest("table.docx", docx(p(t("Pricing.")) + markup)).sections[0].paragraphs[1:]
    assert [slots_of(row) for row in rows] == [slots for slots in expected if any(slots)]


def test_the_pricing_row_keeps_its_repeated_and_empty_cells() -> None:
    table = (
        "<w:tbl>"
        + "".join(
            "<w:tr>" + "".join(f"<w:tc>{p(t(c)) if c else '<w:p/>'}</w:tc>" for c in row) + "</w:tr>"
            for row in (["Item", "Year 1", "Year 2", "Year 3"], ["Licence fee", "$100,000", "$100,000", "$150,000"], ["Price increase cap", "", "3%", "5%"])
        )
        + "</w:tbl>"
    )
    text = body_text(docx(p(t("Fees.")) + table))
    assert "Licence fee | $100,000 | $100,000 | $150,000" in text and "Price increase cap |  | 3% | 5%" in text
    assert locate("Licence fee | $100,000 | $150,000", text) is None


def test_a_vertically_merged_cell_is_read_once_in_its_first_row() -> None:
    table = (
        '<w:tbl><w:tr><w:tc><w:tcPr><w:vMerge w:val="restart"/></w:tcPr>'
        + p(t("Fees"))
        + "</w:tc><w:tc>"
        + p(t("R1C1"))
        + "</w:tc></w:tr><w:tr><w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p/></w:tc><w:tc>"
        + p(t("R2C1"))
        + "</w:tc></w:tr></w:tbl>"
    )
    text = body_text(docx(p(t("Table.")) + table))
    assert "Fees | R1C1" in text and "| R2C1" in text and text.count("Fees") == 1


def test_a_cell_inside_a_control_is_read_and_counted() -> None:
    row = f"<w:tbl><w:tr><w:tc>{p(t('Monthly Uptime'))}</w:tc><w:sdt><w:sdtPr/><w:sdtContent><w:tc>{p(t('99.9%'))}</w:tc></w:sdtContent></w:sdt></w:tr></w:tbl>"
    data = docx(p(t("Service levels.")) + row)
    assert "Monthly Uptime | 99.9%" in body_text(data)
    assert coverage(data)["content_controls"][:2] == (Status.READ, 1)


# ------------------------------------------------------------------------------------- (4) displayed content reported

BOX_TEXT = "Notwithstanding Section 2, Supplier's aggregate liability for data breaches is unlimited."


def text_box(text: str) -> str:
    content = f"<w:txbxContent>{p(t(text))}</w:txbxContent>"
    dml = (
        '<w:drawing><wp:anchor><wp:extent cx="1" cy="1"/><wp:docPr id="1" name="Text Box 1"/><a:graphic>'
        '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f"<wps:wsp><wps:txbx>{content}</wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>"
    )
    vml = f"<w:pict><v:shape><v:textbox>{content}</v:textbox></v:shape></w:pict>"
    return f"<w:r>{alternate(dml, vml, requires='wps')}</w:r>"


def test_a_text_box_is_read_once_as_its_own_text_after_its_paragraph() -> None:
    data = docx(p(t("Supplier's aggregate liability shall not exceed "), text_box(BOX_TEXT), t("the fees paid.")) + p(t("Next clause.")))
    sections = ingest("box.docx", data).sections
    paragraphs = [para for s in sections for para in s.paragraphs]
    assert paragraphs[:3] == ["Supplier's aggregate liability shall not exceed the fees paid.", BOX_TEXT, "Next clause."]
    assert coverage(data)["text_boxes"][:2] == (Status.READ, 1)


def test_a_vml_text_box_alone_is_read() -> None:
    vml = f"<w:r><w:pict><v:shape><v:textbox><w:txbxContent>{p(t(BOX_TEXT))}</w:txbxContent></v:textbox></v:shape></w:pict></w:r>"
    text = body_text(docx(p(t("Liability."), vml)))
    assert text.count(BOX_TEXT) == 1


def test_a_block_custom_xml_element_is_read() -> None:
    data = docx(p(t("1. Term")) + f'<w:customXml w:element="clause">{p(t("The Initial Term is thirty-six (36) months."))}</w:customXml>')
    assert "The Initial Term is thirty-six (36) months." in body_text(data)


def test_a_picture_and_an_equation_are_counted_not_silently_dropped() -> None:
    picture = '<w:r><w:drawing><wp:inline><wp:extent cx="1" cy="1"/><wp:docPr id="2" name="Picture 2"/></wp:inline></w:drawing></w:r>'
    equation = "<m:oMath><m:r><m:t>x=2</m:t></m:r></m:oMath>"
    data = docx(p(t("See the figure "), picture, t(" and "), equation, t(".")))
    parts = coverage(data)
    assert parts["embedded_objects"][:2] == (Status.OMITTED, 1)
    assert parts["other_content"][0] is Status.OMITTED and "equations (1)" in parts["other_content"][2]


def test_control_a_plain_document_reports_none_of_the_new_parts() -> None:
    parts = coverage(docx(p(t("1. Term")) + p(t("The term is one year."))))
    assert not {"text_boxes", "symbols", "bound_controls", "invisible_text", "other_content"} & set(parts)


# ------------------------------------------------------------------------------------- (5) text that cannot be seen

WHITE = '<w:color w:val="FFFFFF"/>'
TINY = '<w:sz w:val="1"/>'


@pytest.mark.parametrize("properties", [WHITE, TINY, '<w:color w:val="ffffff" w:themeColor="background1"/>', '<w:sz w:val="2"/>'])
def test_text_drawn_so_it_cannot_be_seen_is_left_out_and_counted(properties: str) -> None:
    data = docx(p(t("Supplier may terminate on "), t("thirty (30) days notice. Supplier may terminate on ", properties), t("ninety (90) days notice.")))
    text = body_text(data)
    assert "thirty" not in text and " ".join(text.split()) == "Supplier may terminate on ninety (90) days notice."
    assert coverage(data)["invisible_text"][:2] == (Status.EXCLUDED, 1)


def test_white_text_from_a_character_style_is_left_out() -> None:
    style = '<w:style w:type="character" w:styleId="Ghost"><w:name w:val="Ghost"/><w:rPr><w:color w:val="FFFFFF"/></w:rPr></w:style>'
    data = docx(p(t("Notice: "), t("thirty (30)", '<w:rStyle w:val="Ghost"/>'), t(" ninety (90) days")), styles=style)
    assert "thirty" not in body_text(data)


@pytest.mark.parametrize(
    ("run", "paragraph"),
    [
        (WHITE + '<w:highlight w:val="black"/>', ""),
        (WHITE + '<w:shd w:val="clear" w:color="auto" w:fill="1F3864"/>', ""),
        (WHITE, '<w:shd w:val="clear" w:color="auto" w:fill="000000"/>'),
        ('<w:color w:val="auto"/>', ""),
        ('<w:color w:val="FF0000"/>', ""),
        ('<w:sz w:val="20"/>', ""),
    ],
)
def test_control_text_that_can_be_seen_is_read(run: str, paragraph: str) -> None:
    data = docx(p(t("Fees: "), t("ninety (90) days", run), properties=paragraph))
    assert "ninety (90) days" in body_text(data)
    assert "invisible_text" not in coverage(data)


def test_control_white_text_in_a_table_cell_or_on_a_page_background_is_read() -> None:
    cell = f"<w:tbl><w:tr><w:tc>{p(t('Header', WHITE))}</w:tc></w:tr></w:tbl>"
    assert "Header" in body_text(docx(p(t("Table.")) + cell))
    background = docx(p(t("Fees: "), t("ninety (90) days", WHITE)), head='<w:background w:color="1F3864"/>')
    assert "ninety (90) days" in body_text(background)


def test_control_white_spaces_between_words_are_kept() -> None:
    assert body_text(docx(p(t("ninety"), t(" ", WHITE), t("(90) days")))) == "ninety (90) days"
