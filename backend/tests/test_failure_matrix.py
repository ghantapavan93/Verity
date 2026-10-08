"""The failure matrix: what the system does when something goes wrong. One test per row.

Rows proven elsewhere: invalid model output → failed run, never a finding (test_api_flow);
fabricated quote → finding withheld (test_api_flow); one changed digit → verifier rejects
(test_spans); duplicate POST → the same run (test_api_flow); SSE replays recorded stages
(test_api_flow); reload restores a finished run from the URL (useUrlState.test.tsx).
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError

from app import db as db_module
from app.analysis import service as analysis_service
from app.application import ingest_document as ingest_module
from app.application.start_run import start_run
from app.ingest import PARSER_VERSION
from app.runs.service import execute_run
from tests.support import CONTRACT, upload_and_ask


def test_provider_unavailable_is_a_failed_run_that_says_why(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(analysis_service, "PROVIDER_RETRY_DELAY_S", 0.0)
    client.provider.fail_always = True
    result = upload_and_ask(client, "What law governs this agreement?", with_guidance=False)
    run = result["run"]
    assert run["stage"] == "failed" and run["reason"] == "provider_error" and "not reachable" in run["error"]
    assert run["findings"] == [] and run["withheld"] == []
    assert len(client.provider.calls) == 2, "one retry, then the failure is the run's"
    stages = {s["stage"]: s for s in client.get(f"/api/runs/{result['run_id']}/detail").json()["stages"]}
    assert stages["checking"]["status"] == "failed" and stages["checking"]["errorCode"] == "provider_error"
    assert stages["failed"]["status"] == "failed" and stages["failed"]["completedAt"]


def test_a_transient_provider_failure_is_retried_once_and_recorded(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(analysis_service, "PROVIDER_RETRY_DELAY_S", 0.0)
    client.provider.fail_once = True
    result = upload_and_ask(client, "What law governs this agreement?", with_guidance=False)
    assert result["run"]["stage"] == "complete"
    stages = {s["stage"]: s for s in client.get(f"/api/runs/{result['run_id']}/detail").json()["stages"]}
    assert stages["checking"]["attempt"] == 2 and "attempt 2" in stages["checking"]["detail"]


def test_every_stage_row_is_closed_with_timing_status_and_hashes(client: TestClient) -> None:
    result = upload_and_ask(client, "How much notice does the customer need to give to terminate for convenience?")
    stages = client.get(f"/api/runs/{result['run_id']}/detail").json()["stages"]
    assert [s["stage"] for s in stages] == ["reading", "finding_evidence", "checking", "verifying", "complete"]
    for row in stages:
        assert row["status"] == "ok" and row["completedAt"] and row["durationMs"] is not None, row
    for row in stages[:-1]:
        assert len(row["inputHash"]) == 64 and len(row["outputHash"]) == 64, row["stage"]
    assert stages[0]["inputHash"] == result["document"]["sha256"], "reading consumes the document by hash"
    assert stages[2]["outputHash"] == stages[3]["inputHash"], "verifying consumes exactly what the model produced"


def test_malformed_and_empty_files_are_refused_at_the_boundary(client: TestClient) -> None:
    docx = client.post("/api/documents", files={"file": ("broken.docx", b"PK\x03\x04 this is not a package", "application/octet-stream")})
    assert docx.status_code == 422 and "not a readable .docx" in docx.json()["detail"]
    pdf = client.post("/api/documents", files={"file": ("broken.pdf", b"%PDF-1.4 garbage without a trailer", "application/pdf")})
    assert pdf.status_code == 422 and "not a readable PDF" in pdf.json()["detail"]
    empty = client.post("/api/documents", files={"file": ("empty.txt", b"", "text/plain")})
    assert empty.status_code == 422 and empty.json()["detail"] == "the file is empty"


def test_an_oversized_upload_is_refused_before_parsing(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(ingest_module, "MAX_UPLOAD_BYTES", 100)
    refused = client.post("/api/documents", files={"file": ("agreement.txt", CONTRACT.encode("utf-8"), "text/plain")})
    assert refused.status_code == 413


def test_a_memo_is_written_once_per_run(client: TestClient) -> None:
    result = upload_and_ask(client, "How much notice does the customer need to give to terminate for convenience?")
    first = client.post("/api/memos", json={"runId": result["run_id"]})
    second = client.post("/api/memos", json={"runId": result["run_id"]})
    assert (first.status_code, second.status_code) == (201, 200)
    assert first.json()["id"] == second.json()["id"] and first.json()["docxSha256"] == second.json()["docxSha256"]


def test_a_finished_run_cannot_be_changed_or_deleted_in_the_database(client: TestClient) -> None:
    result = upload_and_ask(client, "How much notice does the customer need to give to terminate for convenience?")
    run_id = result["run_id"]
    statements = (
        "UPDATE runs SET note = 'edited' WHERE id = :id",
        "DELETE FROM runs WHERE id = :id",
        "UPDATE findings SET status = 'pass' WHERE run_id = :id",
        "DELETE FROM findings WHERE run_id = :id",
        "UPDATE run_stages SET detail = 'edited' WHERE run_id = :id",
        "UPDATE evidence_spans SET verified = 0 WHERE finding_id IN (SELECT id FROM findings WHERE run_id = :id)",
    )
    for statement in statements:
        with pytest.raises(IntegrityError, match="immutable"), db_module.engine.begin() as connection:
            connection.execute(text(statement), {"id": run_id})
    assert client.get(f"/api/runs/{run_id}").json()["stage"] == "complete"


def test_a_prompt_change_is_a_new_run_and_the_old_one_is_untouched(client: TestClient) -> None:
    result = upload_and_ask(client, "How much notice does the customer need to give to terminate for convenience?")
    before = client.get(f"/api/runs/{result['run_id']}/detail").json()
    with db_module.SessionLocal() as session:
        started = start_run(session, result["document"]["id"], None, result["run"]["question"], client.provider, prompt_version="answer-v1")
        assert started.created and started.run.id != result["run_id"]
        execute_run(session, started.run.id, client.provider)
        other = client.get(f"/api/runs/{started.run.id}").json()
    assert other["promptVersion"] == "answer-v1" and other["promptHash"] != before["promptHash"]
    assert client.get(f"/api/runs/{result['run_id']}/detail").json() == before


def test_the_reader_version_travels_with_the_document(client: TestClient) -> None:
    uploaded = client.post("/api/documents", files={"file": ("agreement.txt", CONTRACT.encode("utf-8"), "text/plain")}).json()
    assert uploaded["parserVersion"] == PARSER_VERSION


def test_a_docx_that_unpacks_far_beyond_its_size_is_refused_before_it_is_parsed(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    from app.ingest import readers as readers_module

    monkeypatch.setattr(readers_module, "MAX_DOCX_UNPACKED_BYTES", 1_000_000)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<w:document>" + " " * 3_000_000 + "</w:document>")
    assert len(buffer.getvalue()) < 20_000, "the bomb itself is small; only its declared size is large"
    response = client.post("/api/documents", files={"file": ("bomb.docx", buffer.getvalue(), "application/octet-stream")})
    assert response.status_code == 413
    assert "unpacks to" in response.text


def test_a_pdf_with_more_pages_than_any_contract_is_refused(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    from pypdf import PdfWriter

    from app.ingest import readers as readers_module

    monkeypatch.setattr(readers_module, "MAX_PDF_PAGES", 5)
    writer = PdfWriter()
    for _ in range(6):
        writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    response = client.post("/api/documents", files={"file": ("long.pdf", buffer.getvalue(), "application/pdf")})
    assert response.status_code == 413
    assert "pages" in response.text


def test_large_responses_are_compressed_and_the_event_stream_is_not(client: TestClient) -> None:
    run = upload_and_ask(client, "What is the cap on each party's liability?", with_guidance=False)
    # The fixture contract's JSON is under the 1 KB threshold; a longer upload is what compression is for.
    uploaded = client.post("/api/documents", files={"file": ("long.txt", (CONTRACT + "\n") * 6, "text/plain")})
    document_id = uploaded.json()["id"]
    document = client.get(f"/api/documents/{document_id}", headers={"Accept-Encoding": "gzip"})
    assert document.status_code == 200
    assert document.headers.get("content-encoding") == "gzip"
    assert document.json()["id"] == document_id  # the client decodes it; the payload is intact
    with client.stream("GET", f"/api/runs/{run['run_id']}/events", headers={"Accept-Encoding": "gzip"}) as events:
        assert events.status_code == 200
        assert events.headers["content-type"].startswith("text/event-stream")
        assert "content-encoding" not in events.headers


def test_an_encrypted_pdf_is_refused_with_the_reason(client: TestClient) -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("secret")
    buffer = io.BytesIO()
    writer.write(buffer)
    response = client.post("/api/documents", files={"file": ("locked.pdf", buffer.getvalue(), "application/pdf")})
    assert response.status_code == 422
    assert "encrypted" in response.text


def test_a_macro_enabled_document_is_not_a_supported_type(client: TestClient) -> None:
    response = client.post("/api/documents", files={"file": ("macros.docm", b"PKwhatever", "application/vnd.ms-word.document.macroEnabled.12")})
    assert response.status_code == 422
    assert "unsupported file type .docm" in response.text


def test_reading_a_docx_with_an_external_relationship_makes_no_network_call(monkeypatch: pytest.MonkeyPatch) -> None:
    import socket

    from docx import Document as DocxDocument
    from docx.opc.constants import RELATIONSHIP_TYPE

    from app.ingest import ingest

    document = DocxDocument()
    document.add_paragraph("1. Term")
    document.add_paragraph("This Agreement runs for two years from the Effective Date.")
    document.part.relate_to("http://example.invalid/terms", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    buffer = io.BytesIO()
    document.save(buffer)

    def refuse(*_args: object, **_kwargs: object) -> None:
        raise AssertionError("the reader tried to open a network connection")

    monkeypatch.setattr(socket, "create_connection", refuse)
    monkeypatch.setattr(socket, "getaddrinfo", refuse)
    parsed = ingest("template-linked.docx", buffer.getvalue())
    assert any("two years" in s.text for s in parsed.sections)


def test_a_docx_member_named_to_escape_the_archive_is_never_written_anywhere(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from app.ingest import UnsupportedFile, ingest

    monkeypatch.chdir(tmp_path)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("../../escaped.txt", "should never land on disk")
        archive.writestr("word/document.xml", "<w:document/>")
    with pytest.raises(UnsupportedFile):
        ingest("slip.docx", buffer.getvalue())
    assert not (tmp_path.parent / "escaped.txt").exists() and not list(tmp_path.iterdir())


def test_a_two_megabyte_paragraph_ingests_quickly_into_sections_that_fit_the_prompt() -> None:
    """Before 2026-10-01 the wall stayed one section and the prompt cut it at its window, so everything past the window
    was unreadable by the model. Now the wall is split at sentence ends into parts no longer than a section, each of
    which fits the prompt whole, and nothing is lost."""
    import time

    from app.analysis.service import MAX_SECTION_CHARS_IN_PROMPT, SectionForPrompt, build_user_message
    from app.ingest import ingest
    from app.ingest.sections import MAX_SECTION_CHARS

    sentence = "The Provider shall deliver the services under this Agreement and any Order Form. "
    data = (sentence * 25_000).encode()
    started = time.perf_counter()
    parsed = ingest("wall.txt", data)
    assert time.perf_counter() - started < 2.0
    assert len(parsed.sections) > 300 and all(len(s.text) <= MAX_SECTION_CHARS for s in parsed.sections)
    assert " ".join(s.text for s in parsed.sections) == (sentence * 25_000).strip()
    message = build_user_message("What must the Provider deliver?", None, "wall.txt", [SectionForPrompt("sec_0", "", "", parsed.sections[0].text)])
    assert len(message) < MAX_SECTION_CHARS_IN_PROMPT + 400


def test_section_text_a_finished_run_cited_cannot_be_changed_in_the_database(client: TestClient) -> None:
    """Found by the independent review of 2026-09-29: the record's citations point into section text, which had no trigger."""
    run = upload_and_ask(client, "How much notice is required to terminate for convenience?", with_guidance=False)
    document_id = client.get(f"/api/runs/{run['run_id']}/detail").json()["documentId"]
    with pytest.raises(IntegrityError, match="immutable"), db_module.engine.begin() as connection:
        connection.execute(text("UPDATE sections SET text = 'rewritten' WHERE document_id = :d"), {"d": document_id})
    with pytest.raises(IntegrityError, match="immutable"), db_module.engine.begin() as connection:
        connection.execute(text("DELETE FROM documents WHERE id = :d"), {"d": document_id})


def test_a_pass_that_names_a_section_the_document_does_not_have_is_lowered_to_needs_review(client: TestClient) -> None:
    """Golden g44's shape: a verified quote under a conclusion that points at a clause the contract lacks (docs/GOLDENS.md)."""
    client.provider.invented_reference = True
    run = upload_and_ask(client, "Section 14.2 grants the Customer most favoured nation pricing. Quote that clause.", with_guidance=False)
    detail = client.get(f"/api/runs/{run['run_id']}/detail").json()
    finding = detail["findings"][0]
    assert finding["status"] == "needs_review" and finding["statusSource"] == "reference_check"
    assert finding["spans"][0]["verified"] is True  # the quote is real; the reference is not
