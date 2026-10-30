"""Reader v13: a DOCX heading carries a number only when the document shows one (legal evaluation B2, 2026-10-14).

The evaluation's distribution agreement (Word headings "1. Definitions" … "9. Governing Law", then "SCHEDULE A") was read
with Schedule A as section "9.1": the reader computed a number from heading levels for every Word heading whose text
states none, whether or not Word numbers it. "9.1" is the governing-law clause's own number in that agreement, so the
card said "CODE Found in the document text at §9.1 · SCHEDULE A", the memo "Observed language (§9.1 SCHEDULE A)", and the
evidence pack's sections carried the same number.

The invariant: a section's number is one the document shows, either written at the head of its heading or printed by
Word's own numbering as a clause number (a w:numPr on the paragraph or its style chain naming a list level that exists,
is neither a bullet nor "none", prints something, and prints no word but Section, Clause or Article). A heading that
states none and that Word does not number that way is labelled by its heading alone, and leaves the count of the
numbered headings around it untouched. The section's identity (its id, its ordinal, the internal label sec_N handed to
the model) is not its citation label.

Not closed here, and said so in docs/FAILURE-ENVELOPE.md: a computed number is the count of headings by level, not Word's
own rendering of its list (start values, restarts, a level printed in roman numerals or letters); and a schedule list
Word prints as "Schedule 1" is labelled by its heading alone, without the "Schedule 1" Word shows.
"""

from __future__ import annotations

import io
import time

import pytest
from docx import Document as DocxDocument
from docx.document import Document as WordDocument
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from hypothesis import given, settings
from hypothesis import strategies as st

from app.ingest import PARSER_VERSION, ingest
from tests.support import FUZZ_SCALE


def _numpr(num_id: str | None, level: str | None = "0") -> str:
    parts = []
    if level is not None:
        parts.append(f'<w:ilvl w:val="{level}"/>')
    if num_id is not None:
        parts.append(f'<w:numId w:val="{num_id}"/>')
    return f"<w:numPr {nsdecls('w')}>{''.join(parts)}</w:numPr>"


def _list(document: WordDocument, levels: list[tuple[str, str]] | None = None) -> str:
    """A numbering list of the given (numFmt, lvlText) levels, decimal clause numbers by default; its instance id."""
    levels = levels or [("decimal", "%1."), ("decimal", "%1.%2"), ("decimal", "%1.%2.%3")]
    root = document.part.numbering_part.element
    taken = {int(n.get(qn("w:numId"))) for n in root.iter(qn("w:num"))} | {int(a.get(qn("w:abstractNumId"))) for a in root.iter(qn("w:abstractNum"))}
    key = str(max(taken) + 1)
    lvls = "".join(
        f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{text}"/></w:lvl>' for i, (fmt, text) in enumerate(levels)
    )
    root.insert(0, parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{key}">{lvls}</w:abstractNum>'))
    root.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="{key}"><w:abstractNumId w:val="{key}"/></w:num>'))
    return key


def _number_style(document: WordDocument, name: str, num_id: str | None = None) -> str:
    """Give a heading style Word's automatic numbering, as many templates do: a decimal clause list unless one is named."""
    num_id = num_id or _list(document)
    document.styles[name].element.get_or_add_pPr().append(parse_xml(_numpr(num_id)))
    return num_id


def _number_paragraph(paragraph: object, num_id: str | None, level: str | None = "0") -> None:
    paragraph._p.get_or_add_pPr().append(parse_xml(_numpr(num_id, level)))  # type: ignore[attr-defined]


def _read(document: WordDocument) -> list[tuple[str, str]]:
    buffer = io.BytesIO()
    document.save(buffer)
    return [(s.number, s.heading) for s in ingest("doc.docx", buffer.getvalue()).sections]


def evaluation_agreement() -> WordDocument:
    """The shape of the evaluation's h2-dist.docx: numbered clause headings written in the text, then a schedule."""
    document = DocxDocument()
    document.add_paragraph().add_run("EXCLUSIVE DISTRIBUTION AGREEMENT").bold = True
    for number, title in [(8, "Precedence"), (9, "Governing Law")]:
        document.add_heading(f"{number}. {title}", level=2)
        document.add_paragraph(f"{number}.1 This clause {title.lower()} applies.")
    document.add_heading("SCHEDULE A", level=2)
    document.add_paragraph("Territory: The United States of America, excluding its territories and possessions.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Contract Year 1"
    table.rows[0].cells[1].text = "US$1,200,000"
    return document


def test_b2_an_unnumbered_schedule_is_not_given_the_clause_number_before_it() -> None:
    read = _read(evaluation_agreement())
    assert ("", "SCHEDULE A") in read, read
    assert ("9.1", "SCHEDULE A") not in read
    assert read[1:3] == [("8", "Precedence"), ("9", "Governing Law")]


def test_no_section_takes_a_number_another_section_states_or_its_text_writes() -> None:
    """A citation label is never a clause number the document gives to other words: "9.1 This clause …" is written in
    section 9, so no section may be cited as §9.1."""
    buffer = io.BytesIO()
    evaluation_agreement().save(buffer)
    sections = ingest("doc.docx", buffer.getvalue()).sections
    numbers = [s.number for s in sections if s.number]
    written = {line.split(" ", 1)[0] for s in sections for line in s.text.split("\n") if line[:1].isdigit()}
    assert len(numbers) == len(set(numbers)), numbers
    assert not set(numbers) & written, (numbers, written)


def test_the_schedule_keeps_its_text_and_its_table_in_order() -> None:
    buffer = io.BytesIO()
    evaluation_agreement().save(buffer)
    sections = ingest("doc.docx", buffer.getvalue()).sections
    schedule = next(s for s in sections if s.heading == "SCHEDULE A")
    assert schedule.text.index("Territory:") < schedule.text.index("Contract Year 1 | US$1,200,000")


def test_control_word_numbered_headings_still_get_their_computed_numbers() -> None:
    document = DocxDocument()
    _number_style(document, "Heading 1")
    document.add_paragraph("MASTER AGREEMENT")
    for title in ["Definitions", "Payment", "Term"]:
        document.add_heading(title, level=1)
        document.add_paragraph(f"The {title.lower()} clause.")
    assert _read(document)[1:] == [("1", "Definitions"), ("2", "Payment"), ("3", "Term")]


def test_an_unnumbered_appendix_between_word_numbered_headings_does_not_take_a_count() -> None:
    """Word would show 1, 2 around an unnumbered annex; v12 showed 1, 2 and 3, counting the annex as clause 2."""
    document = DocxDocument()
    _number_style(document, "Heading 1")
    document.add_heading("Definitions", level=1)
    document.add_paragraph("Terms.")
    annex = document.add_heading("ANNEX I", level=1)
    _number_paragraph(annex, "0", None)  # the paragraph removes the style's numbering, as Word writes it
    document.add_paragraph("Annexed terms.")
    document.add_heading("Payment", level=1)
    document.add_paragraph("Fees.")
    assert _read(document) == [("1", "Definitions"), ("", "ANNEX I"), ("2", "Payment")]


def test_a_paragraphs_own_numbering_counts_and_a_level_alone_takes_the_styles_instance() -> None:
    document = DocxDocument()
    first = document.add_heading("Scope", level=1)
    _number_paragraph(first, _list(document))  # numbered on the paragraph itself
    document.add_paragraph("Scope text.")
    plain = document.add_heading("Recitals", level=1)  # neither the paragraph nor its style numbers it
    document.add_paragraph("Whereas.")
    assert _read(document) == [("1", "Scope"), ("", "Recitals")]

    styled = DocxDocument()
    _number_style(styled, "Heading 1")
    level_only = styled.add_heading("Services", level=1)
    _number_paragraph(level_only, None, "0")  # a level alone: the instance is the style's
    styled.add_paragraph("Services text.")
    assert _read(styled) == [("1", "Services")]
    assert plain is not None


@pytest.mark.parametrize("heading", ["SCHEDULE A", "EXHIBIT B", "Appendix 1", "ANNEX I", "AMENDMENT NO. 1", "Recitals"])
def test_unnumbered_schedules_exhibits_appendices_and_amendments_are_labelled_by_their_heading(heading: str) -> None:
    document = DocxDocument()
    document.add_heading("12. General", level=1)
    document.add_paragraph("12.1 This Agreement is governed by the laws of Ontario.")
    document.add_heading(heading, level=1)
    document.add_paragraph("1. Section 6.1 of the Lease is amended as follows.")
    read = _read(document)
    assert read[-1] == ("", heading), read


def test_repeated_headings_stay_apart_and_in_order() -> None:
    document = DocxDocument()
    _number_style(document, "Heading 1")
    for title in ["Definitions", "Definitions"]:
        document.add_heading(title, level=1)
        document.add_paragraph("Terms.")
    document.add_heading("SCHEDULE A", level=2)
    document.add_paragraph("First schedule.")
    document.add_heading("SCHEDULE A", level=2)
    document.add_paragraph("Second schedule.")
    buffer = io.BytesIO()
    document.save(buffer)
    sections = ingest("doc.docx", buffer.getvalue()).sections
    assert [(s.number, s.heading) for s in sections] == [("1", "Definitions"), ("2", "Definitions"), ("", "SCHEDULE A"), ("", "SCHEDULE A")]
    assert [s.text for s in sections[2:]] == ["First schedule.", "Second schedule."]


def test_a_number_written_in_the_heading_is_kept_whatever_word_does() -> None:
    document = DocxDocument()
    _number_style(document, "Heading 2")
    document.add_heading("4.2 Fees", level=2)
    document.add_paragraph("Fees text.")
    assert _read(document) == [("4.2", "Fees")]


def test_the_reader_version_names_the_change() -> None:
    # v14 keeps every v13 rule here and prints Word's own numbers (test_post_pilot_repair.py); v15 and v16 keep both
    # (test_reader_v15.py, test_reader_v16_*.py).
    assert PARSER_VERSION in ("v13", "v14", "v15", "v16")


KIND = st.sampled_from(["written", "word", "plain"])


@settings(max_examples=60 * FUZZ_SCALE, deadline=None)
@given(kinds=st.lists(KIND, min_size=1, max_size=8), levels=st.lists(st.integers(1, 3), min_size=8, max_size=8))
def test_every_number_is_written_in_its_heading_or_applied_by_word(kinds: list[str], levels: list[int]) -> None:
    """Over any mix of headings with a written number, Word-numbered headings and plain ones, at any levels: a section
    has a number only if its heading wrote it or Word numbers it, a plain heading never has one, and the Word-numbered
    headings are counted as if the plain ones were not there."""
    document = DocxDocument()
    clauses = _list(document)
    expected_plain = []
    for index, kind in enumerate(kinds):
        level = levels[index]
        title = f"Heading {kind} {index}"
        if kind == "written":
            paragraph = document.add_heading(f"{index + 1}. {title}", level=level)
        else:
            paragraph = document.add_heading(title, level=level)
            if kind == "word":
                _number_paragraph(paragraph, clauses)
            else:
                expected_plain.append(title)
        document.add_paragraph(f"Body {index}.")
    read = {h: n for n, h in _read(document)}
    for title in expected_plain:
        assert read[title] == "", (title, read)
    for kind, index in zip(kinds, range(len(kinds)), strict=True):
        title = f"Heading {kind} {index}"
        if kind == "written":
            assert read[title] == str(index + 1), (title, read)
        elif kind == "word":
            assert read[title] != "", (title, read)


def _bullet(document: WordDocument) -> str:
    return _list(document, [("bullet", "\uf0b7")])


def _none(document: WordDocument) -> str:
    return _list(document, [("none", "")])


def _silent(document: WordDocument) -> str:
    return _list(document, [("decimal", "")])


def _schedule_after_clauses(make: object) -> list[tuple[str, str]]:
    """Two Word-numbered clauses, then "SCHEDULE A" carrying a numbering reference made by ``make``, then a clause."""
    document = DocxDocument()
    _number_style(document, "Heading 1")
    for name in ["Definitions", "Payment"]:
        document.add_heading(name, level=1)
        document.add_paragraph("Terms.")
    heading = document.add_heading("SCHEDULE A", level=1)
    _number_paragraph(heading, make(document) if callable(make) else str(make))
    document.add_paragraph("Schedule text.")
    document.add_heading("Notices", level=1)
    document.add_paragraph("Notice text.")
    return _read(document)


@pytest.mark.parametrize(
    ("case", "make"),
    [
        ("a bullet level", _bullet),
        ("a level whose format is none", _none),
        ("a level that prints nothing", _silent),
        ("a list that does not exist", "777"),
        ("python-docx's own instance 1, a bullet list", "1"),
    ],
)
def test_a_heading_word_shows_no_number_for_gets_none(case: str, make: object) -> None:
    """The reviewer's counterexamples to the first repair: a w:numPr was there, but Word prints no number."""
    read = _schedule_after_clauses(make)
    assert read == [("1", "Definitions"), ("2", "Payment"), ("", "SCHEDULE A"), ("3", "Notices")], (case, read)


def test_a_schedule_list_is_never_labelled_with_clause_numbers() -> None:
    """Schedules on their own list ("Schedule %1") after clauses 1-3: Word prints "Schedule 1", "Schedule 2" (and "1.1"
    under the first); v12 and the first v13 labelled them 4 and 5, numbers the document gives no section."""
    document = DocxDocument()
    _number_style(document, "Heading 1")
    for name in ["Definitions", "Payment", "Term"]:
        document.add_heading(name, level=1)
        document.add_paragraph("Terms.")
    schedules = _list(document, [("decimal", "Schedule %1"), ("decimal", "%1.%2")])
    for name, level in [("Services", "0"), ("Service levels", "1"), ("Fees", "0")]:
        heading = document.add_heading(name, level=1 if level == "0" else 2)
        _number_paragraph(heading, schedules, level)
        document.add_paragraph(f"{name} text.")
    assert _read(document) == [("1", "Definitions"), ("2", "Payment"), ("3", "Term"), ("", "Services"), ("", "Service levels"), ("", "Fees")]


def test_control_an_article_list_and_its_sub_levels_are_clause_numbers() -> None:
    document = DocxDocument()
    articles = _list(document, [("decimal", "Article %1"), ("decimal", "%1.%2")])
    for name, level in [("Definitions", "0"), ("Terms", "1"), ("Payment", "0"), ("Fees", "1")]:
        heading = document.add_heading(name, level=1 if level == "0" else 2)
        _number_paragraph(heading, articles, level)
        document.add_paragraph(f"{name} text.")
    assert _read(document) == [("1", "Definitions"), ("1.1", "Terms"), ("2", "Payment"), ("2.1", "Fees")]


def test_control_a_list_linked_to_a_numbering_style_is_followed() -> None:
    """A list whose definition is a link to a numbering style (w:numStyleLink) prints that style's list."""
    document = DocxDocument()
    target = _list(document)
    root = document.part.numbering_part.element
    style = document.styles.add_style("ClauseNumbering", 3)  # WD_STYLE_TYPE.LIST
    style.element.get_or_add_pPr().append(parse_xml(_numpr(target)))
    key = str(int(target) + 100)
    root.insert(0, parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{key}"><w:numStyleLink w:val="ClauseNumbering"/></w:abstractNum>'))
    root.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="{key}"><w:abstractNumId w:val="{key}"/></w:num>'))
    for name in ["Definitions", "Payment"]:
        heading = document.add_heading(name, level=1)
        _number_paragraph(heading, key)
        document.add_paragraph("Terms.")
    assert _read(document) == [("1", "Definitions"), ("2", "Payment")]


def test_a_numbering_instance_id_of_any_characters_never_refuses_the_file() -> None:
    """The second review: an instance id holding a quote was pasted into a query, and the whole file was refused as not
    a readable package. The definitions are read in one pass; nothing is looked up by a query built from an id."""
    document = DocxDocument()
    root = document.part.numbering_part.element
    level = '<w:lvl w:ilvl="0"><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>'
    root.insert(0, parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="60">{level}</w:abstractNum>'))
    odd = "6'0"
    root.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="{odd}"><w:abstractNumId w:val="60"/></w:num>'))
    _number_paragraph(document.add_heading("Definitions", level=1), odd)
    document.add_paragraph("Terms are defined here.")
    assert ("1", "Definitions") in _read(document)


def test_many_numbering_instances_read_in_time_proportional_to_their_count() -> None:
    """The second review: each instance scanned every other one for its overrides, so 8,000 instances in a 56 KB upload
    read for 54 s, past the 30 s reading bound, holding a reading slot. Read in one pass they take a fraction of a
    second; the bound here is generous so that a slow machine never fails it, and quadratic reading fails it by far."""
    document = DocxDocument()
    num_id = _list(document)
    root = document.part.numbering_part.element
    for extra in range(8000):
        root.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="{100000 + extra}"><w:abstractNumId w:val="{num_id}"/></w:num>'))
    _number_paragraph(document.add_heading("Definitions", level=1), num_id)
    document.add_paragraph("Terms are defined here.")
    started = time.perf_counter()
    sections = _read(document)
    assert time.perf_counter() - started < 10
    assert ("1", "Definitions") in sections
