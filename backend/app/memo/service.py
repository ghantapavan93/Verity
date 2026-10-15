"""The one artifact: a review memo built from verified findings, as HTML and as DOCX.

The DOCX is a working Word document, not a flattened export: built-in paragraph styles, a
sources table each citation links to (internal hyperlinks to bookmarks), a link from every
source back to the exact run and finding in the workbench, core properties naming the run, and a
custom-properties part carrying the run id, its fingerprint, the document hash, the model and
the prompt version and hash, so the file can be traced to its record from inside Word.
"""

from __future__ import annotations

import json
import re
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import datetime
from html import escape
from pathlib import Path
from xml.sax.saxutils import escape as xml_escape

from docx import Document as DocxDocument
from docx.document import Document as WordDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph
from lxml import etree

from ..config import settings
from ..document_names import run_document_name
from ..hashing import sha256_file
from ..models import EvidenceSpan, Finding, Run, as_utc, iso, utcnow

# Characters XML 1.0 cannot carry (control characters other than tab, newline and return). python-docx refuses them
# with a ValueError; a form feed pasted from a PDF into the guidance made the memo a 500 (hostile review, 2026-10-01).
_XML_UNSAFE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def _suggests(finding: Finding) -> bool:
    """A memo carries the model's suggested change only beside a finding that needs review. Elsewhere the clause meets
    the point or no review was asked for, and a change it does not need is noise a reviewer would have to undo; the
    suggestion stays on the record, shown in the workbench as the model's (2026-10-09 evaluation: net 30 days suggested
    against the buyer's own 45)."""
    return bool(finding.suggested_position) and finding.status == "needs_review"


def xml_safe(text: str) -> str:
    return _XML_UNSAFE.sub("", text)


# "Not found" is a statement about the sections the model was handed, never about the agreement (prompt rule 8).
# "(model's view)": only the model decides that a point is absent; code never checks an absence (authority review, 2026-10-14).
STATUS_LABELS = {"needs_review": "Needs review", "missing": "Not found in the sections read (model's view)", "unresolved": "Unresolved"}

# What the memo says about its flags (2026-10-14, the legal pilot). Until then it headed the flagged findings "Issues
# requiring review" and, with none, said "None. Every finding was within guidance or a plain answer.": on real output a
# wrong answer that 30 days met a guidance of 30 Business Days carried exactly that line. A flag is what an automated
# check raised, no more; an unflagged finding is the model's answer, not a cleared one; and no finding is a person's
# decision unless a review is shown.
FLAGS_HEAD = "Flagged by automated checks"
NO_FLAGS = (
    "No automated check flagged a finding. The findings below are the model's answers: code verified only that their "
    "quoted passages appear in the document, not that the answers are right."
)
NOT_REVIEWED = "Not reviewed by a person."
# The status sources under which code, not the model, set the status (application/explain_run DETERMINISTIC_SOURCES).
DECIDED_BY_CODE = frozenset({"computed_days", "position_check", "ambiguous_fact", "reference_check"})


def flags_of(findings: list[Finding], flagged: dict[str, str] | None) -> list[tuple[Finding, str]]:
    """The flagged findings and why. The caller reads each finding's account (runs.account: a status other than a pass,
    or a pass with anything code raised against it) and hands the memo, by finding id, why each flagged one is. Until
    2026-10-14 only a needs-review or not-found finding counted, so a pass whose recorded reason said "code does not
    confirm that as a pass" sat under "No automated check flagged a finding" (authority review). Without it (a caller
    that has none), the status alone decides, as before."""
    if flagged is None:
        return [(f, flag_reason(f)) for f in findings if f.status in ("needs_review", "missing")]
    return [(f, flagged[f.id]) for f in findings if f.id in flagged]


# A proposal whose quotes were not found in the sections the model was handed is kept on the run as withheld, never as a
# finding. Until the first-visitor acceptance test (2026-10-14) the memo left such proposals out altogether: on an amended
# PDF, code withheld the model's renewal-window answer (it quoted the superseded clause and stated 210 and 150 days that
# no verified quote held), and the memo printed "No automated check flagged a finding" above the one finding left. Each is
# now listed as flagged with code's reason, and the model's sentence is printed as the model's. Its quotes are not
# printed: every quote in a memo was found in the document.
WITHHELD_HEAD = "Withheld from the findings"
WITHHELD_LEDE = (
    "The model proposed these, but no quote it gave for them was found in the sections handed to it, so they are not "
    "findings. The model's sentence is printed so that the memo does not read as if it said nothing; it is not an answer, "
    "and code checked none of it."
)
WITHHELD_REASON = "no quote of it was found in the sections handed to the model, so it is withheld"
# A statement of the contract's own about a quoted provision that raises nothing (runs.relations "context"): a later
# replacement the answer uses, or one a question about the original agreement need not. Concerns are under "Why flagged".
RELATED_HEAD = "The contract on this provision (located by code)."


def memo_flags(findings: list[Finding], withheld: Sequence[Finding], flagged: dict[str, str] | None) -> list[tuple[Finding, str]]:
    """Every flag a memo lists: the findings' (`flags_of`), then each withheld proposal, with what code raised about it."""
    return [*flags_of(findings, flagged), *((w, (flagged or {}).get(w.id) or WITHHELD_REASON) for w in withheld)]


def flag_title(finding: Finding) -> str:
    return f"{finding.topic} (withheld)" if finding.status == "unresolved" else finding.topic


_QUOTE_MARKS = re.compile(r"[\"'“”‘’]")


def _plain(text: str) -> str:
    return " ".join(_QUOTE_MARKS.sub("", text).casefold().split())


def in_guidance(reference: str, guidance: str | None) -> bool:
    """Whether the model's reference to the guidance is the guidance's own words (spacing, case and quote marks aside)."""
    return bool(guidance) and bool(_plain(reference)) and _plain(reference) in _plain(guidance or "")


def guidance_line(finding: Finding, guidance: str | None) -> str:
    """The model's reference to the guidance, said to be the model's, and whether code found it in the guidance: the memo
    printed it as "Legal guidance.", and a sentence the guidance does not contain was printed as guidance (2026-10-14)."""
    reference = finding.guidance_reference or ""
    found = "found in the guidance" if in_guidance(reference, guidance) else "not found in the guidance: the model's words"
    return f"{reference} ({found})"


def model_reading(finding: Finding) -> str:
    """The positions the model stated, as its words: 'observed 15 days · required at least 30 days'."""
    return " · ".join(
        part for part in (f"observed {finding.observed}" if finding.observed else "", f"required {finding.required}" if finding.required else "") if part
    )


def flag_reason(finding: Finding) -> str:
    """Why a finding is flagged, in the words that decided it: code's own reason where code set the status, otherwise
    whose call the status is. The reading before accounts (runs.account), for a caller that has none."""
    if finding.status_source in DECIDED_BY_CODE and finding.status_reason:
        return f"automated check: {finding.status_reason}"
    if finding.status == "missing":
        return "the model reported it not found in the sections it was handed"
    if finding.status == "needs_review":
        return "the model asked for review" + (f"; automated check: {finding.status_reason}" if finding.status_reason else "")
    return ""


# The model's own sentence behind a not-found finding, in the memo as on the card: labelled as the model's, after the
# product's sentence, never as the finding (models.Finding.withheld_conclusion; m10, 2026-10-12).
WITHHELD_NOTE = "(not checked by code; withheld from the finding because the model reported the point as not found in the sections it was given)"
# A finding whose quoted sentence points at sections the model was not handed (runs.status.basis): the memo says so beside
# the passages, as the card does, since its status was set without them (release integration, 2026-10-13). The caller
# reads the status and hands the memo what it found, by finding id (`unseen`): the memo sits below the run's rules.
UNSEEN_HEAD = "Not shown to the model."


def unseen_note(unseen: str) -> str:
    return f"Code found that {unseen}. The answer may depend on it, and nothing checked that."


def status_label(status: str, with_guidance: bool, source: str = "computed_days") -> str:
    """A pass is "within guidance" only when the run had guidance to be within; without it the model gave a plain
    answer (prompt rule 5). Found while tracing on 2026-09-29: every pass used to be headed "Within guidance".
    When the pass is the model's own hint rather than a comparison code made, the label says so: the sweep of
    2026-10-01 found "Within guidance" on findings that said the contract provides no notice period at all."""
    if status == "pass":
        if not with_guidance:
            # Code evaluated nothing: no guidance, so the pass is the model's answer and is named as such (2026-10-08).
            return "Model's answer"
        if source == "model_hint":
            return "Within guidance (model's view)"
        # The model proposed the pass and code confirmed it (policy-v2): not a decision code made, and not headed as one.
        return "Within guidance (confirmed by code)" if source == "confirmed_days" else "Within guidance"
    return STATUS_LABELS.get(status, status)


AUTHOR = "Contract Workbench"
CUSTOM_PROPERTIES_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.custom-properties+xml"
CUSTOM_PROPERTIES_FMTID = "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}"


@dataclass(frozen=True)
class Source:
    number: int
    finding: Finding
    span: EvidenceSpan
    title: str


def deep_link(run: Run, finding: Finding) -> str:
    """The workbench URL that reopens this run at this finding, with the cited passage highlighted."""
    return f"{settings.app_url}/?document={run.document_id}&run={run.id}&finding={finding.id}"


def _section_label(span: EvidenceSpan) -> str:
    return span.cited_section_label or ""


def _review_line(finding: Finding) -> str:
    """'Confirmed by A. Reviewer on 28 September 2026: note' from the finding's current review."""
    review = finding.review
    if review is None:
        return ""
    # Stored in UTC and read back without a zone: say it is UTC before converting. Read as local time, a review
    # made in the evening west of Greenwich was dated the next day (2026-10-02).
    when = as_utc(review.created_at).astimezone().strftime("%d %B %Y")
    note = f": {review.note}" if review.note else ""
    return f"{review.verdict.capitalize()} by {review.reviewer} on {when}{note}"


def _located(span: EvidenceSpan) -> str:
    if span.verified:
        return f"found in the document text · {span.method} · characters {span.start} to {span.end}"
    return "not found in the sections handed to the model; withheld"  # only those were searched, never the agreement


def _evidence_head(finding: Finding) -> str:
    return "Closest provision read" if finding.status == "missing" else "Observed language"


def reading_line(run: Run) -> str:
    """How much of the agreement the model was given, and whether those sections were chosen for the question. A memo
    leaves the workbench; until 2026-10-02 it said neither, and "not found" in it read as a search of the agreement."""
    candidates = json.loads(run.candidates_json or "[]")
    read = len(candidates)
    linked = sum(1 for c in candidates if c.get("linked"))
    total = len(run.document.sections)
    scope = "Nothing in this memo is a statement about the other sections."
    if run.retrieval_mode == "opening_fallback":
        return f"{read} of {total} sections: retrieval ranked no section for this question, so these are the opening sections, not ones chosen for it. {scope}"
    if run.retrieval_mode is None:  # a run recorded before the mode was a field: how its sections were chosen is not on the record
        return f"{read} of {total} sections. {scope}"
    if linked:  # retrieval-v2: sections the contract points to from the ranked ones (a definition, an amendment, a clause cited)
        return (
            f"{read} of {total} sections: {read - linked} chosen by retrieval for the question, and {linked} the contract points to from "
            f"them (a definition they use, an amendment of them, or a clause they cite). {scope}"
        )
    return f"{read} of {total} sections, chosen by retrieval for the question. {scope}"


def _sources(findings: list[Finding], section_titles: dict[str, str], clauses: dict[str, str] | None = None) -> list[Source]:
    """``clauses`` names, by span id, the clause a passage sits in when the document's own text names one; the section's
    title follows it, since a clause written inline is body under the reading's last standalone heading (2026-10-08)."""
    sources: list[Source] = []
    for finding in findings:
        for span in finding.spans:
            title = section_titles.get(span.section_id or "", _section_label(span))
            if clauses and (clause := clauses.get(span.id)):
                title = f"{clause} · within {title}"
            sources.append(Source(len(sources) + 1, finding, span, title))
    return sources


# ----------------------------------------------------------------------------- HTML


def memo_html(
    run: Run,
    findings: list[Finding],
    section_titles: dict[str, str],
    review_head: str = "",
    clauses: dict[str, str] | None = None,
    unseen: dict[str, str] | None = None,
    flagged: dict[str, str] | None = None,
    withheld: Sequence[Finding] = (),
    related: dict[str, list[str]] | None = None,
) -> str:
    now = datetime.now().astimezone()
    date = f"{now.day} {now.strftime('%B %Y')}"
    flags = memo_flags(findings, withheld, flagged)
    issues = [f for f, _ in flags]
    why = {f.id: reason for f, reason in flags}
    guidance_text = run.guidance.text if run.guidance else None
    sources = _sources(findings, section_titles, clauses)
    by_span = {id(s.span): s for s in sources}
    parts = [
        "<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
        "<title>Contract review memo</title>",
        # The memo is set like the paper in the workbench: a serif page, sans labels, the cited passage on a brass rule.
        "<style>"
        "html{background:#ece7dc}"
        "body{font:16px/1.65 'Source Serif 4','Source Serif Pro',Georgia,serif;max-width:760px;margin:48px auto;padding:56px 64px 64px;"
        "background:#f6f2e9;color:#1d1b17;box-shadow:0 0 0 1px #e1d9c8,0 30px 60px -30px rgba(0,0,0,.35);border-radius:3px}"
        ".kicker{font:600 11px/1 system-ui,'Segoe UI',sans-serif;letter-spacing:.16em;text-transform:uppercase;color:#8a7a55}"
        "h1{font-size:30px;font-weight:500;letter-spacing:-.01em;line-height:1.15;margin:10px 0 26px}"
        "h2{font-size:18px;font-weight:600;margin:34px 0 8px;padding-top:18px;border-top:1px solid #e1d9c8}"
        "dl{display:grid;grid-template-columns:140px 1fr;gap:8px 16px;margin:0;font-size:14.5px}"
        "dt{font:600 11px/1.9 system-ui,'Segoe UI',sans-serif;letter-spacing:.12em;text-transform:uppercase;color:#6d675c}"
        "dd{margin:0}"
        "blockquote{margin:10px 0 14px;padding:10px 14px;background:#efe8d8;border-left:2px solid #b0852c;border-radius:2px}"
        "small{color:#6d675c;font-family:system-ui,'Segoe UI',sans-serif;font-size:12.5px}"
        "strong{font-family:system-ui,'Segoe UI',sans-serif;font-size:13px;font-weight:600}"
        "table{width:100%;border-collapse:collapse;font:13px/1.5 system-ui,'Segoe UI',sans-serif}"
        "td,th{padding:7px 8px;border-bottom:1px solid #e1d9c8;text-align:left;vertical-align:top}"
        "th{font-size:11px;letter-spacing:.1em;text-transform:uppercase;color:#6d675c}"
        "a{color:#7a5a1c;text-decoration-color:#c9a964;text-underline-offset:3px}"
        "hr{border:0;border-top:1px solid #e1d9c8;margin:36px 0 14px}"
        "@media (max-width:700px){body{margin:0;padding:28px 20px;border-radius:0}dl{grid-template-columns:1fr}}"
        "</style></head><body>",
        "<div class='kicker'>Verity · contract review memo</div><h1>Contract Review Memo</h1>",
        f"<dl><dt>Agreement</dt><dd>{escape(run_document_name(run))}</dd><dt>Date</dt><dd>{date}</dd><dt>Question</dt><dd>{escape(run.question)}</dd>",
        f"<dt>Guidance</dt><dd>{escape(run.guidance.text) if run.guidance else 'none supplied'}</dd>"
        f"<dt>Sections read</dt><dd>{escape(reading_line(run))}</dd></dl>",
        f"<h2>{FLAGS_HEAD} ({len(issues)})</h2>",
    ]
    if issues:
        items = "".join(f"<li>{escape(flag_title(f))}" + (f" — {escape(why[f.id])}" if why[f.id] else "") + "</li>" for f in issues)
        parts.append(f"<ol>{items}</ol>")
    else:
        parts.append(f"<p>{escape(NO_FLAGS)}</p>")
    for f in findings:
        parts.append(f"<h2>{escape(f.topic)} · {status_label(f.status, run.guidance is not None, f.status_source)}</h2>")
        parts.append(f"<p>{escape(f.shown_conclusion)}</p>")
        if f.withheld_conclusion:
            also = escape(f.withheld_conclusion)
            parts.append(f"<p class='withheld'><strong>The model also wrote</strong> <small>{escape(WITHHELD_NOTE)}</small>: {also}</p>")
        for span in f.spans:
            source = by_span[id(span)]
            parts.append(
                f"<p><strong>{_evidence_head(f)}</strong> <small>({escape(source.title)})</small> "
                f"<a href='#src-{source.number}'>[{source.number}]</a></p><blockquote>{escape(span.quote)}</blockquote>"
            )
        if unseen and f.id in unseen:
            parts.append(f"<p class='unseen'><strong>{UNSEEN_HEAD}</strong> {escape(unseen_note(unseen[f.id]))}</p>")
        if f.guidance_reference:
            parts.append(f"<p><strong>The model's reference to the guidance.</strong> {escape(guidance_line(f, guidance_text))}</p>")
        if model_reading(f):
            parts.append(f"<p><strong>The model's reading.</strong> {escape(model_reading(f))}</p>")
        if why.get(f.id):
            parts.append(f"<p><strong>Why flagged.</strong> {escape(why[f.id])}</p>")
        for sentence in (related or {}).get(f.id, []):
            parts.append(f"<p><strong>{RELATED_HEAD}</strong> {escape(sentence)}</p>")
        if _suggests(f):
            parts.append(f"<p><strong>Suggested position (the model's).</strong> {escape(f.suggested_position or '')}</p>")
        if f.review is not None:
            parts.append(f"<p><strong>Reviewed.</strong> {escape(_review_line(f))}</p>")
        else:
            parts.append(f"<p><strong>Person.</strong> {NOT_REVIEWED}</p>")
    if withheld:
        parts.append(f"<h2>{WITHHELD_HEAD} ({len(withheld)})</h2><p><small>{escape(WITHHELD_LEDE)}</small></p>")
        for w in withheld:
            parts.append(
                f"<p class='withheld'><strong>{escape(w.topic)}.</strong> <small>The model wrote, withheld:</small> {escape(w.conclusion)}</p>"
                f"<p><strong>Why withheld.</strong> {escape(why[w.id])}</p>"
            )
    parts.append("<h2>Sources</h2><table><tr><th>#</th><th>Section</th><th>Located</th><th></th></tr>")
    for source in sources:
        parts.append(
            f"<tr id='src-{source.number}'><td>{source.number}</td><td>{escape(source.title)}</td><td>{escape(_located(source.span))}</td>"
            f"<td><a href='{escape(deep_link(run, source.finding))}'>Open in the workbench</a></td></tr>"
        )
    parts.append("</table>")
    parts.append(
        f"<hr><p><small>Generated from run {escape(run.id)} · review state {escape(review_head[:8]) or 'none'} · model {escape(run.model)} · "
        f"prompt {escape(run.prompt_version)} ({escape(run.prompt_hash[:12])}) · "
        "every quoted passage was found in the document text before this memo was written; that a passage supports its finding is the "
        "model's reading and the reviewer's decision.</small></p></body></html>"
    )
    return "".join(parts)


# ----------------------------------------------------------------------------- DOCX


def _ensure_hyperlink_style(document: WordDocument) -> None:
    if "Hyperlink" in [style.name for style in document.styles]:
        return
    style = document.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
    style.font.color.rgb = RGBColor(0x0B, 0x5C, 0xAD)
    style.font.underline = True


def _bookmark(paragraph: Paragraph, name: str, text: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.append(start)
    paragraph.add_run(xml_safe(text))
    paragraph._p.append(end)


def _hyperlink_run(text: str) -> etree._Element:
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    properties.append(style)
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = xml_safe(text)
    node.set(qn("xml:space"), "preserve")
    run.append(node)
    return run


def _internal_link(paragraph: Paragraph, anchor: str, text: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.append(_hyperlink_run(text))
    paragraph._p.append(hyperlink)


def _external_link(paragraph: Paragraph, url: str, text: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    hyperlink.append(_hyperlink_run(text))
    paragraph._p.append(hyperlink)


def _custom_properties(document: WordDocument, values: dict[str, str]) -> None:
    """A docProps/custom.xml part: typed name/value pairs Word shows under File → Info → Properties → Advanced."""
    rows = "".join(
        f'<property fmtid="{CUSTOM_PROPERTIES_FMTID}" pid="{index}" name="{xml_escape(name, {chr(34): "&quot;"})}">'
        f"<vt:lpwstr>{xml_escape(value)}</vt:lpwstr></property>"
        for index, (name, value) in enumerate(values.items(), start=2)
    )
    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">' + rows + "</Properties>"
    )
    package = document.part.package
    part = Part(PackURI("/docProps/custom.xml"), CUSTOM_PROPERTIES_CONTENT_TYPE, xml.encode("utf-8"), package)
    package.relate_to(part, RT.CUSTOM_PROPERTIES)


def _properties(document: WordDocument, run: Run, verified: int, total: int, review_head: str = "") -> None:
    core = document.core_properties
    core.title = xml_safe(f"Contract review memo · {run_document_name(run)}")
    core.subject = xml_safe(run.question)
    core.author = AUTHOR
    core.last_modified_by = AUTHOR
    core.category = "Review memo"
    core.keywords = "contract review; quoted passages; contract workbench"
    core.identifier = run.id
    core.version = run.prompt_version
    core.comments = f"run {run.id} · document sha256 {run.document_sha256} · model {run.model} · prompt {run.prompt_version} {run.prompt_hash}"
    _custom_properties(
        document,
        {
            "workbench_run_id": run.id,
            "workbench_run_fingerprint": run.fingerprint or "",
            "workbench_document_id": run.document_id,
            "workbench_document_sha256": run.document_sha256,
            "workbench_guidance_sha256": run.guidance_sha256 or "",
            "workbench_model": run.model,
            "workbench_prompt_version": run.prompt_version,
            "workbench_prompt_hash": run.prompt_hash,
            "workbench_citations_verified": f"{verified} of {total}",
            "workbench_review_head": review_head,
            "workbench_generated_at": iso(utcnow()) or "",
            "workbench_link": f"{settings.app_url}/?document={run.document_id}&run={run.id}",
        },
    )


def memo_docx(
    run: Run,
    findings: list[Finding],
    section_titles: dict[str, str],
    out_dir: Path,
    review_head: str = "",
    clauses: dict[str, str] | None = None,
    unseen: dict[str, str] | None = None,
    flagged: dict[str, str] | None = None,
    withheld: Sequence[Finding] = (),
    related: dict[str, list[str]] | None = None,
) -> tuple[Path, str]:
    out_dir.mkdir(parents=True, exist_ok=True)
    document = DocxDocument()
    normal = document.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(11)
    _ensure_hyperlink_style(document)
    sources = _sources(findings, section_titles, clauses)
    by_span = {id(s.span): s for s in sources}
    _properties(document, run, sum(1 for s in sources if s.span.verified), len(sources), review_head)

    document.add_heading("Contract Review Memo", level=1)
    date = datetime.now().astimezone().strftime("%d %B %Y")
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in (
        ("Agreement", run_document_name(run)),
        ("Date", date),
        ("Question", run.question),
        ("Guidance", run.guidance.text if run.guidance else "none supplied"),
        ("Sections read", reading_line(run)),
    ):
        row = table.add_row().cells
        row[0].text = label
        row[1].text = xml_safe(value)

    flags = memo_flags(findings, withheld, flagged)
    issues = [f for f, _ in flags]
    why = {f.id: reason for f, reason in flags}
    guidance_text = run.guidance.text if run.guidance else None
    document.add_heading(f"{FLAGS_HEAD} ({len(issues)})", level=2)
    if issues:
        for f in issues:
            document.add_paragraph(xml_safe(flag_title(f) + (f" — {why[f.id]}" if why[f.id] else "")), style="List Number")
    else:
        document.add_paragraph(NO_FLAGS)

    for f in findings:
        document.add_heading(xml_safe(f"{f.topic} · {status_label(f.status, run.guidance is not None, f.status_source)}"), level=2)
        document.add_paragraph(xml_safe(f.shown_conclusion))
        if f.withheld_conclusion:
            also = document.add_paragraph()
            also.add_run("The model also wrote ").bold = True
            also.add_run(xml_safe(f"{WITHHELD_NOTE}: ")).italic = True
            also.add_run(xml_safe(f.withheld_conclusion))
        for span in f.spans:
            source = by_span[id(span)]
            p = document.add_paragraph()
            p.add_run(f"{_evidence_head(f)} ").bold = True
            p.add_run(xml_safe(f"({source.title}) "))
            _internal_link(p, f"src_{source.number}", f"[{source.number}]")
            quote = document.add_paragraph(xml_safe(span.quote))
            quote.paragraph_format.left_indent = Pt(18)
        if unseen and f.id in unseen:
            p = document.add_paragraph()
            p.add_run(f"{UNSEEN_HEAD} ").bold = True
            p.add_run(xml_safe(unseen_note(unseen[f.id])))
        if f.guidance_reference:
            p = document.add_paragraph()
            p.add_run("The model's reference to the guidance. ").bold = True
            p.add_run(xml_safe(guidance_line(f, guidance_text)))
        if model_reading(f):
            p = document.add_paragraph()
            p.add_run("The model's reading. ").bold = True
            p.add_run(xml_safe(model_reading(f)))
        if why.get(f.id):
            p = document.add_paragraph()
            p.add_run("Why flagged. ").bold = True
            p.add_run(xml_safe(why[f.id]))
        for sentence in (related or {}).get(f.id, []):
            p = document.add_paragraph()
            p.add_run(f"{RELATED_HEAD} ").bold = True
            p.add_run(xml_safe(sentence))
        if _suggests(f):
            p = document.add_paragraph()
            p.add_run("Suggested position (the model's). ").bold = True
            p.add_run(xml_safe(f.suggested_position or ""))
        p = document.add_paragraph()
        if f.review is not None:
            p.add_run("Reviewed. ").bold = True
            p.add_run(xml_safe(_review_line(f)))
        else:
            p.add_run("Person. ").bold = True
            p.add_run(NOT_REVIEWED)

    if withheld:
        document.add_heading(f"{WITHHELD_HEAD} ({len(withheld)})", level=2)
        document.add_paragraph(WITHHELD_LEDE).runs[0].italic = True
        for w in withheld:
            p = document.add_paragraph()
            p.add_run(xml_safe(f"{w.topic}. ")).bold = True
            p.add_run("The model wrote, withheld: ").italic = True
            p.add_run(xml_safe(w.conclusion))
            p = document.add_paragraph()
            p.add_run("Why withheld. ").bold = True
            p.add_run(xml_safe(why[w.id]))

    document.add_heading("Sources", level=2)
    sources_table = document.add_table(rows=1, cols=4)
    sources_table.style = "Table Grid"
    header = sources_table.rows[0].cells
    for cell, text in zip(header, ("#", "Section", "Located", "Record"), strict=True):
        cell.text = text
    for source in sources:
        row = sources_table.add_row().cells
        _bookmark(row[0].paragraphs[0], f"src_{source.number}", str(source.number), source.number)
        row[1].text = xml_safe(source.title)
        row[2].text = _located(source.span)
        _external_link(row[3].paragraphs[0], deep_link(run, source.finding), "Open in the workbench")

    footer = document.add_paragraph()
    footer.add_run(
        f"Generated from run {run.id} · review state {review_head[:8] or 'none'} · model {run.model} · prompt {run.prompt_version} ({run.prompt_hash[:12]}). "
        "Every quoted passage was found in the document text before this memo was written; that a passage supports its finding is "
        "the model's reading and the reviewer's decision. "
        "The run id, document hash, model and prompt are also in this file's document properties."
    ).italic = True
    path = out_dir / f"memo-{run.id}.docx"
    document.save(str(path))
    return path, sha256_file(path)
