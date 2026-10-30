"""Read DOCX, PDF and TXT into blocks. Nothing here calls a model.

DOCX is read as Word shows it with every tracked change accepted: inserted text is kept, deleted
and moved-away text is dropped, runs marked hidden are dropped, deleted table rows are skipped and
a deleted paragraph mark joins its paragraph to the next. python-docx's `paragraph.text` does none
of this; it walks only the runs that hang directly off the paragraph, so on negotiated paper it
loses both the old wording and the new. What the reader left out is counted and stored, so the
interface can say which view of the file the model read.
"""

from __future__ import annotations

import codecs
import io
import logging
import re
import zipfile
from dataclasses import dataclass, field
from typing import cast

from docx import Document as DocxDocument
from docx.document import Document as WordDocument
from docx.oxml.ns import qn
from docx.oxml.text.paragraph import CT_P
from docx.styles.style import ParagraphStyle
from docx.text.paragraph import Paragraph
from lxml import etree
from pypdf import PageObject, PdfReader

from .coverage import DocxWalkCounts, PartCoverage, PdfTally, Status, docx_coverage, docx_more_coverage, pdf_coverage, txt_coverage
from .invisible import remove_format_characters, visible_text
from .pdf_layout import page_text
from .sections import (
    HEADING_NUMBER,
    MAX_TOP_STEP,
    PASSAGE_ENDS,
    Block,
    continues_numbering,
    goes_on,
    looks_like_heading,
    page_edge_furniture,
    shouting,
    states,
)
from .word_display import (
    ALTERNATE,
    CHOICE,
    FALLBACK,
    REPLACEMENT,
    Bindings,
    Compatibility,
    Fields,
    fields_balance,
    in_symbol_font,
    namespace_of,
    symbol_font,
    symbol_text,
)

log = logging.getLogger(__name__)

SUPPORTED = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pdf": "application/pdf",
    ".txt": "text/plain",
}
# The extension a document is stored and read again under, from the media type recorded when its bytes were checked;
# never what is left of its name. A name over 255 characters was cut through its ".txt" (QA campaign, 2026-10-08), and
# the stored bytes, the reading again and the evidence pack all took their type from the name.
SUFFIX_OF = {media_type: suffix for suffix, media_type in SUPPORTED.items()}


# Refused before reading begins: a package that declares more than this unpacked. A declaration is not a bound on what
# zipfile inflates (it inflates a chunk, then truncates), so this only spares reading what no real file declares; what
# reading costs is bounded where it runs. Images do not compress, so a real file unpacks to about its own size; the
# largest of 323 real contracts unpacks to 6.4 MB (2026-10-13; 256 MB before). What reading itself may cost is bounded
# where it runs (ingest/bounded.py).
MAX_DOCX_UNPACKED_BYTES = 64 * 1024 * 1024
MAX_PDF_PAGES = 2000


class UnsupportedFile(ValueError):
    pass


@dataclass
class ReadResult:
    blocks: list[Block]
    pages: int | None
    title: str
    # DOCX only. Revisions the file carried (insertions, deletions, moves, including paragraph
    # marks and table rows) and runs marked hidden. Both are left out of the blocks.
    tracked_changes: int = 0
    hidden_runs: int = 0
    # What this reading did with each part of the file (ingest.coverage); anything unread is listed as unread.
    coverage: list[PartCoverage] = field(default_factory=list)
    # Invisible format characters removed from the text (ingest.invisible), counted so the reading can say so.
    format_characters: int = 0
    # Parts a reading accounts for beyond coverage.PARTS (reader v16), reported after them only when the file has any.
    more_coverage: list[PartCoverage] = field(default_factory=list)


PDF_SIGNATURE = b"%PDF-"
PDF_SIGNATURE_WINDOW = 1024  # readers accept a header anywhere in the first kilobyte, so the check does too
ZIP_SIGNATURE = b"PK\x03\x04"
UNREADABLE = {".docx": "the file is not a readable .docx package", ".pdf": "the file is not a readable PDF"}


def content_type(data: bytes) -> str | None:
    """What the bytes are, whatever the name says: ".pdf", ".docx" for a zip package, or None (text, or neither)."""
    if PDF_SIGNATURE in data[:PDF_SIGNATURE_WINDOW]:
        return ".pdf"
    if data.startswith(ZIP_SIGNATURE):
        return ".docx"
    return None


def file_type(filename: str, data: bytes) -> str:
    """The type a file is read as: the name claims it, the bytes must agree. A document is its bytes and the reading of
    them, so the same bytes must only ever be read one way. Before 2026-10-08 the name alone chose the reader: a PDF
    uploaded as .txt became a "document" of its own binary, and the same bytes uploaded later as .pdf were handed that
    reading, because the store knew them by hash (QA campaign, CUAD pair)."""
    suffix = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if suffix not in SUPPORTED:
        raise UnsupportedFile(f"unsupported file type {suffix or '(none)'}; upload .docx, .pdf or .txt")
    if not data:  # an empty file was "not a readable PDF" or "no readable text found", which sent the reader looking at the file's form
        raise UnsupportedFile("the file is empty")
    found = content_type(data)
    if suffix == ".txt" and found is not None:
        raise UnsupportedFile(f"the file is a {'PDF' if found == '.pdf' else 'zip package such as a .docx'}, not plain text; upload it as {found}")
    if suffix != ".txt" and found != suffix:
        raise UnsupportedFile(UNREADABLE[suffix])
    return suffix


def read(filename: str, data: bytes) -> ReadResult:
    """Every file of a supported type ends as a reading or as UnsupportedFile with a reason. The parsers fail on a damaged
    file in many shapes (pypdf: AssertionError, KeyError, TypeError, NotImplementedError and its own errors; python-docx
    and zipfile: zlib.error among others); 112 of 300 mutations of a real PDF escaped as a 500 before 2026-10-08."""
    suffix = file_type(filename, data)
    result = _read_as(suffix, data)
    # The text a person reads is the text that is checked: each reader removes invisible format characters where its
    # text is first read (so headings are found on the visible text), and this pass guarantees it for every block.
    removed = result.format_characters
    for block in result.blocks:
        block.text, count = remove_format_characters(block.text)
        removed += count
    result.title = visible_text(result.title)
    result.format_characters = removed
    result.coverage.append(
        PartCoverage(
            "format_characters",
            Status.EXCLUDED if removed else Status.ABSENT,
            removed,
            "invisible characters (zero-width, soft hyphen, direction controls) removed; the text is shown in stored order"
            if removed
            else "no invisible format characters",
        )
    )
    result.coverage.extend(result.more_coverage)
    return result


def _read_as(suffix: str, data: bytes) -> ReadResult:
    try:
        if suffix == ".docx":
            return read_docx(data)
        if suffix == ".pdf":
            return read_pdf(data)
        return read_txt(data)
    except UnsupportedFile:
        raise
    except Exception as error:
        # pypdf's LimitReachedError is among these: it guards decompression and also flags damaged structure ("Invalid CID
        # width range"), so it is not evidence of size; a file the parsers cannot finish is unreadable, whatever the shape.
        log.warning("reader refused a damaged %s: %s: %s", suffix, type(error).__name__, str(error)[:200])
        raise UnsupportedFile(UNREADABLE.get(suffix, "the file is not readable plain text")) from error


# --------------------------------------------------------------------------- DOCX

_PARAGRAPH = qn("w:p")
_TABLE = qn("w:tbl")
_RUN = qn("w:r")
_TEXT = qn("w:t")
# Content that Word removes when the change is accepted.
_DELETED = frozenset({qn("w:del"), qn("w:moveFrom")})
# Property bags hold no visible text. Skipping them also keeps mark-level revision elements
# (w:pPr/w:rPr/w:ins, w:trPr/w:del) out of the text walk.
_PROPERTIES = frozenset({qn(tag) for tag in ("w:pPr", "w:rPr", "w:sdtPr", "w:sdtEndPr", "w:tblPr", "w:tblGrid", "w:trPr", "w:tcPr")})
_RUN_BREAKS = {qn("w:tab"): "\t", qn("w:ptab"): "\t", qn("w:br"): "\n", qn("w:cr"): "\n", qn("w:noBreakHyphen"): "-"}
_REVISIONS_XPATH = ".//w:ins | .//w:del | .//w:moveFrom | .//w:moveTo"
_DELETED_MARK = f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:del')}"
# Reader v14: a numbered paragraph inside a text box, or with a tracked change to its properties, leaves its list's
# printed numbers to Word (`_Numbering.printed_numbers`).
_TEXT_BOX = qn("w:txbxContent")
_PARAGRAPH_PROPERTIES_CHANGE = f"{qn('w:pPr')}/{qn('w:pPrChange')}"
_MARK_REMOVED = (_DELETED_MARK, f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:moveFrom')}")
_SECTION_BREAK = f"{qn('w:pPr')}/{qn('w:sectPr')}"
_DELETED_ROW = f"{qn('w:trPr')}/{qn('w:del')}"
_SDT = qn("w:sdt")
_SDT_CONTENT = qn("w:sdtContent")
_ROW = qn("w:tr")
_CELL = qn("w:tc")
_CUSTOM_XML = qn("w:customXml")
_VAL = qn("w:val")
_SIZE = qn("w:sz")
_COLOUR = qn("w:color")
_SHADING = qn("w:shd")
_HIGHLIGHT = qn("w:highlight")
_FIELD_CHAR = qn("w:fldChar")
_SYM = qn("w:sym")
_MATH = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_MATH_TEXT = f"{{{_MATH}}}t"
# Reader v16: what a run holds that Word draws as a picture or a box, what it holds that Word never displays, the
# containers a paragraph's runs may sit in, and the direction controls Word holds for w:bdo and w:dir.
_DRAWINGS = frozenset({qn("w:drawing"), qn("w:pict"), qn("w:object")})
_NOT_DISPLAYED = frozenset(
    qn(f"w:{tag}")
    for tag in (
        "rPr",
        "instrText",
        "delText",
        "delInstrText",
        "footnoteReference",
        "endnoteReference",
        "footnoteRef",
        "endnoteRef",
        "commentReference",
        "annotationRef",
        "lastRenderedPageBreak",
        "softHyphen",
        "separator",
        "continuationSeparator",
    )
)
_INLINE_CONTAINERS = frozenset(qn(f"w:{tag}") for tag in ("ins", "moveTo", "hyperlink", "smartTag", "customXml", "fldSimple", "sdtContent"))
_DIRECTIONS = {qn("w:bdo"): {"rtl": "\u202e", "ltr": "\u202d"}, qn("w:dir"): {"rtl": "\u202b", "ltr": "\u202a"}}
# Elements of a table, a row or the body that hold no text of their own.
_STRUCTURE = frozenset(qn(f"w:{tag}") for tag in ("tblPr", "tblGrid", "trPr", "tblPrEx", "tcPr", "sectPr"))
# A title or a heading is a line, not a page: a paragraph longer than this becomes body text under a shortened heading,
# so a 100,000-character paragraph no longer rides into every prompt whole as the section's name (sweep, 2026-10-01).
MAX_TITLE_CHARS = 200


@dataclass
class _Walk:
    """What the walk over a DOCX body met and left out, and what decides what Word displays (reader v16: markup
    compatibility, fields, data-bound controls, symbols, text drawn so it cannot be seen; ingest/word_display.py)."""

    styles: _Styles
    compat: Compatibility
    bindings: Bindings
    hidden_runs: int = 0
    # Runs hidden only through a style (a character style named Hidden, a paragraph style whose runs vanish) or
    # through w:specVanish; counted apart so the coverage report can say which door was closed.
    style_hidden_runs: int = 0
    fields: Fields = field(default_factory=Fields)
    # Runs Word draws so that they cannot be seen: white with nothing shaded beneath, or 1 pt or smaller. Read until
    # v16, so a quote of them verified (DECISIONS 2026-10-30, open item one).
    invisible_runs: int = 0
    page_shaded: bool = False  # the document has a background of its own, on which white text may be seen
    symbols: int = 0
    unknown_symbols: int = 0
    # Text boxes met in the paragraph being read: each is read once, as blocks after that paragraph, never inside its
    # sentence. Until v16 none was read and coverage said nothing.
    boxes: list[etree._Element] = field(default_factory=list)
    text_boxes: int = 0
    drawn: int = 0  # pictures, shapes without text and embedded objects
    controls: int = 0
    tables: int = 0
    unread: dict[str, int] = field(default_factory=dict)  # displayed content the walk does not render, by kind

    def skip(self, element: etree._Element) -> None:
        """An element the walk does not render: counted, by its name, when it holds text Word would display."""
        if any(node.text and node.text.strip() for node in element.iter(_TEXT, _MATH_TEXT)):
            kind = "equations" if namespace_of(element.tag) == _MATH else etree.QName(element).localname
            self.unread[kind] = self.unread.get(kind, 0) + 1

    def symbol(self, mark: etree._Element) -> str:
        code = mark.get(qn("w:char")) or ""
        try:
            shown, known = symbol_text(mark.get(qn("w:font")), int(code, 16))
        except ValueError:
            shown, known = REPLACEMENT, False
        self.symbols += 1
        self.unknown_symbols += not known
        return shown

    def draw(self, drawing: etree._Element) -> None:
        """A drawing, a VML picture or an embedded object in a run: its text boxes are kept to be read after the
        paragraph, once each (a box inside another is read with it); one without text is counted as drawn."""
        boxes = [box for box in drawing.iter(_TEXT_BOX) if not _inside_box(box, drawing)]
        if boxes:
            self.boxes.extend(boxes)
        else:
            self.drawn += 1


_STYLE_DEPTH = 12


class _Styles:
    """The document's styles, read once: for each style id, the style it is based on and whether its run
    properties hide text. Word hides text through styles as readily as through a run's own properties, and a
    reader that resolves only the run's own w:vanish hands style-hidden text to the model (reader v6)."""

    def __init__(self, styles: etree._Element | None) -> None:
        self._based_on: dict[str, str | None] = {}
        self._hidden: dict[str, bool | None] = {}
        self._run_properties: dict[str, etree._Element | None] = {}
        self._paragraph_properties: dict[str, etree._Element | None] = {}
        self.default_hidden = False
        self.defaults: etree._Element | None = None
        if styles is None:
            return
        defaults = styles.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")
        self.defaults = defaults
        self.default_hidden = _vanish_state(defaults) is True
        for style in styles.iterfind(qn("w:style")):
            style_id = style.get(qn("w:styleId"))
            if not style_id:
                continue
            based = style.find(qn("w:basedOn"))
            self._based_on[style_id] = based.get(qn("w:val")) if based is not None else None
            self._hidden[style_id] = _vanish_state(style.find(qn("w:rPr")))
            self._run_properties[style_id] = style.find(qn("w:rPr"))
            self._paragraph_properties[style_id] = style.find(qn("w:pPr"))

    def hides(self, style_id: str | None) -> bool | None:
        """True or False when a style in the chain says; None when none does."""
        seen = 0
        while style_id and seen < _STYLE_DEPTH:
            state = self._hidden.get(style_id)
            if state is not None:
                return state
            style_id = self._based_on.get(style_id)
            seen += 1
        return None

    def stated(self, style_id: str | None, tag: str, paragraph: bool = False) -> etree._Element | None:
        """The first ``tag`` a style's chain states in its run properties (or, with ``paragraph``, its paragraph
        properties): w:color, w:sz, w:shd, w:highlight (reader v16)."""
        bags = self._paragraph_properties if paragraph else self._run_properties
        seen = 0
        while style_id and seen < _STYLE_DEPTH:
            bag = bags.get(style_id)
            found = bag.find(tag) if bag is not None else None
            if found is not None:
                return found
            style_id = self._based_on.get(style_id)
            seen += 1
        return None


def _vanish_state(properties: etree._Element | None) -> bool | None:
    """What a run-properties bag says about hiding: w:vanish (on unless its value is off), or w:specVanish; None when silent."""
    if properties is None:
        return None
    vanish = properties.find(qn("w:vanish"))
    if vanish is not None:
        return vanish.get(qn("w:val")) not in ("0", "false")
    if properties.find(qn("w:specVanish")) is not None:
        return True
    return None


def _style_id(properties: etree._Element | None, tag: str) -> str | None:
    style = properties.find(qn(tag)) if properties is not None else None
    return style.get(qn("w:val")) if style is not None else None


def _paragraph_hides(paragraph: etree._Element, styles: _Styles) -> bool:
    """Whether the paragraph's style chain hides its runs. The paragraph mark's own run properties (w:pPr/w:rPr)
    speak for the mark alone, not for the runs, and are not consulted."""
    return styles.hides(_style_id(paragraph.find(qn("w:pPr")), "w:pStyle")) is True


def _hidden_by(run: etree._Element, paragraph_hides: bool, styles: _Styles) -> str | None:
    """How this run is hidden, if it is: "run" by its own properties, "style" by a style or the paragraph's; None when shown.
    Precedence is Word's: the run's own w:vanish decides either way; then its character style; then the paragraph's style;
    then the document default."""
    properties = run.find(qn("w:rPr"))
    own = _vanish_state(properties)
    if own is not None:
        return "run" if own else None
    by_character_style = styles.hides(_style_id(properties, "w:rStyle"))
    if by_character_style is not None:
        return "style" if by_character_style else None
    if paragraph_hides or styles.default_hidden:
        return "style"
    return None


@dataclass(frozen=True)
class _Look:
    """What a paragraph decides for the runs in it: whether its style hides them, its style (for their colour and size),
    and whether something may be shaded beneath them (its own shading or its style's, a table cell, a text box)."""

    hides: bool = False
    style_id: str | None = None
    shaded: bool = False


def _look(paragraph: etree._Element, styles: _Styles) -> _Look:
    properties = paragraph.find(qn("w:pPr"))
    style_id = _style_id(properties, "w:pStyle")
    own = properties.find(_SHADING) if properties is not None else None
    shading = own if own is not None else styles.stated(style_id, _SHADING, paragraph=True)
    beneath = next((a for a in paragraph.iterancestors(_CELL, _TEXT_BOX)), None) is not None
    return _Look(_paragraph_hides(paragraph, styles), style_id, beneath or _shades(shading))


# Text of this size or smaller is a line, not a word: 2 half-points (1 pt). The stress test hid "thirty (30)" at half a point.
MAX_INVISIBLE_HALF_POINTS = 2


def _shades(mark: etree._Element | None) -> bool:
    """Whether a w:shd or w:highlight puts a colour other than white beneath the text."""
    if mark is None:
        return False
    if mark.tag == _HIGHLIGHT:
        return (mark.get(_VAL) or "none") != "none"
    return (mark.get(qn("w:fill")) or "auto").upper() not in ("AUTO", "FFFFFF") or (mark.get(_VAL) or "clear") not in ("clear", "nil")


def _invisible(run: etree._Element, look: _Look, walk: _Walk) -> str | None:
    """How Word draws this run so that it cannot be seen, if it does: "tiny" at 1 pt or less, "white" in white with
    nothing shaded beneath (its own shading or highlight, its paragraph's, a table cell's, a text box's or the page's,
    each of which is taken to show it). Each property is Word's: the run's own, then its character style's chain, then its
    paragraph style's, then the document's defaults; a theme colour is read by the value Word stores beside it."""
    properties = run.find(qn("w:rPr"))
    character_style = _style_id(properties, "w:rStyle")

    def stated(tag: str) -> etree._Element | None:
        for found in (
            properties.find(tag) if properties is not None else None,
            walk.styles.stated(character_style, tag),
            walk.styles.stated(look.style_id, tag),
            walk.styles.defaults.find(tag) if walk.styles.defaults is not None else None,
        ):
            if found is not None:
                return found
        return None

    size = stated(_SIZE)
    if size is not None and _int_or(size.get(_VAL), MAX_INVISIBLE_HALF_POINTS + 1) <= MAX_INVISIBLE_HALF_POINTS:
        return "tiny"
    colour = stated(_COLOUR)
    if colour is None or (colour.get(_VAL) or "").upper() != "FFFFFF" or look.shaded or walk.page_shaded:
        return None
    if _shades(stated(_SHADING)) or _shades(stated(_HIGHLIGHT)):
        return None
    return "white"


def _run_text(run: etree._Element, walk: _Walk, shown: bool = True) -> str:
    """Displayed text of one run: w:t (in a symbol font, the glyphs that font draws), tabs and breaks, a w:sym's glyph,
    a form field's box or entry at its end, and nothing while a field's instruction is open (reader v16). A run that is
    not ``shown`` still moves the field state, as Word's hidden field marks do. Drawings hand their text boxes to the
    walk; deleted text and instructions carry none."""
    properties = run.find(qn("w:rPr"))
    fonts = properties.find(qn("w:rFonts")) if properties is not None else None
    table = symbol_font(fonts.get(qn("w:ascii")) or fonts.get(qn("w:hAnsi"))) if fonts is not None else None
    parts: list[str] = []
    for child in walk.compat.children(run):
        tag = child.tag
        if tag == _FIELD_CHAR:
            drawn = walk.fields.mark(child)
            if shown:
                parts.append(drawn)
        elif not shown or not walk.fields.showing or tag in _NOT_DISPLAYED:
            continue
        elif tag == _TEXT:
            text = child.text or ""
            if table is not None:
                text, unknown = in_symbol_font(text, table)
                walk.symbols += 1
                walk.unknown_symbols += unknown
            parts.append(text)
        elif tag in _RUN_BREAKS:
            parts.append(_RUN_BREAKS[tag])
        elif tag == _SYM:
            parts.append(walk.symbol(child))
        elif tag in _DRAWINGS:
            walk.draw(child)
        else:
            walk.skip(child)
    return "".join(parts)


def accepted_text(element: etree._Element, walk: _Walk, look: _Look | None = None) -> str:
    """Text of a paragraph, as Word shows it with all changes accepted.

    Runs are collected in document order through every container Word renders (w:ins, w:moveTo, w:hyperlink,
    w:fldSimple, w:smartTag, w:sdtContent, w:customXml, w:bdo, w:dir), with markup compatibility applied (one branch of
    an mc:AlternateContent, an ignorable element Word does not know left out); deleted and moved-away containers,
    hidden runs and runs drawn so they cannot be seen are left out. A direction override (w:bdo) or embedding (w:dir) is
    written as the control characters Word holds for it, which the reading then removes and counts, as it does in a
    text file. Any other element is not read, and counted when it holds text (reader v16: until then every element
    with children was read as a container, both branches of an mc:AlternateContent among them).
    """
    if look is None:
        look = _look(element, walk.styles) if element.tag == _PARAGRAPH else _Look()
    parts: list[str] = []
    for child in walk.compat.children(element):
        tag = child.tag
        if tag in _DELETED or tag in _PROPERTIES:
            continue
        if tag == _RUN:
            hidden = _hidden_by(child, look.hides, walk.styles)
            if hidden is None and _has_text(child) and (invisible := _invisible(child, look, walk)):
                walk.invisible_runs += 1
                _run_text(child, walk, shown=False)
                parts.append(" " if invisible == "white" else "")  # white text still takes its width on the line
                continue
            if hidden == "run":
                walk.hidden_runs += 1
            elif hidden == "style":
                walk.style_hidden_runs += 1
            parts.append(_run_text(child, walk, shown=hidden is None))
        elif tag == _SDT:
            parts.append(_control_text(child, walk, look))
        elif tag in _DIRECTIONS:
            opening = _DIRECTIONS[tag].get(child.get(_VAL) or "", "")
            inner = accepted_text(child, walk, look)
            parts.append(f"{opening}{inner}\u202c" if opening and inner else inner)
        elif tag in _INLINE_CONTAINERS:
            parts.append(accepted_text(child, walk, look))
        else:
            walk.skip(child)
    return "".join(parts)


def _control_text(control: etree._Element, walk: _Walk, look: _Look) -> str:
    """An inline content control's displayed text: its content, or what its data binding displays (word_display.Bindings)."""
    walk.controls += 1
    content = control.find(_SDT_CONTENT)
    cached = accepted_text(content, walk, look) if content is not None else ""
    bound = walk.bindings.shown(control.find(qn("w:sdtPr")), cached)
    return cached if bound is None else bound


def _has_text(run: etree._Element) -> bool:
    return any(node.text and node.text.strip() for node in run.iter(_TEXT))


def _inside_box(box: etree._Element, drawing: etree._Element) -> bool:
    """Whether a text box sits inside another text box below ``drawing``: it is read with that one."""
    parent = box.getparent()
    while parent is not None and parent is not drawing:
        if parent.tag == _TEXT_BOX:
            return True
        parent = parent.getparent()
    return False


def _plain_text(element: etree._Element) -> str:
    return "".join(node.text or "" for node in element.iter(_TEXT))


def _set_apart(paragraph: Paragraph) -> bool:
    """Whether a paragraph's text is set apart as a title: every run with text bold (its own setting, or its paragraph
    style's when it has none), or every letter a capital."""
    letters = [c for c in paragraph.text if c.isalpha()]
    if letters and all(c.isupper() for c in letters):
        return True
    runs = [run for run in paragraph.runs if run.text.strip()]
    style_bold = bool(paragraph.style is not None and paragraph.style.font.bold)
    return bool(runs) and all(run.bold or (run.bold is None and style_bold) for run in runs)


def _heading_level(paragraph: Paragraph) -> int:
    """1-based outline level from the style name (Heading 2, SchLevel1Heading, Title), the paragraph's own
    w:outlineLvl, or the outline level its style chain declares; 0 for body text."""
    style = paragraph.style
    name = (style.name if style is not None else "") or ""
    if name.lower() == "title":
        return 1
    if re.search(r"heading", name, re.IGNORECASE):
        digit = re.search(r"(\d)", name)
        if digit:
            return int(digit.group(1))
    own = _outline_level(paragraph._p.pPr)
    if own is not None:
        return own + 1
    depth = 0
    while style is not None and depth < 8:
        declared = _outline_level(style.element.find(qn("w:pPr")))
        if declared is not None:
            return declared + 1
        style = style.base_style
        depth += 1
    return 0


# The words a heading may print before its number and still be a clause number (sections.HEADING_NUMBER reads the same
# three): "Article 4" is clause 4; "Schedule 1" is not a clause, and its number is not one.
_CLAUSE_WORDS = frozenset({"section", "clause", "article"})


class _Numbering:
    """The document's numbering definitions (w:numbering), read once: for each numbering instance (w:num), each level's
    format and text as Word prints it, the instance's own level overrides applied and a list that is a link to a
    numbering style followed one step. A reading of the definitions is never a reason to refuse the file: a package
    without them, or with ones it cannot read, numbers nothing."""

    def __init__(self, document: WordDocument) -> None:
        self.levels: dict[str, dict[str, tuple[str, str]]] = {}
        # Reader v14: each level's start value, and the levels whose counting this reader does not reproduce
        # (w:lvlRestart, w:isLgl), per instance; and which list (w:abstractNum) each instance counts in.
        self.starts: dict[str, dict[str, int]] = {}
        self.unsupported: dict[str, set[str]] = {}
        self.list_of: dict[str, str] = {}
        # What each paragraph style's chain states (`_style_numbering`), worked out once per style for this reading.
        self.by_style: dict[str | None, tuple[str | None, str | None]] = {}
        try:
            root = document.part.numbering_part.element
            styles = document.styles.element
        except Exception:  # noqa: BLE001 - no numbering part, or one python-docx cannot read: nothing is numbered
            return
        abstract: dict[str, dict[str, tuple[str, str]]] = {}
        abstract_starts: dict[str, dict[str, int]] = {}
        abstract_unsupported: dict[str, set[str]] = {}
        linked: dict[str, str] = {}
        for definition in root.iterfind(qn("w:abstractNum")):
            key = definition.get(qn("w:abstractNumId")) or ""
            abstract[key] = {(level.get(qn("w:ilvl")) or "0"): _level_of(level) for level in definition.iterfind(qn("w:lvl"))}
            abstract_starts[key] = {(level.get(qn("w:ilvl")) or "0"): _start_of(level) for level in definition.iterfind(qn("w:lvl"))}
            abstract_unsupported[key] = {(level.get(qn("w:ilvl")) or "0") for level in definition.iterfind(qn("w:lvl")) if _counting_unsupported(level)}
            link = definition.find(qn("w:numStyleLink"))
            if link is not None and link.get(qn("w:val")):
                linked[key] = link.get(qn("w:val")) or ""
        instances: dict[str, str] = {}
        overrides: dict[str, list[etree._Element]] = {}
        for instance in root.iterfind(qn("w:num")):  # one pass: each instance's list and its own level overrides
            num_id = instance.get(qn("w:numId")) or ""
            ref = instance.find(qn("w:abstractNumId"))
            instances[num_id] = (ref.get(qn("w:val")) or "") if ref is not None else ""
            overrides[num_id] = instance.findall(qn("w:lvlOverride"))
        style_numbers = {
            style.get(qn("w:styleId")) or "": numbering.get(qn("w:val")) or ""
            for style in styles.iterfind(qn("w:style"))
            for numbering in style.iterfind(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
        }
        for num_id, abstract_id in instances.items():
            list_id = abstract_id
            levels = abstract.get(abstract_id, {})
            if not levels and abstract_id in linked:  # a list that is a link to a numbering style: that style's list
                list_id = instances.get(style_numbers.get(linked[abstract_id], ""), "")
                levels = abstract.get(list_id, {})
            levels = dict(levels)
            starts = dict(abstract_starts.get(list_id, {}))
            unsupported = set(abstract_unsupported.get(list_id, set()))
            for override in overrides[num_id]:
                at = override.get(qn("w:ilvl")) or "0"
                level = override.find(qn("w:lvl"))
                if level is not None:
                    levels[at] = _level_of(level)
                    starts[at] = _start_of(level)
                    if _counting_unsupported(level):
                        unsupported.add(at)
                restart = override.find(qn("w:startOverride"))
                if restart is not None:
                    starts[at] = _int_or(restart.get(qn("w:val")), starts.get(at, 1))
            self.levels[num_id] = levels
            self.starts[num_id] = starts
            self.unsupported[num_id] = unsupported
            self.list_of[num_id] = list_id

    def shows_clause_number(self, num_id: str, level: str) -> bool:
        """Whether Word prints a clause number for this level of this instance: the level exists, its format is neither
        "none" nor a bullet, its text prints something, and neither it nor its list's first level prints words other
        than Section, Clause or Article ("Schedule %1" numbers schedules, not clauses, and so do its sub-levels)."""
        levels = self.levels.get(num_id)
        if not levels or level not in levels:
            return False
        number_format, text = levels[level]
        if number_format in ("none", "bullet") or not text.strip():
            return False
        for printed in (text, levels.get("0", ("", ""))[1]):
            words = re.findall(r"[^\W\d_]+", re.sub(r"%\d", "", printed))
            if any(word.casefold() not in _CLAUSE_WORDS for word in words):
                return False
        return True

    def printed_numbers(self, document: WordDocument, body: etree._Element) -> dict[etree._Element, str | None]:
        """The clause number Word prints for each paragraph it numbers (reader v14), from the list's own counters: every
        numbered paragraph of the body in document order counts, tables included; a level starts at its start value (an
        instance's override first), counts up, and restarts whenever a level above it counts. The number is the level's
        text with each %N replaced by level N's count in its format ("%1.%2" is "4.2", "(%2)" is "(b)", "%1." with upper
        roman is "IV"). Where the reader cannot say exactly what Word prints, the paragraph gets None, never a guess: two
        instances of one list in use (whether the second continues the first or restarts is Word's own rule), a level
        whose counting it does not reproduce, a format other than decimal, roman or letters, a count referred to before
        it has started, a numbered paragraph in a text box or with a tracked change to its paragraph properties."""
        events: list[tuple[etree._Element, str, str]] = []
        doubtful: set[str] = set()
        for element in body.iter(_PARAGRAPH):
            found = _numbering_of(Paragraph(cast(CT_P, element), document), self.by_style)
            if found is None or found[0] not in self.levels:
                continue
            counted = _counted_by_word(element)
            if counted is False:
                continue
            events.append((element, *found))
            if counted is None or next(element.iterancestors(_TEXT_BOX), None) is not None or element.find(_PARAGRAPH_PROPERTIES_CHANGE) is not None:
                doubtful.add(found[0])
        in_use: dict[str, set[str]] = {}
        for _, num_id, _ in events:
            in_use.setdefault(self.list_of.get(num_id, ""), set()).add(num_id)
        counts: dict[str, dict[int, int]] = {}
        printed: dict[etree._Element, str | None] = {}
        for element, num_id, level in events:
            if not level.isdigit():
                printed[element] = None
                continue
            at = int(level)
            count = counts.setdefault(num_id, {})
            count[at] = count.get(at, self.starts[num_id].get(level, 1) - 1) + 1
            for deeper in [k for k in count if k > at]:
                del count[deeper]
            exact = num_id not in doubtful and len(in_use.get(self.list_of.get(num_id, ""), set())) == 1 and level not in self.unsupported[num_id]
            printed[element] = self._cited(num_id, at, count) if exact and self.shows_clause_number(num_id, level) else None
        return printed

    def _cited(self, num_id: str, at: int, count: dict[int, int]) -> str | None:
        """The number a clause is cited by: what Word prints for it, and for a lettered or roman sub-item that prints only
        its own counter ("(a)"), the number of the clause it sits in before it, as the document's own cross-references
        cite it ("Section 8.1(a)" in Common Paper's agreement; reader v14 showed "§(a)" beside "§8.1"). Built only from
        what Word prints at each level; a level whose text names its parents ("%1.%2") already carries them."""
        own = self._render(num_id, str(at), count)
        if own is None or at == 0 or not _SUB_ITEM.fullmatch(own):
            return own
        if any(int(level) - 1 != at for level in re.findall(r"%([1-9])", self.levels[num_id][str(at)][1])):
            return own
        if (at - 1) not in count or not self.shows_clause_number(num_id, str(at - 1)) or str(at - 1) in self.unsupported[num_id]:
            return own
        parent = self._cited(num_id, at - 1, count)
        return f"{parent}{own}" if parent else own

    def _render(self, num_id: str, level: str, count: dict[int, int]) -> str | None:
        levels = self.levels[num_id]
        text = levels[level][1]
        parts: list[str] = []
        position = 0
        for marker in re.finditer(r"%([1-9])", text):
            referred = int(marker.group(1)) - 1
            value = count.get(referred)
            shape = levels.get(str(referred), ("", ""))[0]
            if value is None or str(referred) in self.unsupported[num_id]:
                return None
            shown = _format_count(value, shape)
            if shown is None:
                return None
            parts += [text[position : marker.start()], shown]
            position = marker.end()
        parts.append(text[position:])
        return _clause_number("".join(parts))


def _counted_by_word(paragraph: etree._Element) -> bool | None:
    """Whether Word counts this numbered paragraph in its list, as Word 16 itself reports (measured on built cases,
    2026-10-14): not when its mark is a tracked deletion or move (the paragraph runs into the next one), nor when it
    holds a section break and no character at all, not even an empty run's; it does count an empty paragraph without a
    section break, a hidden one, and one whose text is all deleted. Common Paper's cover page ends in an empty numbered
    Heading 1 holding the section break: Word prints "1. Service" after it, and counting it said "2". None for a
    section break beside content other than text, tabs and breaks, whose counting was not measured."""
    if any(paragraph.find(path) is not None for path in _MARK_REMOVED):
        return False
    if paragraph.find(_SECTION_BREAK) is None:
        return True
    contents = [child for run in paragraph.iter(_RUN) for child in run if child.tag != qn("w:rPr")]
    if not contents:
        return False
    if all(child.tag in (qn("w:tab"), qn("w:br")) or (child.tag == _TEXT and child.text) for child in contents):
        return True
    return None


def _level_of(level: etree._Element) -> tuple[str, str]:
    """(numFmt, lvlText) of a w:lvl: Word's default format is decimal; a level with no text prints nothing."""
    number_format = level.find(qn("w:numFmt"))
    text = level.find(qn("w:lvlText"))
    return (
        (number_format.get(qn("w:val")) if number_format is not None else None) or "decimal",
        (text.get(qn("w:val")) if text is not None else None) or "",
    )


def _start_of(level: etree._Element) -> int:
    """A level's start value (w:start); Word's default is 1."""
    start = level.find(qn("w:start"))
    return _int_or(start.get(qn("w:val")) if start is not None else None, 1)


def _int_or(value: str | None, default: int) -> int:
    return int(value) if value is not None and value.strip().isdigit() else default


def _counting_unsupported(level: etree._Element) -> bool:
    """A level whose counting reader v14 does not reproduce: one that restarts on its own rule (w:lvlRestart) or prints
    every level in legal decimal form (w:isLgl)."""
    return level.find(qn("w:lvlRestart")) is not None or level.find(qn("w:isLgl")) is not None


_ROMAN = ((1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"), (50, "L"), (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I"))


def _format_count(value: int, shape: str) -> str | None:
    """A count as Word prints it in the level's format, for the formats reader v14 reproduces; None for any other."""
    if value < 1 and shape != "decimal":
        return None
    if shape == "decimal":
        return str(value)
    if shape == "decimalZero":
        return f"{value:02d}"
    if shape in ("upperRoman", "lowerRoman"):
        roman = ""
        for size, letters in _ROMAN:
            while value >= size:
                roman, value = roman + letters, value - size
        return roman if shape == "upperRoman" else roman.lower()
    if shape in ("upperLetter", "lowerLetter"):
        letter = chr(ord("a") + (value - 1) % 26) * ((value - 1) // 26 + 1)  # Word's 27th is "aa", its 28th "bb"
        return letter.upper() if shape == "upperLetter" else letter
    return None


# A sub-item's printed label: "(a)", "(iv)", "(2)".
_SUB_ITEM = re.compile(r"\([0-9A-Za-z]{1,6}\)")


def _clause_number(printed: str) -> str | None:
    """The number a section is cited by, from what Word prints: the clause word and a closing full stop left off
    ("Article 4" and "4." are §4), a lone closing parenthesis too ("4)"); None when nothing number-like is left."""
    number = re.sub(r"^(?:section|clause|article)\s+", "", printed.strip(), flags=re.IGNORECASE).rstrip(".")
    if number.endswith(")") and "(" not in number:
        number = number[:-1]
    number = number.strip()
    return number if number and " " not in number and any(c.isalnum() for c in number) else None


def _numbering_of(paragraph: Paragraph, by_style: dict[str | None, tuple[str | None, str | None]]) -> tuple[str, str] | None:
    """(numbering instance, level) Word applies to the paragraph: its own w:numPr, completed from the first style of its
    chain that has one (a w:numPr naming only a level takes its instance from the style); None when it has none or names
    instance 0, Word's "no numbering". ``by_style`` keeps what each paragraph style's chain states, keyed by the style the
    paragraph names (None for the default style), so a document of many headings walks each chain once."""
    own_id, own_level = _numbering_fields(paragraph._p.pPr.find(qn("w:numPr")) if paragraph._p.pPr is not None else None)
    style_id = paragraph._p.style
    if style_id not in by_style:
        by_style[style_id] = _style_numbering(paragraph.style)
    style_num_id, style_level = by_style[style_id]
    num_id = own_id if own_id is not None else style_num_id
    level = own_level if own_level is not None else style_level
    if not num_id or num_id == "0":
        return None
    return num_id, level or "0"


def _numbering_fields(numbering: etree._Element | None) -> tuple[str | None, str | None]:
    """(numId, ilvl) a w:numPr states, each None where it states none."""
    if numbering is None:
        return None, None
    found_id, found_level = numbering.find(qn("w:numId")), numbering.find(qn("w:ilvl"))
    num_id = found_id.get(qn("w:val")) if found_id is not None else None
    level = found_level.get(qn("w:val")) if found_level is not None else None
    return (num_id.strip() if num_id is not None else None), (level.strip() if level is not None else None)


def _style_numbering(style: ParagraphStyle | None) -> tuple[str | None, str | None]:
    """The first numId and the first ilvl stated along a paragraph style's chain: the style, then its bases."""
    num_id: str | None = None
    level: str | None = None
    depth = 0
    while style is not None and depth < _STYLE_DEPTH:
        properties = style.element.find(qn("w:pPr"))
        found_id, found_level = _numbering_fields(properties.find(qn("w:numPr")) if properties is not None else None)
        num_id = num_id if num_id is not None else found_id
        level = level if level is not None else found_level
        style = style.base_style
        depth += 1
    return num_id, level


def _uses_heading_styles(document: WordDocument, body: etree._Element) -> bool:
    """Whether any paragraph is a heading by its style or its own outline level: read once per paragraph style in use, by
    one paragraph of each. Asked of every paragraph, a body of six million empty paragraphs (a bounds fixture) took more
    than five minutes; the styles in use are a handful."""
    if body.find(f".//{_OUTLINE_LEVEL}") is not None:
        return True
    by_style: dict[str, etree._Element] = {}
    for style in body.iterfind(f".//{_PARAGRAPH_STYLE}"):
        properties = style.getparent()  # w:pStyle → w:pPr → w:p
        paragraph = properties.getparent() if properties is not None else None
        if paragraph is not None:
            by_style.setdefault(style.get(qn("w:val")) or "", paragraph)
    unstyled = next((p for p in body.iterfind(f".//{_PARAGRAPH}") if p.find(_PARAGRAPH_STYLE_PATH) is None), None)
    examples = [*by_style.values(), *([unstyled] if unstyled is not None else [])]
    return any(element.tag == _PARAGRAPH and _heading_level(Paragraph(cast(CT_P, element), document)) for element in examples)


_OUTLINE_LEVEL = qn("w:outlineLvl")
_PARAGRAPH_STYLE = qn("w:pStyle")
_PARAGRAPH_STYLE_PATH = f"{qn('w:pPr')}/{qn('w:pStyle')}"


def _outline_level(properties: etree._Element | None) -> int | None:
    """w:outlineLvl as Word stores it: 0–8 are heading levels, 9 is body text."""
    outline = properties.find(qn("w:outlineLvl")) if properties is not None else None
    value = outline.get(qn("w:val")) if outline is not None else None
    if value is None:
        return None
    try:
        level = int(value)
    except ValueError:
        return None
    return level if 0 <= level <= 8 else None


def _docx_pages(data: bytes) -> int | None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            app_xml = archive.read("docProps/app.xml").decode("utf-8", errors="replace")
    except (KeyError, zipfile.BadZipFile):
        return None
    match = re.search(r"<Pages>(\d{1,9})</Pages>", app_xml)
    return int(match.group(1)) if match else None


class TooLargeToRead(UnsupportedFile):
    """Well-formed, but it would cost more memory or time than any contract needs; refused before parsing."""


def _unpacked_size(data: bytes) -> int | None:
    """The declared uncompressed size of a zip package; zipfile never reads a member past its declaration."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile:
        return None


def read_docx(data: bytes) -> ReadResult:
    unpacked = _unpacked_size(data)
    if unpacked is not None and unpacked > MAX_DOCX_UNPACKED_BYTES:
        raise TooLargeToRead(f"the .docx unpacks to {unpacked // (1024 * 1024)} MB; the limit is {MAX_DOCX_UNPACKED_BYTES // (1024 * 1024)} MB")
    # A valid zip with a broken or hostile package inside fails in many shapes; `read` refuses every one of them.
    document = DocxDocument(io.BytesIO(data))
    body = document.element.body
    try:
        styles_element: etree._Element | None = document.styles.element
    except (KeyError, AttributeError):  # a package without a styles part: nothing hides through a style
        styles_element = None
    compat = Compatibility(document.element)
    walk = _Walk(styles=_Styles(styles_element), compat=compat, bindings=Bindings(_bindable_parts(data)))
    walk.fields.enabled = fields_balance(_field_marks(body, compat))
    background = document.element.find(qn("w:background"))
    walk.page_shaded = background is not None and (background.get(qn("w:color")) or "FFFFFF").upper() not in ("FFFFFF", "AUTO")
    numbering = _Numbering(document)
    printed = numbering.printed_numbers(document, body)
    # A DOCX whose headings carry no heading style (manual numbers, bold paragraphs, list numbering) had no heading at all:
    # one block of text cut into parts with no clause numbers, so no citation could name a clause (retrieval study,
    # 2026-10-14: 0 of 293 headings). Where no paragraph has a heading style or outline level, a paragraph is read as a
    # plain text line is (sections.looks_like_heading), with the number Word prints for it when it has one.
    styled = _uses_heading_styles(document, body)
    last_number = 0  # the top-level number of the last such heading: a number that does not continue it is not a clause
    blocks: list[Block] = []
    title = ""
    # Text of paragraphs whose mark is deleted: once accepted, they run into the next paragraph.
    carry = ""

    removed = 0

    def add(text: str, level: int = 0, word_number: str | None = None) -> None:
        nonlocal title, removed
        text, count = remove_format_characters(text)
        removed += count
        if not text:
            return
        if level:
            # A heading carries the number Word prints for it, or none; no number is computed from levels (reader v14).
            blocks.append(Block("heading", text, level, implicit_number=False, word_number=word_number))
            return
        if not title and not blocks:
            title = shorten_title(text)
        blocks.append(Block("text", text))

    def wrapped(container: etree._Element) -> list[etree._Element]:
        """The children of a table, a row, a cell or the body as Word renders them: markup compatibility applied, and a
        content control (its content) or a w:customXml around them entered as the children it wraps (reader v16: a
        block w:customXml, and a control or a w:customXml around a row's cells, were skipped while coverage said read)."""
        out: list[etree._Element] = []
        for child in compat.children(container):
            if child.tag == _SDT:
                walk.controls += 1
                content = child.find(_SDT_CONTENT)
                if content is not None:
                    out.extend(wrapped(content))
            elif child.tag == _CUSTOM_XML:
                out.extend(wrapped(child))
            else:
                out.append(child)
        return out

    def cell_text(cell: etree._Element) -> str:
        """A cell's paragraphs, the rows of any table nested in it, and the content of any control in it."""
        parts: list[str] = []
        for child in wrapped(cell):
            if child.tag == _PARAGRAPH:
                parts.append(accepted_text(child, walk).strip())
            elif child.tag == _TABLE:
                parts.extend(table_rows(child))
            elif child.tag not in _STRUCTURE:
                walk.skip(child)
        return " ".join(part for part in parts if part)

    def row_slots(row: etree._Element) -> list[str]:
        """A row's text, one slot per grid column (reader v16): w:gridBefore's columns empty, each cell's text in the
        slot of its first column and empty slots for the rest of its span, a vertical or horizontal merge's
        continuation empty, as Word shows it; an empty cell empty and a value repeated across cells repeated; trailing
        empty slots dropped. Until v16 a row kept only cells whose text was new in it, so every value after an empty
        cell or a repeated one moved into the column before ("Licence fee | $100,000 | $150,000" under Year 1 | Year 2
        | Year 3, the Year 2 fee read as $150,000), a dedup that python-docx's row.cells had needed and an XML walk did
        not (stress test, 2026-10-30)."""
        properties = row.find(qn("w:trPr"))
        before = properties.find(qn("w:gridBefore")) if properties is not None else None
        slots = [""] * _int_or(before.get(_VAL) if before is not None else None, 0)
        for cell in wrapped(row):
            if cell.tag != _CELL:
                if cell.tag not in _STRUCTURE:
                    walk.skip(cell)
                continue
            cell_properties = cell.find(qn("w:tcPr"))
            span = cell_properties.find(qn("w:gridSpan")) if cell_properties is not None else None
            width = max(1, _int_or(span.get(_VAL) if span is not None else None, 1))
            merged = [cell_properties.find(qn(f"w:{tag}")) if cell_properties is not None else None for tag in ("vMerge", "hMerge")]
            continued = any(m is not None and (m.get(_VAL) or "continue") != "restart" for m in merged)
            text = "" if continued else cell_text(cell)
            slots.extend([text] + [""] * (width - 1))
        while slots and not slots[-1]:
            slots.pop()
        return slots

    def table_rows(table: etree._Element) -> list[str]:
        """Each row as its grid's slots joined with ' | ' (`row_slots`), a deleted row not at all, a row inside a
        content control like any other. Nested tables were dropped before 2026-10-01; coverage said they were read."""
        walk.tables += 1
        rows: list[str] = []
        for child in wrapped(table):
            if child.tag != _ROW:
                if child.tag not in _STRUCTURE:
                    walk.skip(child)
                continue
            if child.find(_DELETED_ROW) is not None:
                continue
            slots = row_slots(child)
            if any(slots):
                rows.append(" | ".join(slots))
        return rows

    def read_boxes() -> None:
        """The text boxes met in the paragraph or table just read, each read once as blocks of its own, with fields of
        its own (a text box is a story of its own in Word)."""
        pending, walk.boxes = walk.boxes, []
        for box in pending:
            walk.text_boxes += 1
            outer, walk.fields.open = walk.fields.open, []
            walk_container(box)
            walk.fields.open = outer

    def _unstyled_heading_level(text: str, number: str | None, after_text: bool) -> int:
        """The level of a paragraph read as a plain heading line in a document with no heading style, or 0: a numbered
        title whose number continues the clause numbering, Word's list number before a title, or a short all-caps line
        after the document's first line (that one is its title)."""
        nonlocal last_number
        line = f"{number} {text}" if number and number[:1].isdigit() else text
        if not looks_like_heading(line):
            return 0
        numbered = HEADING_NUMBER.match(line)
        if numbered is None:
            return 1 if after_text else 0
        if not continues_numbering(numbered.group(1), last_number):
            return 0
        last_number = int(numbered.group(1).split(".")[0])
        return 1 if "." not in numbered.group(1) else 2

    def _styled_top(text: str, number: str | None) -> int | None:
        """The top-level clause number a styled heading states ("10" for "10. Liability" or Word's "10."), or None."""
        line = f"{number} {text}" if number and number[:1].isdigit() else text
        numbered = HEADING_NUMBER.match(line)
        top = numbered.group(1).split(".", 1)[0] if numbered else ""
        return int(top) if top.isdigit() else None

    def _typed_heading_level(paragraph: Paragraph, text: str, number: str | None) -> int:
        """In a document whose headings carry heading styles, a clause heading its author typed instead, or 0. The
        first-visitor acceptance test (2026-10-14) uploaded a services agreement whose "11. Termination" was a bold
        Normal paragraph between the styled "10 Liability" and "13 Notices": read as body, the termination clause sat
        inside §10, and a quote from 11.2(b)(ii) was "found at §10 · Liability" on the card. A typed heading is a short
        numbered title set apart (all bold, or in capitals) whose whole number is the next clause number after the last
        heading, at most one clause written inline between. A numbered list item fails at least one of those."""
        nonlocal last_number
        line = f"{number} {text}" if number and number[:1].isdigit() else text
        numbered = HEADING_NUMBER.match(line)
        if numbered is None or not numbered.group(1).isdigit() or not looks_like_heading(line) or not _set_apart(paragraph):
            return 0
        top = int(numbered.group(1))
        if not 0 < top - last_number <= MAX_TOP_STEP:
            return 0
        last_number = top
        return 1

    def walk_container(container: etree._Element) -> None:
        """Paragraphs and tables in document order, entering block-level content controls (w:sdt), which wrap a
        heading, a clause or a whole table and were skipped before 2026-10-01 while coverage called them read, and
        block-level w:customXml (reader v16). A control bound to a part of the package is read as the value Word
        displays for it (word_display.Bindings). The text boxes a paragraph or a table holds are read after it."""
        nonlocal carry, last_number
        for child in compat.children(container):
            if child.tag == _SDT:
                walk.controls += 1
                content = child.find(_SDT_CONTENT)
                bound = walk.bindings.shown(child.find(qn("w:sdtPr")), _plain_text(content) if content is not None else "")
                if bound is not None:
                    add(carry.strip())
                    carry = ""
                    add(bound)
                elif content is not None:
                    walk_container(content)
            elif child.tag == _CUSTOM_XML:
                walk_container(child)
            elif child.tag == _PARAGRAPH:
                text = accepted_text(child, walk)
                if child.find(_DELETED_MARK) is not None:
                    carry += text
                    continue
                text = (carry + text).strip()
                carry = ""
                if text:
                    paragraph = Paragraph(cast(CT_P, child), document)
                    level = _heading_level(paragraph)
                    if level and styled:
                        top = _styled_top(text, printed.get(child))
                        last_number = top if top is not None else last_number
                    elif styled:
                        level = _typed_heading_level(paragraph, text, printed.get(child))
                    else:
                        level = _unstyled_heading_level(text, printed.get(child), bool(blocks))
                    add(text, level, printed.get(child) if level > 0 else None)
                read_boxes()
            elif child.tag == _TABLE:
                add(carry.strip())
                carry = ""
                for row in table_rows(child):
                    add(row)
                read_boxes()
            elif child.tag not in _STRUCTURE:
                walk.skip(child)

    walk_container(body)
    add(carry.strip())
    tracked = len(body.xpath(_REVISIONS_XPATH))
    counts = DocxWalkCounts(
        tables=walk.tables,
        content_controls=walk.controls,
        form_fields=walk.fields.forms,
        fields_balanced=walk.fields.enabled,
        bound_read=walk.bindings.read,
        bound_unknown=walk.bindings.unknown,
        symbols=walk.symbols,
        unknown_symbols=walk.unknown_symbols,
        text_boxes=walk.text_boxes,
        drawn=walk.drawn,
        invisible_runs=walk.invisible_runs,
        unread=tuple(sorted(walk.unread.items())),
    )
    return ReadResult(
        blocks=blocks,
        pages=_docx_pages(data),
        title=title,
        tracked_changes=tracked,
        hidden_runs=walk.hidden_runs + walk.style_hidden_runs,
        format_characters=removed,
        coverage=docx_coverage(data, etree.tostring(body, encoding="unicode"), tracked, walk.hidden_runs, walk.style_hidden_runs, counts),
        more_coverage=docx_more_coverage(counts),
    )


def _bindable_parts(data: bytes) -> dict[str, bytes]:
    """The parts a content control may be bound to: the custom XML items and their properties, and the document's core
    and extended properties."""
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        names = [n for n in package.namelist() if n.startswith("customXml/") or n in ("docProps/core.xml", "docProps/app.xml")]
        return {name: package.read(name) for name in names}


def _field_marks(body: etree._Element, compat: Compatibility) -> list[tuple[str | None, bool]]:
    """Each w:fldChar of the body Word displays, in document order, with whether a tracked deletion removes it: not one in
    a branch of an mc:AlternateContent Word does not display."""
    marks: list[tuple[str | None, bool]] = []
    for mark in body.iter(_FIELD_CHAR):
        deleted, shown = False, True
        for ancestor in mark.iterancestors():
            if ancestor.tag in _DELETED:
                deleted = True
            elif ancestor.tag in (CHOICE, FALLBACK):
                parent = ancestor.getparent()
                shown = shown and parent is not None and parent.tag == ALTERNATE and compat.branch(parent) is ancestor
        if shown:
            marks.append((mark.get(qn("w:fldCharType")), deleted))
    return marks


# ---------------------------------------------------------------------------- PDF


def read_pdf(data: bytes) -> ReadResult:
    # A damaged file fails anywhere below; `read` refuses every such failure as an unreadable PDF.
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise UnsupportedFile("the PDF is encrypted; remove the password and upload it again")
    page_count = len(reader.pages)
    if page_count > MAX_PDF_PAGES:
        raise TooLargeToRead(f"the PDF has {page_count} pages; the limit is {MAX_PDF_PAGES}")
    pages = list(reader.pages)
    lines: list[str] = []
    line_pages: list[int] = []  # the page each line was printed on: page furniture is a line printed on most of them
    textless = 0
    removed = 0
    tally = PdfTally()
    for number, page in enumerate(pages):
        # Each page is read from its own drawing operations, laid out as pypdf's layout mode lays out lines, which keeps
        # the blank lines between paragraphs, the only structure most PDFs carry (reader v16, ingest/pdf_layout.py:
        # columns, form XObjects, glyph order, text that cannot be seen).
        try:
            shown = page_text(page)
            text, drawn = shown.text, shown.drawn
            tally.add(shown)
        except Exception:  # noqa: BLE001 - a page the walk cannot finish is read as pypdf's layout mode reads it, and counted
            tally.as_given += 1
            try:
                text = page.extract_text(extraction_mode="layout") or ""
            except Exception:  # noqa: BLE001 - then the plain extractor, as before v16
                text = page.extract_text() or ""
            drawn = bool(text.strip())
        tally.annotations += _annotation_texts(page)
        text, count = remove_format_characters(text)
        removed += count
        if not text.strip() and not drawn:
            textless += 1  # an image-only (scanned) page: the record says so instead of "unknown"
        lines.extend(re.sub(r"[ \t]{2,}", " ", line) for line in text.splitlines())
        lines.append("")
        line_pages.extend([number] * (len(lines) - len(line_pages)))
    # A page's running header and footer are left out of the reading and counted, as a Word file's are (reader v15).
    furniture = page_edge_furniture(lines, line_pages)
    headers = [lines[i].strip() for i in sorted(furniture.headers)]
    footers = [lines[i].strip() for i in sorted(furniture.footers)]
    lines = ["" if i in furniture.headers or i in furniture.footers else line for i, line in enumerate(lines)]
    return ReadResult(
        coverage=pdf_coverage(page_count, textless, headers, footers, tally),
        blocks=_blocks_from_lines(lines, line_pages),
        pages=len(pages),
        title=_first_line(lines),
        format_characters=removed,
    )


def _annotation_texts(page: PageObject) -> int:
    """How many annotations on the page show text the reading does not hold (reader v16): a form field with a value (a
    filled text field, a ticked box) and a free-text note or stamp, unless the annotation is flagged hidden. Neither of
    pypdf's extractions reads them: "Initial Term:" was kept and its filled-in "thirty-six (36) months" dropped."""
    count = 0
    try:
        annotations = page.get("/Annots")
        for reference in annotations.get_object() if annotations is not None else []:
            annotation = reference.get_object()
            if int(annotation.get("/F", 0)) & (2 | 32):  # Hidden, NoView
                continue
            subtype = annotation.get("/Subtype")
            if subtype == "/Widget":
                value = annotation.get("/V")
                parent = annotation.get("/Parent")
                if value is None and parent is not None:
                    value = parent.get_object().get("/V")
                shown = annotation.get("/AS")
                if (isinstance(value, str) and value.strip()) or (shown is not None and str(shown) != "/Off"):
                    count += 1
            elif subtype in ("/FreeText", "/Stamp") and str(annotation.get("/Contents") or "").strip():
                count += 1
    except Exception:  # noqa: BLE001 - annotations this reader cannot walk are not counted; the page's own text is read
        return count
    return count


# ---------------------------------------------------------------------------- TXT


_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def decode_text(data: bytes) -> tuple[str, str]:
    """The text of a plain-text file and the name of the encoding it was read with. A byte-order mark decides; then
    UTF-16 when every other byte is a NUL; then strict UTF-8; then Windows-1252, the encoding of text saved from Word
    on a Western-locale machine. Before 2026-10-01 everything was read as UTF-8 with replacement characters, so a
    UTF-16 file became NUL-interleaved garbage with a 201 and "€1,500" became "\ufffd1,500"."""
    if data.startswith(codecs.BOM_UTF8):
        return data[len(codecs.BOM_UTF8) :].decode("utf-8", errors="replace"), "UTF-8"
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16", errors="replace"), "UTF-16"
    sample = data[:4096]
    if sample.count(b"\x00") > len(sample) // 4:
        # Every other byte a NUL: UTF-16 without its mark, which strict UTF-8 would accept as NUL-interleaved text.
        little_endian = sample[1:2] == b"\x00"
        return data.decode("utf-16-le" if little_endian else "utf-16-be", errors="replace"), "UTF-16"
    try:
        return data.decode("utf-8"), "UTF-8"
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace"), "Windows-1252"


def read_txt(data: bytes) -> ReadResult:
    text, encoding = decode_text(data)
    if "\x00" in text:
        # Text never carries NUL once decoded (UTF-16 is decoded above); binary does, and was read as text before
        # 2026-10-08, its control characters stripped and the rest kept as clauses.
        raise UnsupportedFile("the file is not plain text: it holds binary data")
    text, removed = remove_format_characters(text)
    lines = _CONTROL_CHARS.sub("", text).splitlines()
    return ReadResult(coverage=txt_coverage(encoding), blocks=_blocks_from_lines(lines), pages=None, title=_first_line(lines), format_characters=removed)


def _first_line(lines: list[str]) -> str:
    for line in lines:
        if line.strip():
            return shorten_title(line.strip())
    return ""


def shorten_title(text: str) -> str:
    """The text as a title: whole when it is a line, else its first words and an ellipsis."""
    if len(text) <= MAX_TITLE_CHARS:
        return text
    cut = text.rfind(" ", MAX_TITLE_CHARS // 2, MAX_TITLE_CHARS)
    return text[: cut if cut > 0 else MAX_TITLE_CHARS].rstrip() + "…"


# Reader v16: a line at least this share of the length of the line after it is one of a paragraph's wrapped lines.
WRAPPED_LINE_SHARE = 0.6
MAX_TITLE_CONTINUATION_WORDS = 6


def _continues_title(line: str, heading: str, candidate: bool) -> bool:
    """Whether a capitals line run straight on from a heading is that heading's next line: short, stating nothing, not
    stopping on a comma, and either after a heading that goes on ("REPRESENTATIONS AND") or no title of its own (it opens
    with a function word, "OF LOAN REQUEST", or stops on a semicolon, "CONTRIBUTIONS; CAPITAL ACCOUNTS;")."""
    if not shouting(line) or len(line.split()) > MAX_TITLE_CONTINUATION_WORDS or line.rstrip().endswith(",") or states(line):
        return False
    return goes_on(heading) or not candidate


# A line that names an Article, a Part or a Chapter and nothing else: "ARTICLE II", "Article 2", "PART IV".
_PART_LINE = re.compile(r"^(?:ARTICLE|Article|PART|Part|CHAPTER|Chapter)\s+(?:[IVXLC]{1,7}|\d{1,3})\.?$")


def _blocks_from_lines(lines: list[str], pages: list[int] | None = None) -> list[Block]:
    """Paragraphs are separated by blank lines; short numbered or all-caps lines are headings. ``pages`` holds the page
    each line was printed on, for a format that has pages (PDF, which ends each page with a blank line); every block
    carries the page of its first line."""
    blocks: list[Block] = []
    paragraph: list[str] = []
    paragraph_page: int | None = None

    def flush() -> None:
        if paragraph:
            blocks.append(Block("text", " ".join(paragraph), page=paragraph_page))
            paragraph.clear()

    seen_heading = False
    last_number = 0  # the top-level number of the last numbered heading: clause numbers count on from it
    after_heading = False  # the line before was a heading's, with no blank line between
    for index, raw in enumerate(lines):
        line = raw.strip()
        page = pages[index] if pages is not None else None
        if not line:
            flush()
            after_heading = False
            continue
        if not paragraph:
            paragraph_page = page
        candidate = looks_like_heading(line)
        unnumbered = HEADING_NUMBER.match(line) is None and not _PART_LINE.match(line)
        if after_heading and unnumbered and blocks and blocks[-1].kind == "heading" and _continues_title(line, blocks[-1].text, candidate):
            # A title that wraps is one heading (reader v16): "ARTICLE IV REPRESENTATIONS AND" / "WARRANTIES", "RESOLUTION
            # OF" / "CLAIMS", "NOTICE" / "OF LOAN REQUEST", "ARTICLE V CAPITAL" / "CONTRIBUTIONS; CAPITAL ACCOUNTS;".
            blocks[-1] = Block("heading", f"{blocks[-1].text} {line}", blocks[-1].level, implicit_number=False, page=blocks[-1].page)
            continue
        if candidate and unnumbered:
            # Reader v16: an unnumbered title line inside a paragraph is the paragraph's next line, where the line
            # before it goes on (it is in capitals and ends no passage, or it stops on a comma or a function word): a
            # wrapped capitals clause's last line ("IN BREACH OF ITS PAYMENT OBLIGATIONS.") was a heading, and the
            # condition it carried left the quotable text; so were "GAAP." after "in accordance with" and a party's
            # name after "between" (stress test, 2026-10-30).
            last = paragraph[-1] if paragraph else ""
            if last and not PASSAGE_ENDS.search(last) and (shouting(last) or goes_on(last)):
                paragraph.append(line)
                after_heading = False
                continue
            following = lines[index + 1].strip() if index + 1 < len(lines) else ""
            wrapped = following and shouting(following) and not looks_like_heading(following) and len(line) >= WRAPPED_LINE_SHARE * len(following)
            if wrapped and not PASSAGE_ENDS.search(line):
                # A capitals paragraph's first line, run straight into a capitals line that is no title and about as
                # long ("EXCEPT FOR A BREACH OF SECTION" / "9, IN NO EVENT WILL EITHER"): the paragraph's, not a
                # heading. A title is shorter than the line of body after it ("CAPITAL" / "CONTRIBUTIONS; ...").
                flush()
                paragraph_page = page
                paragraph.append(line)
                after_heading = False
                continue
        after_heading = False
        if candidate:
            numbered = HEADING_NUMBER.match(line)
            if numbered and not continues_numbering(numbered.group(1), last_number):
                # "15285 Minnetonka Blvd.", "2024 Fee Schedule", "4552374 CUDDL DUDS": a number that does not continue the
                # clause numbering is an address, a year or a registration number, part of the text around it.
                paragraph.append(line)
                continue
            flush()
            if not numbered and not seen_heading and not blocks:
                # An all-caps first line is the title, not clause 1.
                blocks.append(Block("text", line, page=page))
                continue
            seen_heading = True
            if not numbered and blocks and blocks[-1].kind == "heading" and _PART_LINE.match(blocks[-1].text) and not _PART_LINE.match(line):
                # "ARTICLE II" on one line and "PURCHASE AND SALE" on the next are one heading: apart, the first was an
                # empty section and the Article's clauses sat under a title with no number (retrieval study, 2026-10-14).
                blocks[-1] = Block("heading", f"{blocks[-1].text} {line}", blocks[-1].level, implicit_number=False, page=blocks[-1].page)
                after_heading = True
                continue
            if numbered:
                last_number = int(numbered.group(1).split(".")[0])
            # A plain-text or PDF heading carries its number in its text or has none: no format numbers it implicitly,
            # so a number computed for it would be a number the document never shows (QA campaign, 2026-10-08).
            blocks.append(Block("heading", line, 1 if "." not in line.split(" ", 1)[0].rstrip(".") else 2, implicit_number=False, page=page))
            after_heading = True
            continue
        paragraph.append(line)
    flush()
    return blocks
